"""Create synthetic JET inputs. No client or legacy files are changed."""
from pathlib import Path
import json
from openpyxl import Workbook

root = Path(__file__).resolve().parent
wb = Workbook()
ws = wb.active
ws.title = "Synthetic"
ws.append(["ProbeId", "公司", "凭证日期", "凭证编号", "科目编码", "科目名称", "借方金额", "贷方金额", "摘要", "制单人", "审核人", "测试数值", "测试文本"])
values = [
    (0, None), (1, ""), (2, "调整"), (3, "回款调整"), (4, "回款"),
    (5, "  调整  "), (6, "6001"), (7, "6051"), (8, "7001"), (9, "1002"),
    (10, "2001"), (11, "CFO"), (12, "cfo"), (13, "[空]"), (14, "nan"),
    (15, "None"), (16, "  "), (17, "1,000"), (18, "1000"), (19, "1000.0"),
]
amounts = [None, 0, -1, 1, 6, 7, 8, 22, 23, 1000, "1,000", "abc", "1000", "", 0.01, -0.01, 1000000, 1000001, "NaN", " 7 "]
for (index, value), amount in zip(values, amounts):
    ws.append([f"P{index:02}", "合成测试公司", "2026-12-31", f"T{index:02}", "6001", "测试科目", 0, 0, "合成数据", "甲", "乙", amount, value])
wb.save(root / "synthetic.xlsx")
for row in range(2,ws.max_row+1):
    ws.cell(row,7,amounts[row-2])
    ws.cell(row,9,values[row-2][1])
wb.save(root / "synthetic-mcp.xlsx")
(root / "synthetic-values.json").write_text(json.dumps([
    {"id": f"P{index:02}", "测试数值": amount, "测试文本": value}
    for (index, value), amount in zip(values, amounts)
], ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Wrote {len(values)} synthetic rows")
