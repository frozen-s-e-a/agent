"""Isolated, bounded audit executor. Every output is new and published with a manifest."""
from __future__ import annotations
import csv, hashlib, json, os, re, shutil, sys, time, zipfile, uuid
from pathlib import Path
from decimal import Decimal, InvalidOperation
from datetime import datetime, date
import duckdb
from openpyxl import load_workbook, Workbook

BASE=Path(__file__).resolve().parents[1]
META=['_source_file','_source_sheet','_source_row','_row_id']
def emit(kind,**payload): print(json.dumps({'type':kind,**payload},ensure_ascii=False,default=str),flush=True)
def digest(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def inside(root,value):
    p=(root/value).resolve(strict=True)
    if not p.is_relative_to(root):raise ValueError('文件超出项目授权目录')
    return p
def qi(value): return '"'+str(value).replace('"','""')+'"'
def lit(value): return "'"+str(value).replace("'","''")+"'"
def names(value): return [x.strip() for x in re.split('[,，]',value or '') if x.strip()]
def val(x):
    if x is None:return None
    if isinstance(x,(datetime,date)):return x.isoformat()
    return str(x)
def tables(p,params):
    if p.suffix.lower() in ['.csv','.tsv']:
        with p.open('r',encoding=params.get('encoding','utf-8-sig'),newline='') as f:
            rows=csv.reader(f,delimiter='\t' if p.suffix.lower()=='.tsv' else ',',strict=True)
            for _ in range(int(params.get('headerRow',1))-1):next(rows,None)
            yield p.name,rows
    elif p.suffix.lower() in ['.xlsx','.xlsm']:
        wb=load_workbook(p,read_only=True,data_only=False,keep_links=False)
        try:
            chosen=names(params.get('sheets',''))
            if any(s not in wb.sheetnames for s in chosen):raise ValueError('指定的工作表不存在')
            for ws in wb:
                if chosen and ws.title not in chosen:continue
                ws.reset_dimensions()
                yield ws.title,ws.iter_rows(min_row=int(params.get('headerRow',1)),values_only=True)
        finally:wb.close()
    else:raise ValueError('此表格工具支持 CSV、TSV、XLSX、XLSM；其他格式仍待迁移。')

def import_tables(db,files,root,params):
    columns=[];count=0;db.execute('CREATE TABLE data (_source_file VARCHAR,_source_sheet VARCHAR,_source_row BIGINT,_row_id BIGINT)')
    for p in files:
        published=False
        manifest_path=p.parent/'manifest.json'
        if p.is_relative_to(root/'outputs') and manifest_path.is_file():
            manifest=json.loads(manifest_path.read_text(encoding='utf-8'))
            output=next((x for x in manifest.get('outputs',[]) if x.get('name')==p.name),None)
            if output and manifest.get('status')=='succeeded':
                if digest(p)!=output.get('sha256'):raise ValueError('已发布成果校验失败，不能继续处理被修改的结果')
                published=True
        for sheet,rows in tables(p,params):
            header=next(rows,None)
            if header is None:continue
            header=[str(v).strip() if v is not None else f'未命名列{i+1}' for i,v in enumerate(header)]
            original_header=header[:]
            indices=list(range(len(header)))
            reserved=any(x in META for x in header)
            if len(set(header))!=len(header) or any(x.startswith('__f') for x in header):raise ValueError('存在重名表头或保留的来源列，请先确认表头行')
            if reserved:
                if not published or not set(META).issubset(header):raise ValueError('存在保留的来源列；只有校验通过的已发布成果可以继续处理')
                # Recreate provenance against the immediate input. The parent result's
                # manifest and table preserve the preceding step of the lineage.
                indices=[i for i,name in enumerate(header) if name not in META]
                header=[original_header[i] for i in indices]
            if len(header)>512:raise ValueError('工作表超过 512 列，当前适配器拒绝超宽输入')
            for col in header:
                if col not in columns:db.execute(f'ALTER TABLE data ADD COLUMN {qi(col)} VARCHAR');columns.append(col)
            cols=META+header;sql='INSERT INTO data ('+','.join(map(qi,cols))+') VALUES ('+','.join('?' for _ in cols)+')';batch=[];batch_rows=max(10,min(1000,12000//len(cols)))
            for rownum,row in enumerate(rows,start=int(params.get('headerRow',1))+1):
                if len(row)>len(original_header) and any(v is not None and str(v)!='' for v in row[len(original_header):]):raise ValueError(f'{p.name}/{sheet}/{rownum}：数据列超出表头，未截断')
                values=[val(row[i]) if i<len(row) else None for i in indices]
                count+=1;batch.append([p.relative_to(root).as_posix(),sheet,rownum,count]+values)
                if len(batch)>=batch_rows:
                    prefix,_=sql.rsplit(' VALUES ',1)
                    # Every value is NULL or a quoted SQL literal. This avoids per-value
                    # optional pandas/numpy import probes in the minimal DuckDB runtime.
                    db.execute(prefix+' VALUES '+','.join('('+','.join('NULL' if v is None else lit(v) for v in r)+')' for r in batch))
                    batch=[];emit('progress',phase='导入数据',rows=count)
            if batch:
                prefix,_=sql.rsplit(' VALUES ',1)
                db.execute(prefix+' VALUES '+','.join('('+','.join('NULL' if v is None else lit(v) for v in r)+')' for r in batch))
    if not columns:raise ValueError('没有可读取的表头或数据')
    return columns,count

def require(columns,needed):
    missing=[x for x in needed if x not in columns]
    if missing:raise ValueError('找不到字段：'+', '.join(missing))
def decimal_column(db,column):
    # Reject bad values, excessive precision and overflow rather than silently coercing to zero.
    c=qi(column);raw=f"replace(trim({c}), ',', '')"
    invalid=db.execute(f"SELECT _source_file,_source_row,{c} FROM data WHERE coalesce({c},'')<>'' AND (NOT regexp_full_match({raw},'[+-]?[0-9]+(\\.[0-9]{{1,6}})?') OR try_cast({raw} AS DECIMAL(38,6)) IS NULL) LIMIT 1").fetchone()
    if invalid:raise ValueError(f'金额字段 {column} 无效或超过 6 位小数，位置：{invalid[0]} 第 {invalid[1]} 行')
    return f"coalesce(cast(nullif({raw},'') AS DECIMAL(38,6)),0::DECIMAL(38,6))"
def numeric_param(value,minimum='0'):
    try:d=Decimal(str(value))
    except InvalidOperation:raise ValueError('数值参数无效')
    if not d.is_finite() or d<Decimal(minimum):raise ValueError('数值参数超出范围')
    return format(d,'f')
def date_expr(column):
    return f"try_cast(substr(replace({qi(column)}, '/', '-'),1,10) AS DATE)"

def run_sql(db,tool,p,columns):
    cols=names(p.get('columns'));keys=names(p.get('keys'));allcols=','.join(map(qi,columns));metadata=','.join(map(qi,META))
    if tool in ['select_column','fill_column','subject_clean','merge_column']:
        if not cols:raise ValueError('至少选择一个字段')
        require(columns,cols)
    if tool in ['duplicates','voucher_check']:
        if not keys:raise ValueError('请指定完整分组键')
        require(columns,keys)
    if tool in ['add_column','merge_column']:
        if not p.get('target') or p['target'] in columns+META:raise ValueError('新列名为空或已存在')
    if tool in ['inspect','merge_files']:return 'SELECT * FROM data ORDER BY _row_id'
    if tool=='select_column':return f'SELECT {metadata},'+','.join(map(qi,cols))+' FROM data ORDER BY _row_id'
    if tool=='subject_clean':return f'SELECT {metadata},'+','.join(f'trim({qi(c)}) AS {qi(c)}' if c in cols else qi(c) for c in columns)+' FROM data ORDER BY _row_id'
    if tool=='fill_column':
        expr=[f"last_value(nullif({qi(c)},'') IGNORE NULLS) OVER(PARTITION BY _source_file,_source_sheet ORDER BY _row_id ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) AS {qi(c)}" if c in cols else qi(c) for c in columns]
        return f'SELECT {metadata},'+','.join(expr)+' FROM data ORDER BY _row_id'
    if tool=='merge_column':return f'SELECT *,concat_ws({lit(p.get("separator","_"))},'+','.join(map(qi,cols))+f') AS {qi(p["target"])} FROM data ORDER BY _row_id'
    if tool=='add_column':return f'SELECT *,{lit(p.get("value",""))} AS {qi(p["target"])} FROM data ORDER BY _row_id'
    if tool in ['dropsummary','text_match']:
        require(columns,[p.get('column')]);c=qi(p['column'])
        if tool=='dropsummary':
            vs=names(p.get('values'))
            if not vs:raise ValueError('请指定合计行匹配值')
            return f"SELECT * FROM data WHERE coalesce({c},'') NOT IN ("+','.join(map(lit,vs))+') ORDER BY _row_id'
        return f"SELECT * FROM data WHERE contains(coalesce({c},''),{lit(p.get('keyword',''))}) ORDER BY _row_id"
    if tool=='duplicates':return 'SELECT * EXCLUDE(重复次数) ,重复次数 FROM (SELECT *,count(*) OVER(PARTITION BY '+','.join(map(qi,keys))+') AS 重复次数 FROM data) WHERE 重复次数>1 ORDER BY _row_id'
    if tool in ['voucher_check','monthly_analysis']:
        require(columns,[p.get('debit'),p.get('credit')]);debit=decimal_column(db,p['debit']);credit=decimal_column(db,p['credit'])
        if tool=='voucher_check':
            tolerance=numeric_param(p.get('tolerance','0.01'));group=','.join(map(qi,keys))
            return f'SELECT {group},sum({debit}) AS 借方合计,sum({credit}) AS 贷方合计,sum({debit})-sum({credit}) AS 差额,count(*) AS 分录数,min(_row_id) AS 首条记录 FROM data GROUP BY {group} HAVING abs(sum({debit})-sum({credit}))>{tolerance} ORDER BY {group}'
        require(columns,[p.get('date'),p.get('subject')]);subjects=names(p.get('subjects'))
        if not subjects:raise ValueError('请确认需要分析的科目，不能默认分析全部')
        dx=date_expr(p['date'])
        if db.execute(f'SELECT count(*) FROM data WHERE {dx} IS NULL').fetchone()[0]:raise ValueError('存在无效日期，请先清理；未排除这些数据')
        return f"SELECT {qi(p['subject'])} AS 科目代码,strftime({dx},'%Y-%m') AS 月份,sum({debit}) AS 借方合计,sum({credit}) AS 贷方合计,count(*) AS 分录数 FROM data WHERE {qi(p['subject'])} IN ("+','.join(map(lit,subjects))+') GROUP BY 1,2 ORDER BY 1,2'
    if tool=='jet_test':return jet_query(db,p,columns)
    raise ValueError('尚未实现该工具')

def condition(expression,rule):
    parts=names(str(rule));conditions=[]
    for s in parts:
        if s=='[空]':c=f"coalesce(cast({expression} AS VARCHAR),'')=''"
        elif s=='![空]':c=f"coalesce(cast({expression} AS VARCHAR),'')<>''"
        elif re.fullmatch(r'(>=|<=|>|<|=)[+-]?\d+(\.\d+)?',s):
            m=re.match(r'(>=|<=|>|<|=)(.*)',s);c=f'try_cast({expression} AS DECIMAL(38,6)){m[1]}{m[2]}'
        elif s.startswith('!^'):c=f"NOT starts_with(coalesce(cast({expression} AS VARCHAR),''),{lit(s[2:])})"
        elif s.startswith('^'):c=f"starts_with(coalesce(cast({expression} AS VARCHAR),''),{lit(s[1:])})"
        elif s.startswith('!'):c=f"NOT contains(coalesce(cast({expression} AS VARCHAR),''),{lit(s[1:])})"
        elif s.startswith('#'):c=f'cast({expression} AS VARCHAR)={lit(s[1:])}'
        elif re.fullmatch(r'[+-]?\d+(\.\d+)?',s):c=f'try_cast({expression} AS DECIMAL(38,6))={s}'
        else:c=f"contains(coalesce(cast({expression} AS VARCHAR),''),{lit(s)})"
        conditions.append('('+c+')')
    # Negative alternatives in the original template mean exclusion of every prefix.
    join=' AND ' if all(x.startswith('!') for x in parts) else ' OR '
    return '('+join.join(conditions)+')'

def jet_query(db,p,columns):
    rules=json.loads((BASE/'resources'/'jet-rules.json').read_text(encoding='utf-8'))['rules']
    selected=names(p.get('rules'));mapping=p.get('mapping',{})
    if not selected:raise ValueError('请明确选择 JET 规则')
    unknown=set(selected)-{r['name'] for r in rules}
    if unknown:raise ValueError('未知 JET 规则：'+','.join(unknown))
    def col(name):
        c=mapping.get(name,name);require(columns,[c]);return qi(c)
    def money(name):return decimal_column(db,mapping.get(name,name)) if col(name) else ''
    features={};needed={k for r in rules if r['name'] in selected for k in r['conditions']}
    for k in needed:
        if k=='凭证月份':e=f'month({date_expr(mapping.get("凭证日期","凭证日期"))})';col('凭证日期')
        elif k=='跨月编制':e=f"date_diff('month',try_cast({col('凭证日期')} AS DATE),try_cast({col('制单日期')} AS DATE))"
        elif k=='制单小时数':e=f"coalesce(hour(try_cast({col('制单时间')} AS TIMESTAMP)),hour(try_cast({col('制单时间')} AS TIME)))"
        elif k=='是否周末':e=f"CASE WHEN dayofweek(try_cast({col('凭证日期')} AS DATE)) IN (0,6) THEN '是' ELSE '否' END"
        elif k in ['借方余数','贷方余数']:e=f"({money('借方金额' if k=='借方余数' else '贷方金额')} % 10000)"
        elif k=='尾数后3位':e=f"right(cast(cast(greatest(abs({money('借方金额')}),abs({money('贷方金额')})) AS BIGINT) AS VARCHAR),3)"
        elif k=='摘要长度':e=f"length(coalesce({col('摘要')},''))"
        elif k=='凭证差额':
            ks=names(p.get('keys'))
            if not ks:raise ValueError('借贷不平规则须在高级参数 keys 指定含主体、期间的完整凭证键')
            require(columns,ks);e=f"abs(sum({money('借方金额')}-{money('贷方金额')}) OVER(PARTITION BY "+','.join(map(qi,ks))+'))'
        elif k in ['科目频率','制单人频率']:e=f"count(*) OVER(PARTITION BY {col('科目编码' if k=='科目频率' else '制单人')})"
        elif k=='行重复':e="CASE WHEN count(*) OVER(PARTITION BY "+','.join(map(qi,columns))+ ")>1 THEN '是' ELSE '否' END"
        elif k=='制单审核一致':e=f"CASE WHEN {col('制单人')}={col('审核人')} THEN '是' ELSE '否' END"
        elif k in ['借方金额','贷方金额']:e=money(k)
        else:e=col(k)
        features[k]=e
    db.execute('CREATE VIEW featured AS SELECT *,'+','.join(f'{e} AS {qi("__f"+str(i))}' for i,e in enumerate(features.values()))+' FROM data')
    expressions={k:qi('__f'+str(i)) for i,k in enumerate(features)}
    queries=[]
    for r in rules:
        if r['name'] not in selected:continue
        cond=r['conditions'].copy();cond.update(p.get('ruleOverrides',{}).get(r['name'],{}))
        if set(cond)-set(expressions):raise ValueError('规则覆盖不得添加未声明特征')
        queries.append('SELECT '+','.join(map(qi,META+columns))+f',{lit(r["name"])} AS 命中规则 FROM featured WHERE '+' AND '.join(condition(expressions[k],v) for k,v in cond.items()))
    return 'SELECT * FROM ('+' UNION ALL '.join(queries)+') ORDER BY _row_id,命中规则'

def export_result(db,query,stage):
    db.execute('CREATE TABLE result AS '+query)
    cur=db.execute('SELECT * FROM result');headers=[x[0] for x in cur.description];count=0;preview=[];files=[];part=0;handle=None;writer=None;excel_safe=True;preview_bytes=0
    try:
        while True:
            batch=cur.fetchmany(1000)
            if not batch:break
            for row in batch:
                if count%100000==0:
                    if handle:handle.close()
                    part+=1;name=f'结果-{part:03d}.csv';handle=(stage/name).open('w',encoding='utf-8-sig',newline='');writer=csv.writer(handle);writer.writerow(headers);files.append(name)
                values=[val(v) for v in row]
                if any(v and len(v)>32767 for v in values):excel_safe=False
                if len(preview)<50:
                    bounded=[v[:500] if v is not None else None for v in values]
                    size=len(json.dumps(bounded,ensure_ascii=False).encode('utf-8'))
                    if preview_bytes+size<=256000:preview.append(bounded);preview_bytes+=size
                # Spreadsheet formula injection is escaped in data exports only.
                safe=[("'"+v) if v and re.match(r'^[\s]*[=+@-]',v) and not re.fullmatch(r'[+-]?\d+(\.\d+)?',v) else v for v in values]
                writer.writerow(safe);count+=1
            emit('progress',phase='导出结果',rows=count)
        if not files:
            name='结果-001.csv'
            with (stage/name).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerow(headers)
            files.append(name)
    finally:
        if handle:handle.close()
    if count<=50000 and excel_safe:
        wb=Workbook(write_only=True);ws=wb.create_sheet('结果');ws.append(headers)
        from openpyxl.cell import WriteOnlyCell
        cur=db.execute('SELECT * FROM result')
        while True:
            batch=cur.fetchmany(1000)
            if not batch:break
            for row in batch:
                cells=[]
                for v in row:
                    text=val(v)
                    if text and len(text)>32767:raise ValueError('结果包含超过 Excel 单元格限制的文本，请使用支持该格式的工具')
                    numeric=isinstance(v,(int,float,Decimal)) and len(str(v).replace('.','').replace('-','').lstrip('0'))<=15
                    cell=WriteOnlyCell(ws,value=v if numeric else text)
                    if not numeric:cell.data_type='s'
                    cells.append(cell)
                ws.append(cells)
        wb.save(stage/'结果.xlsx');files.append('结果.xlsx')
    db.execute(f"COPY result TO {lit(str(stage/'完整结果.parquet'))} (FORMAT PARQUET)");files.append('完整结果.parquet')
    return {'columns':headers,'preview':preview,'previewOnly':True,'previewCellLimit':500,'rowCount':count,'files':files,'excelSkippedForLongText':not excel_safe}

def run(request):
    schema=json.loads((BASE/'contracts'/'task.schema.json').read_text(encoding='utf-8'))
    if any(k not in request for k in schema['required']) or request['protocolVersion']!=1:raise ValueError('执行协议不匹配')
    if not re.fullmatch(schema['properties']['taskId']['pattern'],request['taskId']):raise ValueError('任务编号无效')
    if request['mode'] not in schema['properties']['mode']['enum']:raise ValueError('处理模式无效')
    root=Path(request['root']).resolve(strict=True);files=[inside(root,x) for x in request['files']]
    if not files or len(files)>1000 or len(set(files))!=len(files):raise ValueError('文件列表为空、重复或超出上限')
    p=request['parameters'];tool=request['tool'];memory=int(request['limits']['memoryMiB']);threads=int(request['limits']['threads'])
    layout_inspect=tool=='inspect' and any(f.suffix.lower() in ['.xlsx','.xlsm'] for f in files)
    rule_version='inspect-layout-1' if layout_inspect else 'internal-1'
    if not 256<=memory<=8192 or not 1<=threads<=8:raise ValueError('资源预算无效')
    if not 1<=int(p.get('headerRow',1))<=10000:raise ValueError('表头行超出范围')
    if tool=='__structure':
        from file_structure import structure
        return structure(files[0],p)
    if tool=='__attachment':
        from attachment_extract import extract_attachment
        return extract_attachment(files[0])
    if tool=='__preview':
        if files[0].suffix.lower()=='.parquet':
            connection=duckdb.connect(':memory:')
            try:
                connection.execute("SET memory_limit='128MB'")
                cursor=connection.execute('SELECT * FROM read_parquet(?) LIMIT 10',[str(files[0])])
                if len(cursor.description)>512:raise ValueError('结果超过预览列数预算')
                return {'columns':[str(c[0]) for c in cursor.description],'preview':[[val(x)[:200] if x is not None else None for x in row] for row in cursor.fetchall()],'previewOnly':True}
            finally:connection.close()
        if files[0].suffix.lower() in ['.pdf','.docx']:
            from attachment_extract import extract_attachment
            return {**extract_attachment(files[0]),'previewOnly':True}
        for sheet,rows in tables(files[0],p):
            header=next(rows,None)
            if header is None:continue
            import itertools
            if len(header)>512:raise ValueError('工作表超过预览列数预算')
            return {'columns':[str(x)[:200] if x is not None else f'未命名列{i+1}' for i,x in enumerate(header)],'sheet':sheet,'preview':[[val(x)[:200] if x is not None else None for x in row[:512]] for row in itertools.islice(rows,10)],'previewOnly':True}
        raise ValueError('文件中没有可预览的数据')
    total=sum(f.stat().st_size for f in files);expanded=total
    for f in files:
        if f.suffix.lower() in ['.xlsx','.xlsm']:
            with zipfile.ZipFile(f) as z:expanded+=sum(i.file_size for i in z.infolist())
    estimate=max(total*12,expanded*4);light_limit=min(512*1024**2,memory*1024**2//4)
    mode=request['mode'];reason='用户指定'
    if mode=='auto':mode='local-batch' if estimate>light_limit or len(files)>5 else 'local-light';reason=f'估算展开与处理中间数据 {estimate//1024**2} MiB，轻量预算 {light_limit//1024**2} MiB'
    if mode=='local-light' and estimate>light_limit:mode='local-batch';reason='估算超过轻量预算，已转为大批量'
    if expanded>memory*1024**2*4 and any(f.suffix.lower() in ['.xlsx','.xlsm'] for f in files):raise ValueError('Excel 解压体量超过此适配器预算，请先转换为 CSV')
    output=root/'outputs';output.mkdir(exist_ok=True)
    if not output.resolve().is_relative_to(root):raise ValueError('输出目录指向项目外部')
    final=output/request['taskId'];stage=output/('.'+request['taskId']+'.'+uuid.uuid4().hex+'.staging')
    snapshots=[{'path':f.relative_to(root).as_posix(),'sha256':digest(f),'size':f.stat().st_size} for f in files]
    signature=hashlib.sha256(json.dumps({'tool':tool,'parameters':p,'inputs':snapshots,'ruleVersion':rule_version},sort_keys=True).encode()).hexdigest()
    if final.exists():
        if not final.resolve().is_relative_to(root):raise ValueError('结果目录指向项目外部')
        manifest=json.loads((final/'manifest.json').read_text(encoding='utf-8'))
        if manifest['signature']!=signature:raise ValueError('任务输入版本已改变，不能复用旧输出')
        for x in manifest['outputs']:
            if digest(inside(final,x['name']))!=x['sha256']:raise ValueError('已发布结果校验失败')
        return manifest
    stage.mkdir(exist_ok=False)
    if shutil.disk_usage(output).free<max(128*1024**2,total*3):raise ValueError('输出磁盘空间不足')
    emit('strategy',mode=mode,reason=reason,estimatedMiB=estimate//1024**2)
    db=duckdb.connect(':memory:' if mode=='local-light' else str(stage/'analysis.duckdb'))
    db.execute(f"SET memory_limit='{max(64,memory//2)}MB'");db.execute(f'SET threads={threads}');db.execute(f"SET temp_directory={lit(str(stage/'spill'))}");db.execute("SET max_temp_directory_size='2GB'")
    started=time.time();warnings=[]
    try:
        if layout_inspect:
            from inspect_layout import inspect_layout
            count=inspect_layout(db,files,root,p,stage);query='SELECT * FROM data'
            warnings.append('这是工作簿结构概览和有限行列样本，不是全量业务数据导入；首行不自动认定为表头。')
        elif tool in ['file_inventory','pdf_text','link_extract']:
            db.execute('CREATE TABLE data (文件 VARCHAR,位置 VARCHAR,内容 VARCHAR,大小 BIGINT,SHA256 VARCHAR)');count=0
            for f in files:
                rel=f.relative_to(root).as_posix()
                if tool=='file_inventory':records=[(rel,'',f.name,f.stat().st_size,digest(f))]
                elif tool=='pdf_text':
                    from pypdf import PdfReader
                    if f.suffix.lower()!='.pdf':raise ValueError('请选择 PDF 文件')
                    def pages():
                        for i,page in enumerate(PdfReader(f).pages,1):
                            text=page.extract_text() or ''
                            if not text.strip():warnings.append(f'{f.name} 第 {i} 页未提取到文本，需要 OCR 或人工检查')
                            yield(rel,f'第 {i} 页',text,None,None)
                    records=pages()
                else:
                    if f.suffix.lower() not in ['.xlsx','.xlsm']:raise ValueError('请选择 Excel 工作簿')
                    def links():
                        wb=load_workbook(f,read_only=False,data_only=False,keep_links=True)
                        try:
                            for ws in wb:
                                for row in ws:
                                    for c in row:
                                        if c.hyperlink:yield(rel,ws.title+'!'+c.coordinate,c.hyperlink.target or c.hyperlink.location,None,None)
                                        elif c.data_type=='f' and '[' in str(c.value):yield(rel,ws.title+'!'+c.coordinate,str(c.value),None,None)
                        finally:wb.close()
                    records=links()
                for r in records:db.execute('INSERT INTO data VALUES(?,?,?,?,?)',r);count+=1
                emit('progress',phase='处理文件',rows=count)
            query='SELECT * FROM data'
        else:
            columns,count=import_tables(db,files,root,p);emit('progress',phase='执行计算',rows=count)
            query=run_sql(db,tool,p,columns)
        result=export_result(db,query,stage)
        if layout_inspect:
            result['files'].append('结构概览.json');result.update(structureOnly=True,inputUnit='工作表',resultUnit='结构记录')
        if result['excelSkippedForLongText']:warnings.append('包含超过 Excel 单元格限制的长文本，完整内容已保存在 CSV 与 Parquet，未生成 Excel。')
        # Evidence remains separate from user-visible business data.
        if not layout_inspect and tool not in ['file_inventory','pdf_text','link_extract']:
            db.execute(f"COPY data TO {lit(str(stage/'来源数据.parquet'))} (FORMAT PARQUET)");result['files'].append('来源数据.parquet')
        for f,snapshot in zip(files,snapshots):
            if digest(f)!=snapshot['sha256']:raise ValueError('执行中输入文件发生变化，结果未发布')
        manifest={**result,'taskId':request['taskId'],'tool':tool,'status':'succeeded','completeness':'structure_sample' if layout_inspect else 'complete','signature':signature,'ruleVersion':rule_version,'compatibility':'尚未通过原程序全参数对照，不代表完成迁移','mode':mode,'reason':reason,'inputRows':count,'inputs':snapshots,'warnings':warnings[:100],'warningCount':len(warnings),'durationSeconds':round(time.time()-started,3),'outputs':[{'name':name,'sha256':digest(stage/name),'size':(stage/name).stat().st_size} for name in result['files']]}
        manifest['outputDir']=str(final)
        (stage/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
        db.close();db=None
        # Publish only the validated artifact set; analysis files remain in the registered staging directory.
        publish=stage/'publish';publish.mkdir(exist_ok=True)
        for name in result['files']+['manifest.json']:os.replace(stage/name,publish/name)
        os.rename(publish,final)
        return manifest
    finally:
        if db:db.close()

if __name__=='__main__':
    try:
        raw=sys.stdin.buffer.readline(2_000_001)
        if len(raw)>2_000_000:raise ValueError('请求超过大小限制')
        request=json.loads(raw.decode('utf-8'))
        import threading
        def monitor():
            import psutil
            proc=psutil.Process()
            limit=int(request.get('limits',{}).get('memoryMiB',2048))*1024**2
            while True:
                if proc.memory_info().rss>limit:
                    emit('error',message='计算进程超过内存预算，任务已停止，原文件未修改',code='RESOURCE_LIMIT');os._exit(2)
                time.sleep(.25)
        threading.Thread(target=monitor,daemon=True).start()
        emit('result',result=run(request))
    except Exception as e:
        emit('error',message=str(e),code=type(e).__name__);sys.exit(1)
