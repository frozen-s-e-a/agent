from pathlib import Path
import zipfile,itertools,re

def sheet_layout(z,ws):
    # Read a bounded XML tail; ZipExtFile seeks by streaming decompression without
    # building millions of cell objects for formatting-only worksheet ranges.
    info=z.getinfo(ws._worksheet_path);limit=2*1024**2
    with z.open(info) as xml:
        start=xml.read(8192).decode('utf-8','replace')
        xml.seek(max(0,info.file_size-limit));tail=xml.read(limit).decode('utf-8','replace')
    refs=re.findall(r'<(?:\w+:)?mergeCell\s[^>]*?\bref="([A-Z]+\d+(?::[A-Z]+\d+)?)"',tail)
    count=re.search(r'<(?:\w+:)?mergeCells\s[^>]*?\bcount="(\d+)"',tail)
    complete=info.file_size<=limit or count is not None
    total=int(count.group(1)) if count else (len(refs) if complete else None)
    dimension=re.search(r'<(?:\w+:)?dimension\s[^>]*?\bref="([A-Z]+\d+(?::[A-Z]+\d+)?)"',start)
    return {'reportedRange':dimension.group(1) if dimension else None,'mergedRanges':refs[:50],'mergeCount':total,'mergedRangesComplete':total is not None and total==len(refs) and len(refs)<=50,'rangeNote':'保存的范围可能包含空白格式区域，不代表有效数据范围'}

def structure(file,params,all_sheets=False):
    p=Path(file);ext=p.suffix.lower()
    if ext not in ['.xlsx','.xlsm']:
        from worker import tables,val
        for name,rows in tables(p,params):
            sample=list(itertools.islice(rows,6))
            return {'file':p.name,'sampleOnly':True,'sheet':name,'headerRow':int(params.get('headerRow',1)),'columns':[str(v)[:200] if v is not None else '' for v in (sample[0] if sample else [])[:40]],'preview':[[val(v) for v in row[:40]] for row in sample[1:]],'note':'表头及前 5 行、前 40 列；不是全表核查'}
        raise ValueError('没有可读取的表格')
    from openpyxl import load_workbook
    from worker import val
    wb=load_workbook(p,read_only=True,data_only=False,keep_links=False)
    try:
        requested=params.get('sheets','');chosen=[x.strip() for x in requested.split(',') if x.strip()]
        if any(x not in wb.sheetnames for x in chosen):raise ValueError('指定的工作表不存在')
        selected=[ws for ws in wb if not chosen or ws.title in chosen]
        if not all_sheets:selected=selected[:5]
        previews=[]
        for ws in selected:
            with zipfile.ZipFile(p) as z:
                layout=sheet_layout(z,ws)
            ws.reset_dimensions();start=int(params.get('headerRow',1));rows=list(itertools.islice(ws.iter_rows(min_row=start,max_col=40),11 if all_sheets else 6))
            raw_rows=[{'row':start+i,'cells':[{'column':j+1,'value':(val(c.value) or '')[:200],'type':c.data_type,'format':c.number_format} for j,c in enumerate(row)]} for i,row in enumerate(rows)]
            previews.append({'sheet':ws.title,'hidden':ws.sheet_state,**layout,'headerRow':start,'headerConfirmed':'headerRow' in params,'columns':[str(c.value)[:200] if c.value is not None else '' for c in (rows[0] if rows else [])],'preview':[[{'value':(val(c.value) or '')[:200],'type':c.data_type,'format':c.number_format} for c in row] for row in rows[1:]],'rawRows':raw_rows})
        return {'file':p.name,'sheetCount':len(wb.sheetnames),'sheets':[{'name':ws.title,'state':ws.sheet_state} for ws in wb.worksheets],'previews':previews,'sampleOnly':True,'note':('已读取全部指定工作表的结构；每表最多前 11 行' if all_sheets else '本次最多读取 5 个指定工作表、每表前 6 行')+'、前 40 列，每格最多 200 字符。保留空白行和原始行号；首行内容不自动认定为表头。合并范围最多列出 50 项，见完整性标记。保存的范围可能包含格式空白；公式未计算。'}
    finally:wb.close()
