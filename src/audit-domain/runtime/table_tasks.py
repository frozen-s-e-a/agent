from pathlib import Path
import json
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from recovered import table_files
from recovered.kernels import parse_auxiliary_item, filter_leaf_accounts, clean_numeric_prefix
from .context import EXCEL, split, json_value, number


def transform(ctx, name):
    a = ctx.args
    files = ctx.files(a.get('file_path') or a.get('folder_path'), EXCEL, a.get('keyword', ''))
    counts = []
    for source in files:
        p = ctx.copy(source)
        if source.suffix.lower()=='.xlsm' and name in ['fill_column','headers_process','merge_column','dropsummary']:
            p=p.with_suffix('.xlsx')
            book=load_workbook(source,keep_vba=False);book.save(p);book.close()
            ctx.warnings.append(source.name+': 此操作重建数据表，输出为 XLSX；宏和原表样式不包含在成果中。')
        if name == 'add_column':
            if not table_files.merge_same_name_columns(p, a['column_name'], a.get('delimiter', ''), a['max_column_num']):
                raise ValueError(source.name+': 同名列不存在、已合并或超过配置的最大数量')
        elif name == 'asist_split':
            if not table_files.split_auxiliary_columns(p, a['col_name'], a['kv_sep'], a['item_sep']):
                raise ValueError(source.name+': 没有解析到辅助核算项，请检查列名和分隔符')
        elif name == 'fill_column':
            table_files.fill_down_columns(p, a['column_names'])
        elif name == 'headers_process':
            table_files.flatten_headers(p, a['header_rows'])
        elif name == 'merge_column':
            if not table_files.merge_named_columns(p, a['column_names'], a['delimiter'], a['insert_index'], a['insert_name']):
                raise ValueError('待合并列不存在')
        elif name == 'dropsummary':
            df = ctx.table(p)
            ctx.require(df, [a['column_name']])
            table_files.remove_summary_rows(p, a['keywords'], a['column_name'])
        elif name == 'insert_column':
            wb = load_workbook(p, keep_vba=p.suffix.lower()=='.xlsm')
            ws = wb[a['sheet_name']] if a.get('sheet_name') else wb.active
            pos = int(a['insert_position']) if str(a['insert_position']).isdigit() else column_index_from_string(a['insert_position'].upper())
            if not 1<=pos<=ws.max_column+1:
                raise ValueError('插入位置必须在工作表列范围内')
            ws.insert_cols(pos)
            ws.cell(1, pos, a['column_name'])
            for row in range(2, ws.max_row+1):
                ws.cell(row, pos, a['fill_value'])
            wb.save(p)
        ctx.register(p)
        counts.append({'来源文件':str(source.relative_to(ctx.root)), '成果文件':str(p.relative_to(ctx.root)), '行数':len(ctx.table(p))})
    return ctx.result(frame=pd.DataFrame(counts))


def select_column(ctx):
    a = ctx.args
    columns = json_value(a['columns_to_keep'], None) if str(a['columns_to_keep']).strip().startswith('[') else split(a['columns_to_keep'])
    merged = []
    for p in ctx.files(a['folder_path'], EXCEL):
        sheets = pd.read_excel(p, sheet_name=None)
        selected = table_files.filter_sheets_by_keywords(list(sheets), split(a.get('sheet_keywords', '')))
        results = {}
        for name in selected:
            df = sheets[name]
            ctx.require(df, columns)
            df = df[columns].copy() if columns else df.copy()
            if a['should_merge'] in ['是', True]:
                df['_来源文件'], df['_来源工作表'] = str(p.relative_to(ctx.root)), name
                merged.append(df)
            else:
                results[name] = df
        if results:
            ctx.report(results, p.stem+'_列提取.xlsx')
    if merged:
        frame = pd.concat(merged, ignore_index=True)
        ctx.report(frame, '合并列提取结果.xlsx')
        return ctx.result(frame=frame)
    if not ctx.outputs:
        raise ValueError('没有匹配的工作表')
    return ctx.result()


def pickup_value(ctx):
    a = ctx.args
    mapping = json_value(a['field_mappings'])
    if not isinstance(mapping, dict) or not mapping:
        raise ValueError('字段配置应为 JSON 对象，例如 {"公司名称":"A1"}')
    rows = []
    for p in ctx.files(a['folder_path'], EXCEL):
        wb = load_workbook(p, read_only=True, data_only=True)
        for ws in wb:
            row = {'来源文件':str(p.relative_to(ctx.root)), '工作表':ws.title}
            for name, cell in mapping.items():
                if not re.fullmatch(r'[A-Za-z]+[1-9]\d*', str(cell)):
                    raise ValueError('单元格地址无效: '+str(cell))
                row[name] = ws[cell].value
            rows.append(row)
        wb.close()
    frame = pd.DataFrame(rows)
    ctx.report(frame, '固定单元格提取.xlsx')
    return ctx.result(frame=frame)


def preview(ctx, name):
    a = ctx.args
    p = ctx.path(a['file_path'], True)
    if p.is_dir():
        frames = [(f, ctx.table(f, a.get('sheet_name'))) for f in ctx.files(p)]
        columns = list(dict.fromkeys(c for _, frame in frames for c in frame.columns))
        result = {'type':'folder', 'row_count':sum(len(frame) for _,frame in frames), 'file_count':len(frames),
                  'columns':columns, 'files':[{'file':str(f.relative_to(ctx.root)), 'rows':len(frame), 'columns':list(frame.columns)} for f,frame in frames]}
        ctx.json(result, '数据结构.json')
        return ctx.result(result, frames[0][1])
    header = int(a.get('header_row') or 1)
    if name=='tool_preview' and not a.get('header_row'):
        raw = ctx.table(p, a.get('sheet_name'), header=1, nrows=20)
        if not any(re.search('科目|编码|代码', str(c)) for c in raw.columns):
            wb = load_workbook(p, read_only=True)
            ws = wb[a['sheet_name']] if a.get('sheet_name') else wb.active
            candidates = [(i, sum(bool(re.search('科目|编码|代码', str(v))) for v in row)) for i,row in enumerate(ws.values, 1) if i<=20]
            header = max(candidates, key=lambda x:x[1])[0]
            wb.close()
    df = ctx.table(p, a.get('sheet_name'), header)
    result = {'header_row':header, 'row_count':len(df), 'columns':list(df.columns),
              'sample_rows':df.head(min(100, int(a.get('preview_rows',5)))).to_dict('records')}
    if p.suffix.lower() in EXCEL:
        wb=load_workbook(p, read_only=True); result['sheets']=wb.sheetnames; wb.close()
    ctx.json(result, '数据结构.json')
    return ctx.result(result, df)


def subject_execute(ctx):
    a = ctx.args
    df = ctx.table(a['file_path'], header=a['header_row'], dtype={a['code_column']:str})
    code, name = a['code_column'], a['code_name_column']
    ctx.require(df, [code,name,a.get('asist_column'),a.get('grouping_column')])
    df = df[df[code].notna()].copy()
    df[code] = df[code].astype(str).str.strip()
    names = dict(zip(df[code], df[name].fillna('').astype(str)))
    leaf = filter_leaf_accounts(df.copy(), code)
    hierarchies = []
    for _, row in leaf.iterrows():
        if a.get('code_name_split'):
            parts = clean_numeric_prefix(str(row[name]).split(a['code_name_split']))
        else:
            parts = [names[c] for c in sorted(names, key=len) if str(row[code]).startswith(c)]
        hierarchies.append(parts)
    for i in range(max((len(x) for x in hierarchies), default=0)):
        leaf[str(i+1)+'级科目'] = [parts[i] if len(parts)>i else '' for parts in hierarchies]
    if a.get('asist_column'):
        parsed = [parse_auxiliary_item(v, a.get('asist_name_split') or ':', a.get('asist_item_wrapper') or '[]') for v in leaf[a['asist_column']]]
        for k in dict.fromkeys(k for row in parsed for k in row):
            leaf[k] = [row.get(k,'') for row in parsed]
    if a.get('beginning_balance'):
        # Explicit column subtraction only; expressions never run as Python.
        terms = a['beginning_balance'].split('-')
        ctx.require(leaf, terms)
        leaf['期初余额_计算'] = number(leaf[terms[0]])
        for t in terms[1:]:
            leaf['期初余额_计算'] -= number(leaf[t])
    group = a.get('grouping_column')
    groups = leaf.groupby(group, dropna=False) if group else [(a.get('company_name') or '科目余额表',leaf)]
    for company, part in groups:
        ctx.report({'末级科目余额表':part, '原科目清单':df}, str(company)+'_科目处理.xlsx')
    return ctx.result(frame=leaf)


TASKS = {name:(lambda ctx,n=name:transform(ctx,n)) for name in ['add_column','asist_split','fill_column','headers_process','merge_column','dropsummary','insert_column']}
TASKS.update(select_column=select_column, pickup_value=pickup_value, tool_execute=subject_execute)
TASKS.update({name:(lambda ctx,n=name:preview(ctx,n)) for name in ['tool_preview','monthly_preview','detailed_table_inspect']})
