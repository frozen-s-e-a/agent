import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from recovered import jet, accounting
from recovered.kernels import find_parent_codes, pvm_decompose
from recovered.date_rules import clean_and_parse_dates, determine_cutoff_date, deduplicate_and_merge_subjects
from .context import split, json_value, number, EXCEL
from .transport import model
from .table_tasks import preview


def monthly_structure(ctx):
    a = ctx.args
    mapping = json_value(a['column_mapping'], {})
    df = ctx.merged(a['file_path'], a.get('sheet_name'), dtype={mapping['subject_code']:str})
    ctx.require(df, mapping.values())
    code, name = mapping['subject_code'], mapping['subject_name']
    codes = df[code].dropna().astype(str).str.strip().unique().tolist()
    lengths = sorted(set(map(len,codes)))
    parents = set(find_parent_codes(codes.copy()))
    roots = sorted([c for c in codes if not any(c.startswith(p) and p!=c for p in codes)], key=str)
    result = {'code_structure': {'init_length':lengths[0] if lengths else 4,
                                'step':min((b-a for a,b in zip(lengths,lengths[1:])), default=2), 'full_length':0, 'start':1},
              'levels':lengths, 'parent_codes':sorted(parents),
              'subjects':[{'code':c,'name':str(df.loc[df[code].astype(str)==c,name].iloc[0])} for c in roots[:int(a.get('top_n',30))]]}
    ctx.json(result, '科目结构.json')
    return ctx.result(result)


def monthly_source(ctx, path, sheet, mapping, subjects):
    df = ctx.merged(path, sheet, dtype={mapping['subject_code']:str})
    ctx.require(df, [mapping[k] for k in ['date','subject_code','subject_name']]+[mapping[k] for k in ['debit','credit','company'] if mapping.get(k)])
    dates = pd.to_datetime(df[mapping['date']], errors='coerce', format='mixed')
    if dates.isna().any():
        bad = df.loc[dates.isna(), ['_来源文件','_来源行']].to_dict('records')
        ctx.json(bad, '日期异常.json')
        raise ValueError('有 '+str(len(bad))+' 行日期无法识别，不能完成月间分析')
    codes = df[mapping['subject_code']].fillna('').astype(str).str.strip()
    data = []
    company = df[mapping['company']].fillna('未填写主体') if mapping.get('company') else pd.Series('汇总主体',index=df.index)
    debit = number(df[mapping['debit']]) if mapping.get('debit') else pd.Series(0.,index=df.index)
    credit = number(df[mapping['credit']]) if mapping.get('credit') else pd.Series(0.,index=df.index)
    for subject in subjects:
        code = str(subject['code'])
        mask = codes.str.startswith(code)
        selected = pd.DataFrame({'公司名称':company[mask], '月份':dates[mask].dt.month,
                                 '金额':(debit-credit if subject['direction']=='借' else credit-debit)[mask]})
        selected = selected.groupby(['公司名称','月份'],as_index=False,dropna=False)['金额'].sum()
        selected['科目编码'] = code
        selected['科目名称'] = str(df.loc[codes==code,mapping['subject_name']].iloc[0]) if (codes==code).any() else str(subject.get('name',code))
        data.append(selected)
    if not data:
        raise ValueError('没有待分析科目')
    return pd.concat(data, ignore_index=True), df


def monthly_execute(ctx):
    a = ctx.args
    mapping, subjects = json_value(a['column_mapping'], {}), json_value(a['subjects'], [])
    if not subjects or any(not x.get('code') or x.get('direction') not in ['借','贷'] for x in subjects):
        raise ValueError('请选择科目，并明确每个科目的借贷方向')
    this, raw = monthly_source(ctx,a['this_year_file'],a.get('this_year_sheet'),mapping,subjects)
    keys = ['公司名称','科目编码','科目名称','月份']
    this = this.rename(columns={'金额':'本期金额'})
    if a.get('last_year_file'):
        last, _ = monthly_source(ctx,a['last_year_file'],a.get('last_year_sheet'),mapping,subjects)
        result = this.merge(last.rename(columns={'金额':'上期金额'}), on=keys, how='outer').fillna({'本期金额':0,'上期金额':0})
    else:
        result = this.copy(); result['上期金额'] = np.nan
    result['变动额'] = result['本期金额']-result['上期金额']
    result['变动率'] = result['变动额']/result['上期金额'].abs().replace(0,np.nan)
    result['零基数变动'] = np.where((result['上期金额']==0)&(result['本期金额']!=0),'是','否')
    totals = result.groupby(keys[:-1],as_index=False,dropna=False)[['本期金额','上期金额','变动额']].sum(min_count=1)
    totals['变动率'] = totals['变动额']/totals['上期金额'].abs().replace(0,np.nan)
    result['本期月度占比'] = result['本期金额'].abs()/result.groupby(keys[:-1],dropna=False)['本期金额'].transform(lambda x:x.abs().sum()).replace(0,np.nan)
    result['需关注'] = np.where((result['变动额'].abs()>=float(a.get('threshold_amt',100000))) & ((result['变动率'].abs()>=float(a.get('threshold_rate',.3))) | (result['零基数变动']=='是')), '是','否')
    ctx.report({'科目汇总':totals, '月度明细':result, '本期数据来源':raw}, '月间分析结果表.xlsx')
    return ctx.result(frame=result)


def monthly_ai(ctx):
    a = ctx.args
    wb = load_workbook(ctx.path(a['result_file'], True), read_only=True)
    sheet = '月度明细' if '月度明细' in wb.sheetnames else wb.sheetnames[0]; wb.close()
    result = ctx.table(a['result_file'],sheet,dtype={'科目编码':str})
    mapping = json_value(a['column_mapping'],{})
    raw = ctx.merged(a['this_year_file'],a.get('this_year_sheet'),dtype={mapping['subject_code']:str})
    ctx.require(raw,[mapping['subject_code']])
    needed=['科目编码','本期金额','上期金额','变动额','变动率']
    ctx.require(result,needed)
    selected=result[(number(result['变动额']).abs()>=float(a.get('threshold_amt',100000))) & ((number(result['变动率']).abs()>=float(a.get('threshold_rate',.3))) | (number(result['上期金额'])==0))]
    analyses=[]
    for code,group in selected.groupby('科目编码',dropna=False):
        sample=raw[raw[mapping['subject_code']].astype(str).str.startswith(str(code))].head(int(a.get('max_sample_rows',50)))
        response=model(ctx,'根据以下月间分析和有限样本，返回 JSON {"变动原因线索":"...","待补证据":"...","审计应对建议":"..."}。只描述数据支持的线索，不能推断未提供的业务事实。\n月间分析:'+group.to_json(orient='records',force_ascii=False)+'\n样本:'+sample.to_json(orient='records',force_ascii=False),schema=True)
        if not isinstance(response,dict):raise ValueError('模型分析必须返回对象')
        analyses.append({'科目编码':code,**response,'样本行数':len(sample),'结论状态':'待人工复核'})
    ctx.report({'月间分析':result,'AI变动线索':pd.DataFrame(analyses)},'月间分析_AI复核.xlsx')
    return ctx.result(frame=pd.DataFrame(analyses))


def gross_margin(ctx):
    a=ctx.args
    df=ctx.table(a['this_data_path'],a.get('this_data_sheet'))
    if a.get('mode')=='preview':
        return ctx.result(frame=df)
    aliases={'营业收入':['营业收入','销售金额','收入','本期收入'],'营业成本':['营业成本','成本金额','成本','本期成本'],'数量':['数量','销量','销售数量']}
    mapping=json_value(a.get('field_mapping'),{})
    for key,names in aliases.items():
        if not mapping.get(key):
            found=[c for c in df if c in names]
            if len(found)==1:mapping[key]=found[0]
    missing=[k for k in aliases if not mapping.get(k)]
    if missing:raise ValueError('请明确字段映射: '+'、'.join(missing))
    ctx.require(df,mapping.values())
    dimensions=json_value(a.get('dimensions'),None)
    if dimensions is None:
        dimensions=[{'name':c,'column':c,'threshold':.05} for c in df if c not in mapping.values() and any(k in str(c) for k in ['产品','类别','区域','客户'])]
    if not a.get('confirmed'):
        return ctx.result({'status':'needs_confirmation','needs_confirmation':True,'field_mapping':mapping,'dimensions':dimensions,'message':'请核对字段映射和维度后确认执行'},df)
    df=df.rename(columns={v:k for k,v in mapping.items()})
    for k in aliases:df[k]=number(df[k])
    last=None
    if a.get('last_data_path'):
        last=ctx.table(a['last_data_path'],a.get('last_data_sheet'))
        ctx.require(last,mapping.values());last=last.rename(columns={v:k for k,v in mapping.items()})
        for k in aliases:last[k]=number(last[k])
    sheets={}
    dimensions=dimensions or [{'name':'总体','column':None,'threshold':.05}]
    for dim in dimensions:
        column=dim.get('column')
        if column:
            ctx.require(df,[column]);cur=df.groupby(column,dropna=False)[list(aliases)].sum().reset_index()
        else:cur=pd.DataFrame([{**{k:df[k].sum() for k in aliases},'总体':'全部'}]);column='总体'
        cur=cur.rename(columns={k:'本期'+k for k in aliases})
        if last is not None:
            if column!='总体':ctx.require(last,[column])
            prev=last.groupby(column,dropna=False)[list(aliases)].sum().reset_index() if column in last else pd.DataFrame([{**{k:last[k].sum() for k in aliases},'总体':'全部'}])
            cur=cur.merge(prev.rename(columns={k:'上期'+k for k in aliases}),on=column,how='outer').fillna(0)
        for prefix in ['本期','上期'] if last is not None else ['本期']:
            cur[prefix+'毛利']=cur[prefix+'营业收入']-cur[prefix+'营业成本']
            cur[prefix+'毛利率']=cur[prefix+'毛利']/cur[prefix+'营业收入'].replace(0,np.nan)
            cur[prefix+'单价']=cur[prefix+'营业收入']/cur[prefix+'数量'].replace(0,np.nan)
            cur[prefix+'单位成本']=cur[prefix+'营业成本']/cur[prefix+'数量'].replace(0,np.nan)
        if last is not None:
            cur['毛利变动']=cur['本期毛利']-cur['上期毛利']
            cur['毛利率变动']=cur['本期毛利率']-cur['上期毛利率']
            components=[]
            for _,r in cur.iterrows():
                components.append(pvm_decompose(r))
            # All decomposition values originate from the recovered price/volume/cost kernel.
            if components and isinstance(components[0],dict):cur=pd.concat([cur,pd.DataFrame(components)],axis=1)
        sheets[dim.get('name',column)]=cur
    ctx.report(sheets,'毛利率分析.xlsx')
    return ctx.result({'field_mapping':mapping,'dimensions':dimensions},next(iter(sheets.values())))


def jet_inspect(ctx):
    a=ctx.args
    df=ctx.merged(a['file_path'])
    data={'row_count':len(df),'columns':list(df.columns),'files':[str(p.relative_to(ctx.root)) for p in ctx.files(a['file_path'])]}
    if a.get('include_templates',True):data.update(jet.builtin_definitions())
    ctx.json(data,'JET输入检查.json')
    return ctx.result(data,df)


def definitions(value, builtins, name):
    if value is None or value=='' or value=='builtin':return builtins
    values=json_value(value,[])
    if not isinstance(values,list):raise ValueError(name+'必须是 JSON 列表或 builtin')
    result=[]
    for item in values:
        if item=='builtin':result.extend(builtins)
        elif isinstance(item,str):
            found=next((v for v in builtins if v['name']==item),None)
            if found is None:raise ValueError('未知'+name+': '+item)
            result.append(found)
        elif isinstance(item,dict):result.append(item)
        else:raise ValueError(name+'配置格式错误')
    return result


def jet_execute(ctx):
    a=ctx.args
    template=jet.builtin_definitions()
    features=definitions(a.get('features'),template['features'],'特征')
    rules=definitions(a.get('rules'),template['rules'],'规则')
    if not rules:raise ValueError('至少选择一条 JET 规则')
    mapping=json_value(a.get('column_mapping'),{})
    df=ctx.merged(a['file_path'],dtype=str).rename(columns={v:k for k,v in mapping.items()})
    df['_输入记录编号']=range(1,len(df)+1)
    computed={}; failed={}
    for feature in features:
        name=feature['name']
        if name in df:continue
        try:
            df[name]=jet.apply_feature(df,feature['tool'],feature.get('source_cols'),feature.get('param'))
            computed[name]=feature
        except (KeyError,ValueError,TypeError) as e:failed[name]=str(e)
    outputs={};summary=[]
    for rule in rules:
        if not rule.get('conditions'):raise ValueError('规则缺少条件: '+str(rule.get('name')))
        missing=[col for col in rule['conditions'] if col not in df]
        if missing:
            summary.append({'规则':rule['name'],'状态':'未执行','原因':'缺少字段: '+','.join(missing),'命中行数':None})
            continue
        mask=pd.Series(True,index=df.index)
        for col,condition in rule['conditions'].items():
            check=jet.check_condition(df[col],condition)
            if check is not None:mask &= check
        hit=df.loc[mask].copy()
        summary.append({'规则':rule['name'],'状态':'已执行','原因':'','命中行数':len(hit)})
        if len(hit) or json_value(a.get('settings'),{}).get('output_empty',False):outputs[rule['name']]=hit
    skipped=[r for r in summary if r['状态']=='未执行']
    if len(skipped)==len(rules):raise ValueError('所有规则都缺少必需字段，请检查列映射和特征配置')
    if skipped:ctx.warnings.append('有 '+str(len(skipped))+' 条规则未执行；请查看规则汇总中的缺失字段')
    report={'规则汇总':pd.DataFrame(summary),**outputs}
    ctx.report(report,'JET测试结果.xlsx')
    ctx.json({'features':features,'rules':rules,'mapping':mapping,'unavailable_features':failed},'JET执行规则.json')
    return ctx.result({'rules_executed':len(rules)-len(skipped),'rules_skipped':len(skipped),'rule_summary':summary},pd.DataFrame(summary))


def cutoff(ctx):
    wb=ctx.config();c=ctx.kv(wb)
    df=ctx.merged(c['文件或文件夹路径'],c.get('表名（非必填）'))
    date,subject,voucher=c['日期'],c['1级科目'],c['凭证号'];company=c.get('公司（非必填）')
    ctx.require(df,[date,subject,voucher,c['借方金额'],c['贷方金额'],company])
    df[date]=clean_and_parse_dates(df[date])
    start,end=pd.Timestamp(c['开始日期']),pd.Timestamp(c['结束日期'])
    if start>end:raise ValueError('开始日期不能晚于结束日期')
    cutoff_date=determine_cutoff_date(start,end)
    if cutoff_date is None:raise ValueError('截止测试期间必须跨越 12 月 31 日')
    if not company:company='_公司';df[company]='汇总主体'
    df=df[df[date].between(start,end)].copy()
    df['抽样金额']=pd.concat([number(df[c['借方金额']]).abs(),number(df[c['贷方金额']]).abs()],axis=1).max(axis=1)
    df=df.sort_values('抽样金额',ascending=False,kind='stable')
    df=deduplicate_and_merge_subjects(df,company,date,voucher,subject)
    df['截止分组']=np.where(df[date]<=pd.Timestamp(cutoff_date),'年前','年后')
    if c.get('模式','金额')=='金额':
        result=df[df['抽样金额']>=float(c.get('金额') or 0)]
    elif c['模式']=='笔数':
        result=df.groupby([company,'截止分组'],dropna=False,group_keys=False).head(int(c.get('笔数（非必填）') or 10))
    else:raise ValueError('抽样模式必须是金额或笔数')
    ctx.report({'截止样本':result,'候选总体':df},'截止测试样本.xlsx')
    return ctx.result({'population_count':len(df),'sample_count':len(result),'cutoff_date':str(cutoff_date)},result)


def ledger(ctx):
    rows=[];details=[]
    for p in ctx.files(ctx.args['directory'],EXCEL):
        wb=load_workbook(p,read_only=True);sheets=wb.sheetnames;wb.close()
        ledger_sheet=next((s for s in sheets if '序时' in s),None)
        balance_sheet=next((s for s in sheets if '余额' in s and '辅助' not in s),None)
        if not ledger_sheet or not balance_sheet:raise ValueError(p.name+': 需要序时账和科目余额表两张工作表')
        gl=ctx.table(p,ledger_sheet,dtype={'科目代码':str})
        bal=ctx.table(p,balance_sheet,dtype={'科目代码':str})
        ctx.require(gl,['科目代码','借方金额','贷方金额'])
        ctx.require(bal,['科目代码','科目名称','期初余额','期末余额','借方金额','贷方金额','方\n向','方\n向.1'])
        for frame in [gl,bal]:
            for col in ['期初余额','期末余额','借方金额','贷方金额']:
                if col in frame:frame[col]=number(frame[col])
        balanced=accounting.judge_balance(gl);calc,subject=accounting.judge_calc(bal)
        for _,row in bal.iterrows():
            code=str(row['科目代码']);debit,credit=accounting.sum_amounts_by_code(gl,code)
            hierarchy=accounting.judge_father_equal_children(bal,code)
            details.append({'文件':p.name,'科目代码':code,'科目名称':row['科目名称'],'账表借方差异':debit-row['借方金额'],'账表贷方差异':credit-row['贷方金额'],'上下级一致':hierarchy})
        rows.append({'文件':p.name,'借贷平衡':balanced,'余额勾稽':calc,'首个余额差异科目':subject})
    result=pd.DataFrame(rows)
    ctx.report({'账套检查':result,'科目差异':pd.DataFrame(details)},'小账套检查结果.xlsx')
    return ctx.result(frame=result)


TASKS={'monthly_detect_structure':monthly_structure,'monthly_execute':monthly_execute,'monthly_ai_analysis':monthly_ai,
       'gross_margin_analysis':gross_margin,'jet_test_inspect':jet_inspect,'jet_test_execute':jet_execute,
       'cutoff_sampling':cutoff,'ledger_check':ledger}
