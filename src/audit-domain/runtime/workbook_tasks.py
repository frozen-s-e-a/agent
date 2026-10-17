from copy import copy
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.utils import range_boundaries, column_index_from_string

from .context import EXCEL, split, number


def workbook(path):
    return load_workbook(path, keep_vba=path.suffix.lower()=='.xlsm', keep_links=True)


def copy_sheet(source, target):
    for row in source:
        for cell in row:
            if cell.__class__.__name__=='MergedCell':continue
            dest=target.cell(cell.row,cell.column,cell.value)
            if cell.has_style:dest._style=copy(cell._style)
            if cell.comment:dest.comment=copy(cell.comment)
            if cell.hyperlink:dest._hyperlink=copy(cell.hyperlink)
            dest.number_format=cell.number_format
    for key,value in source.row_dimensions.items():target.row_dimensions[key]=copy(value)
    for key,value in source.column_dimensions.items():target.column_dimensions[key]=copy(value)
    for region in source.merged_cells.ranges:target.merge_cells(str(region))
    target.sheet_format=copy(source.sheet_format)
    target.sheet_properties=copy(source.sheet_properties)
    target.page_setup=copy(source.page_setup)
    target.page_margins=copy(source.page_margins)
    target.freeze_panes=source.freeze_panes


def link_extract(ctx):
    wb=ctx.config();paths=[]
    for r in ctx.rows(wb):paths.extend(ctx.files(r['文件或文件夹路径'],EXCEL))
    rows=[]
    for p in dict.fromkeys(paths):
        w=workbook(p)
        for ws in w:
            for row in ws:
                for cell in row:
                    if cell.hyperlink:
                        rows.append({'文件':str(p.relative_to(ctx.root)),'工作表':ws.title,'单元格':cell.coordinate,'类型':'超链接','内容':cell.hyperlink.target or cell.hyperlink.location})
                    if cell.data_type=='f' and ('[' in cell.value or '!' in cell.value):
                        rows.append({'文件':str(p.relative_to(ctx.root)),'工作表':ws.title,'单元格':cell.coordinate,'类型':'公式链接','内容':cell.value})
        w.close()
    result=pd.DataFrame(rows,columns=['文件','工作表','单元格','类型','内容']);ctx.report(result,'工作簿链接清单.xlsx')
    return ctx.result(frame=result)


def replace_links(ctx,name):
    cfg=ctx.config(); jobs=[]
    if name=='link_replace_text':
        ws=cfg['配置表'];folder=ws['B1'].value;keyword=ws['D1'].value or ''
        pairs=[(str(r[0]),str(r[1] or '')) for r in list(ws.values)[2:] if r[0] is not None]
        jobs=[(p,pairs) for p in ctx.files(folder,EXCEL,keyword)]
    elif name=='link_replace_workbook':
        for row in ctx.rows(cfg):
            pairs=[(old,str(row['新链接文本'])) for old in split(row['老链接文本'])]
            jobs.append((ctx.path(row['工作簿路径'],True),pairs))
    counts=[]
    for p,pairs in jobs:
        dest=ctx.copy(p);w=workbook(dest);count=0
        for ws in w:
            for row in ws:
                for cell in row:
                    if cell.data_type=='f':
                        old=cell.value;new=old
                        for before,after in pairs:new=new.replace(before,after)
                        if old!=new:cell.value=new;count+=1
                    if cell.hyperlink and cell.hyperlink.target:
                        old=cell.hyperlink.target;new=old
                        for before,after in pairs:new=new.replace(before,after)
                        if old!=new:cell.hyperlink.target=new;count+=1
        # Excel external-link relationship targets also carry workbook names.
        for link in w._external_links:
            if link.file_link and link.file_link.Target:
                for before,after in pairs:link.file_link.Target=link.file_link.Target.replace(before,after)
        w.save(dest);w.close();ctx.register(dest);counts.append({'文件':p.name,'替换单元格数':count})
    if not jobs:raise ValueError('配置表没有可执行的替换规则')
    return ctx.result(frame=pd.DataFrame(counts))


def replace_range(ctx):
    cfg=ctx.config();counts=[]
    for row in ctx.rows(cfg):
        p=ctx.path(row['现工作簿路径'],True);origin=ctx.path(row['原工作簿路径'],True)
        w=workbook(ctx.copy(p));source=workbook(origin)
        ws=w[row['现工作表名']];src=source[row['原工作表名']]
        x1,y1,x2,y2=range_boundaries(row['现单元格区域']);a1,b1,a2,b2=range_boundaries(row['原单元格区域'])
        if (x2-x1,y2-y1)!=(a2-a1,b2-b1):raise ValueError('来源区域与目标区域大小不一致')
        for dy in range(y2-y1+1):
            for dx in range(x2-x1+1):
                cell=ws.cell(y1+dy,x1+dx);ref=src.cell(b1+dy,a1+dx)
                escaped=str(origin).replace("'","''")
                # A direct workbook link is reviewable and Excel recalculates it.
                cell.value="='"+str(origin.parent).replace("'","''")+"/["+origin.name.replace("'","''")+"]"+src.title.replace("'","''")+"'!"+ref.coordinate
        dest=ctx._copies[p];w.save(dest);ctx.register(dest);counts.append({'目标文件':p.name,'工作表':ws.title,'范围':row['现单元格区域']})
        source.close();w.close()
    return ctx.result(frame=pd.DataFrame(counts))


def ref_replace(ctx):
    a=ctx.args;records=[]
    for p in ctx.files(a['target_path'],EXCEL):
        dest=ctx.copy(p);w=workbook(dest);count=0
        for ws in w:
            if a.get('sheet_keywords') and not any(k in ws.title for k in split(a['sheet_keywords'])):continue
            example=ws[a['example_cell']]
            if example.data_type!='f':raise ValueError('示例单元格必须含有公式')
            axis=str(a['target_axis']).upper();ref=str(a['ref_axis']).upper()
            if axis.isdigit() and ref.isdigit():
                targets=[ws.cell(int(axis),c) for c in range(1,ws.max_column+1)]
                origins=[ws.cell(int(ref),c) for c in range(1,ws.max_column+1)]
            elif axis.isalpha() and ref.isalpha():
                targets=[ws.cell(r,column_index_from_string(axis)) for r in range(1,ws.max_row+1)]
                origins=[ws.cell(r,column_index_from_string(ref)) for r in range(1,ws.max_row+1)]
            else:raise ValueError('参考行与目标行应同时是数字，参考列与目标列应同时是列字母')
            for cell,origin in zip(targets,origins):
                if origin.value is None or cell.row==1:continue
                cell.value=Translator(example.value,origin=example.coordinate).translate_formula(cell.coordinate);count+=1
        w.save(dest);w.close();ctx.register(dest);records.append({'文件':p.name,'公式写入数':count})
    return ctx.result(frame=pd.DataFrame(records))


def collect(ctx,folder_mode=False):
    cfg=ctx.config();tablemap=ctx.rows(cfg,'表匹配关系');records=[]
    for paths in ctx.rows(cfg,'路径匹配'):
        target=ctx.path(paths['目标工作簿'],True);dest=ctx.copy(target);w=workbook(dest)
        sources=ctx.files(paths['来源文件夹'],EXCEL) if folder_mode else [ctx.path(paths['来源工作簿'],True)]
        for mapping in tablemap:
            selected=[p for p in sources if not folder_mode or str(mapping.get('来源工作簿文件关键词') or '') in p.name]
            if len(selected)!=1:raise ValueError('来源工作簿未唯一匹配: '+str(mapping))
            src=workbook(selected[0]);srcname=str(mapping['来源表名']);name=str(mapping['目标表名'])
            if srcname not in src:raise ValueError('来源工作表不存在: '+srcname)
            if name in w:del w[name]
            copy_sheet(src[srcname],w.create_sheet(name));src.close()
            records.append({'目标文件':target.name,'目标表':name,'来源文件':selected[0].name,'来源表':srcname})
        w.save(dest);w.close();ctx.register(dest)
    return ctx.result(frame=pd.DataFrame(records))


def insert_comments(ctx):
    cfg=ctx.config();records=[]
    for config in ctx.rows(cfg):
        for p in ctx.files(config['文件夹或文件路径'],EXCEL):
            dest=ctx.copy(p);w=workbook(dest);ws=w[str(config['工作表'])]
            start=int(config['开始行号']);end=min(int(config['结束行号']),ws.max_row)
            if config.get('是否清除所有批注')=='是':
                for row in ws:
                    for cell in row:cell.comment=None
            regions=[];active=None
            for r in range(start,end+1):
                cells=list(ws[r]);numeric=[c for c in cells if isinstance(c.value,(int,float)) or c.data_type=='f']
                label=any(isinstance(c.value,str) and c.data_type!='f' and c.value.strip() for c in cells)
                if numeric and label:
                    if active is None:active=r
                elif active is not None:
                    regions.append((max(start,active-1),r-1));active=None
            if active is not None:regions.append((max(start,active-1),end))
            for first,last in regions:
                ws.cell(first,1).comment=Comment('表首','AI审计助手')
                ws.cell(last,1).comment=Comment('表尾','AI审计助手')
            w.save(dest);w.close();ctx.register(dest);records.append({'文件':p.name,'表首表尾数量':len(regions),'状态':'请复核自动识别范围'})
    ctx.warnings.append('附注表范围根据连续数值区域识别，需复核表首和表尾批注后再提取。')
    return ctx.result(frame=pd.DataFrame(records))


def disclosure(ctx,scope=False):
    cfg=ctx.config();c=ctx.kv(cfg);sheets={};scopes=[];total=0
    source=c['文件路径'] if scope else c['文件夹或文件路径']
    for p in ctx.files(source,EXCEL):
        w=workbook(p)
        for ws in w:
            if scope and ws.title!=str(c['表名']):continue
            if not scope and c.get('工作表关键词') and not any(k in ws.title for k in split(c['工作表关键词'])):continue
            first=None
            head=str(c.get('表头批注关键词') if scope else c.get('表首关键词') or '表首')
            tail=str(c.get('表尾批注关键词') if scope else c.get('表尾关键词') or '表尾')
            for row in ws:
                for cell in row:
                    if not cell.comment:continue
                    text=cell.comment.text
                    if head in text:first=cell.row
                    if tail in text and first is not None:
                        last=cell.row
                        rows=[[v.value for v in r] for r in ws.iter_rows(min_row=first,max_row=last)]
                        if rows:
                            title=next((str(v) for r in rows[:2] for v in r if isinstance(v,str) and v.strip()),'表格')
                            total+=1;sheets[str(total)+'_'+title[:23]]=pd.DataFrame(rows)
                            for r in rows[1:]:
                                left=next((str(v).strip() for v in r if isinstance(v,str) and v.strip()),'')
                                if left:scopes.append({'来源文件':p.name,'工作表':ws.title,'表格':title,'披露口径':left,'开始行':first,'结束行':last})
                        first=None
        w.close()
    if not total:raise ValueError('没有找到成对的表首和表尾批注，请检查配置或先插入批注')
    if scope:ctx.report(pd.DataFrame(scopes),'披露口径提取结果.xlsx')
    else:ctx.report(sheets,'披露表提取结果.xlsx')
    return ctx.result({'table_count':total},pd.DataFrame(scopes))


TASKS={'link_extractor':link_extract,'link_replace_text':lambda c:replace_links(c,'link_replace_text'),
       'link_replace_workbook':lambda c:replace_links(c,'link_replace_workbook'),'link_replace_range':replace_range,
       'link_ref_replace':ref_replace,'multi_collect_file':lambda c:collect(c),'multi_collect_folder':lambda c:collect(c,True),
       'insert_footnote_comments':insert_comments,'extract_disclosure_scope':lambda c:disclosure(c,True),
       'extract_disclosure_tables':lambda c:disclosure(c)}
