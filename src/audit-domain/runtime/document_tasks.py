from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import json
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

from recovered.matching import process_extract_result, match_and_compare
from .context import EXCEL, split, safe_name
from .transport import model
from .workbook_tasks import workbook

IMAGES={'.png','.jpg','.jpeg','.webp','.bmp','.tif','.tiff'}
DOCUMENTS={'.pdf','.docx','.txt','.md','.csv','.tsv',*EXCEL,*IMAGES}


def document(ctx,p):
    p=ctx.path(p,True)
    if p.suffix.lower() in IMAGES:
        return '',[p]
    if p.suffix.lower()=='.pdf':
        import pymupdf
        doc=pymupdf.open(p);texts=[];images=[]
        if len(doc)>300:raise ValueError('单个 PDF 超过 300 页，请拆分后处理')
        folder=ctx.output/'文档页面';folder.mkdir(exist_ok=True)
        for i,page in enumerate(doc):
            text=page.get_text().strip()
            if len(text)>30:texts.append('第'+str(i+1)+'页\n'+text)
            else:
                target=folder/(safe_name(p.stem)+'_'+str(i+1)+'.png')
                page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(target);images.append(target)
        doc.close()
        return '\n'.join(texts),images
    if p.suffix.lower()=='.docx':
        with zipfile.ZipFile(p) as archive:
            xml=ET.fromstring(archive.read('word/document.xml'))
        return '\n'.join(''.join(n.itertext()) for n in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')),[]
    if p.suffix.lower() in {'.txt','.md'}:
        try:return p.read_text(encoding='utf-8-sig'),[]
        except UnicodeDecodeError:return p.read_text(encoding='gb18030'),[]
    if p.suffix.lower() in {'.csv','.tsv'}:return ctx.table(p).to_csv(index=False),[]
    if p.suffix.lower() in EXCEL:
        sheets=pd.read_excel(p,sheet_name=None,dtype=str)
        return '\n'.join('工作表:'+name+'\n'+df.fillna('').to_csv(index=False) for name,df in sheets.items()),[]
    raise ValueError('不支持的文档格式: '+p.suffix)


def extract_one(ctx,p,prompt):
    text,images=document(ctx,p)
    if not text and not images:raise ValueError(p.name+': 文件中没有可读取内容')
    chunks=[text[i:i+60000] for i in range(0,len(text),60000)] or ['']
    parts=max(len(chunks),(len(images)+7)//8)
    records=[]
    for i in range(parts):
        data=model(ctx,'按指令提取资料，返回 JSON 对象或包含记录列表的 JSON 对象。\n指令:'+prompt+'\n资料来源:'+p.name+'\n正文:'+(chunks[i] if i<len(chunks) else ''),images=images[i*8:(i+1)*8],schema=True)
        for row in process_extract_result(data,str(p)):
            row['_source_file']=str(p.relative_to(ctx.root));row['_提取分段']=i+1;row['_结论状态']='待复核';records.append(row)
    return records


def extract_files(ctx,files,prompt,workers=1):
    results=[];errors=[]
    workers=max(1,min(5,int(workers)))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures={pool.submit(extract_one,ctx,p,prompt):p for p in files}
        for future in as_completed(futures):
            p=futures[future]
            try:results.extend(future.result())
            except Exception as e:errors.append({'文件':str(p.relative_to(ctx.root)),'错误':str(e)})
            ctx.emit({'type':'progress','phase':'文档提取 '+str(len(results))+' 条记录','filesDone':len([f for f in futures if f.done()])})
    if errors:
        ctx.json(errors,'文件处理失败.json')
        ctx.warnings.extend(x['文件']+': '+x['错误'] for x in errors)
    if not results:raise ValueError('所有文件提取失败: '+(errors[0]['错误'] if errors else '没有可提取内容'))
    return results,errors


def info_extract(ctx):
    a=ctx.args;files=ctx.files(a['folder'],DOCUMENTS)
    rows,errors=extract_files(ctx,files,a['role_prompt'],a.get('max_workers',1))
    df=pd.DataFrame(rows);ctx.report(df,'关键信息提取.xlsx')
    return ctx.result({'success':not errors,'status':'partial' if errors else 'success','processed_files':len(files)-len(errors),'failed_files':len(errors)},df)


def info_multi(ctx):
    cfg=ctx.config();base=ctx.kv(cfg,'基础配置');tasks=ctx.rows(cfg,'文件夹配置');rows=[];errors=[]
    root=ctx.path(base['处理文件夹'],True)
    companies=[p for p in root.iterdir() if p.is_dir()]
    for company in companies:
        for task in tasks:
            key=str(task['文件夹名称']);mode=task.get('匹配模式') or '完全'
            dirs=[p for p in company.iterdir() if p.is_dir() and (p.name==key if mode=='完全' else key in p.name)]
            for folder in dirs:
                files=ctx.files(folder,DOCUMENTS)
                data,failed=extract_files(ctx,files,str(task['提示词']));errors.extend(failed)
                if task.get('是否合并')=='是':
                    # Merge structured records by collecting every non-empty field, retaining source provenance.
                    merged={}
                    for item in data:
                        for k,v in item.items():
                            if v is not None and v!='':merged.setdefault(k,[]).append(v)
                    data=[{k:'；'.join(dict.fromkeys(str(x) for x in v)) for k,v in merged.items()}]
                rows.extend({'公司名称':company.name,'文件类型':key,**r} for r in data)
    if not rows:raise ValueError('没有匹配公司、分类文件夹或提取结果')
    df=pd.DataFrame(rows);ctx.report(df,'多公司信息提取.xlsx')
    return ctx.result({'success':not errors,'failed_files':len(errors)},df)


def classification(ctx):
    cfg=ctx.config();c=ctx.kv(cfg,'基础配置');categories=ctx.rows(cfg,'类型配置')
    names=[str(r['文件类别']) for r in categories];rows=[]
    if not names:raise ValueError('请配置文件类别')
    for p in ctx.files(c['处理文件夹'],DOCUMENTS):
        text,images=document(ctx,p)
        if len(text)>60000 or len(images)>8:ctx.warnings.append(p.name+': 分类依据前 60000 字及前 8 张扫描页，分类后请复核。')
        answer=model(ctx,'根据资料从以下类别中选择一个，无法归类填 未分类。返回 JSON {"类别":"...","公司名称":"...","理由":"..."}。类别说明:'+json.dumps(categories,ensure_ascii=False)+'\n文件:'+p.name+'\n资料:'+text[:60000],images=images[:8],schema=True)
        category=str(answer.get('类别','未分类'))
        if category not in names:category='未分类'
        company=safe_name(answer.get('公司名称') or '未识别公司')
        directory=ctx.output/'分类文件'
        if c.get('输出模式')=='公司':directory=directory/company/safe_name(category)
        else:directory=directory/safe_name(category)/(company if c.get('公司模式')=='多个' else '')
        directory.mkdir(parents=True,exist_ok=True);dest=directory/p.name
        if dest.exists():dest=directory/(p.stem+'_'+str(len(rows))+p.suffix)
        shutil.copy2(p,dest);ctx.register(dest)
        rows.append({'来源文件':str(p.relative_to(ctx.root)),'类别':category,'公司名称':company,'理由':str(answer.get('理由','')),'成果文件':str(dest.relative_to(ctx.root))})
    df=pd.DataFrame(rows);ctx.report(df,'文件分类清单.xlsx');return ctx.result(frame=df)


def photo_rename(ctx):
    rows=[]
    for p in ctx.files(ctx.args['directory'],IMAGES):
        answer=model(ctx,ctx.args['user_prompt']+'\n返回 JSON {"文件名":"..."}，文件名不要包含路径。',images=[p],schema=True)
        name=safe_name(answer.get('文件名') or '')
        if name=='结果':raise ValueError(p.name+': 模型未提供文件名')
        if Path(name).suffix.lower() in IMAGES:name=Path(name).stem
        dest=ctx.destination(name+p.suffix)
        i=1
        while dest.exists():dest=ctx.destination(name+'_'+str(i)+p.suffix);i+=1
        shutil.copy2(p,dest);ctx.register(dest);rows.append({'原文件':str(p.relative_to(ctx.root)),'新文件名':dest.name})
    df=pd.DataFrame(rows);ctx.report(df,'图片命名清单.xlsx');return ctx.result(frame=df)


def text_match(ctx):
    cfg=ctx.config();c=ctx.kv(cfg);data=ctx.rows(cfg,'待匹配数据');categories=ctx.rows(cfg,'分类项')
    batch=max(1,min(100,int(c.get('批处理数量') or 20)));results=[]
    for i in range(0,len(data),batch):
        rows=[{'记录编号':i+j,**r} for j,r in enumerate(data[i:i+batch])]
        answer=model(ctx,'逐行完成分类或文本提取。返回 JSON {"records":[{"记录编号":0,"匹配结果":"..."}]}。保留记录编号，不遗漏或重复。任务:'+str(c.get('任务指令及结果要求') or '')+'\n模式:'+str(c.get('任务模式') or '')+'\n分类候选:'+json.dumps(categories,ensure_ascii=False)+'\n待处理:'+json.dumps(rows,ensure_ascii=False),schema=True)
        items=answer.get('records',[])
        if sorted(r.get('记录编号') for r in items)!=list(range(i,i+len(rows))):raise ValueError('模型返回的记录编号不完整或重复')
        lookup={r['记录编号']:r for r in items}
        results.extend({**r,**lookup[i+j]} for j,r in enumerate(data[i:i+batch]))
    if not results:raise ValueError('待匹配数据为空')
    df=pd.DataFrame(results);ctx.report(df,'文本提取匹配结果.xlsx');return ctx.result(frame=df)


def auxiliary(ctx):
    cfg=ctx.config();data=pd.DataFrame(ctx.rows(cfg,'辅助核算配置'));ctx.require(data,['1级科目','辅助核算'])
    rows=[]
    for subject,group in data.groupby('1级科目',sort=False):
        answer=model(ctx,'根据会计科目和辅助核算名称、示例，选择有用的标准核算维度并判断通常余额方向。返回 JSON {"d":"借或贷","s":"辅助项1-辅助项2"}，s 只能选输入中存在的辅助核算项，按具体对象、业务项目、属性类别排序。科目:'+str(subject)+'\n辅助核算:'+group.to_json(orient='records',force_ascii=False),schema=True)
        if answer.get('d') not in ['借','贷']:raise ValueError('模型未返回有效的借贷方向')
        selected=split(answer.get('s',''),sep=r'-')
        if any(x not in set(group['辅助核算'].astype(str)) for x in selected):raise ValueError('模型选择了配置以外的辅助核算项')
        rows.append({'1级科目':subject,'方向':answer['d'],'标准辅助核算1':answer.get('s',''),'结论状态':'待复核'})
    df=pd.DataFrame(rows);ctx.report(df,'辅助核算选择结果.xlsx');return ctx.result(frame=df)


def data_field_check(ctx):
    cfg=ctx.config();ws=cfg['配置表'];values=list(ws.values);rows=[]
    scan=int(values[0][1] or 10)
    for row in values[3:]:
        category,folder=row[:2]
        if not category or not folder:continue
        if category not in cfg:raise ValueError('缺少资料类别字段表: '+str(category))
        required=[str(v) for v in next(cfg[category].values) if v is not None]
        for p in ctx.files(folder,EXCEL):
            raw=ctx.table(p,header=1,nrows=scan)
            candidates=list(raw.columns)+[v for r in raw.values for v in r if isinstance(v,str)]
            answer=model(ctx,'将标准字段与数据表头一一匹配，无法找到填 null。返回 JSON {"mapping":{"标准字段":"原表头或null"}}。禁止编造字段。标准字段:'+json.dumps(required,ensure_ascii=False)+'\n实际表头候选:'+json.dumps(candidates,ensure_ascii=False),schema=True)
            mapping=answer.get('mapping',{})
            for key in required:
                found=mapping.get(key)
                if found and found not in candidates:raise ValueError('模型给出了不存在的数据字段: '+str(found))
                rows.append({'资料类别':category,'文件':str(p.relative_to(ctx.root)),'标准字段':key,'匹配字段':found,'检查结果':'已找到' if found else '缺失'})
    if not rows:raise ValueError('配置表中没有待检查目录')
    df=pd.DataFrame(rows);ctx.report(df,'资料字段检查结果.xlsx');return ctx.result(frame=df)


def paper_verify(ctx):
    a=ctx.args;keys=split(a['match_keys']);records=[];fields=keys
    # Book records may be supplied through the task's standard project template.
    config=ctx.path(a['config_file'],True) if a.get('config_file') else ctx.root/'input'/'纸质信息核对配置表.xlsx'
    if config.exists():
        wb=ctx.config(str(config));sheet=next((s for s in wb.sheetnames if s in ['账面数据','数据源']),None)
        if sheet:records=ctx.rows(wb,sheet);fields=list(records[0]) if records else keys
    if not records:raise ValueError('纸质信息核对配置表的账面数据不能为空，请填写账面记录后执行核对')
    if not fields:raise ValueError('请填写匹配键或在配置表中提供账面字段')
    files=ctx.files(a['file_path'],DOCUMENTS)
    extracted,errors=extract_files(ctx,files,'提取字段:'+','.join(fields)+'。'+str(a.get('extra_prompt','')),a.get('max_workers',1))
    compared=match_and_compare(records,extracted,keys,fields)
    matched=[]
    for row in compared['matched']:
        out={'来源文件':row['file_path']}
        for field in fields:
            out[field+'_账面']=row['book'].get(field);out[field+'_提取']=row['extracted'].get(field);out[field+'_核对']=row['field_results'][field]
        matched.append(out)
    ctx.report({'字段核对':pd.DataFrame(matched),'账面未匹配':pd.DataFrame(compared['unmatched_book']),
                '提取未匹配':pd.DataFrame(compared['unmatched_extracted']),'原始提取':pd.DataFrame(extracted)},'纸质信息核对结果.xlsx')
    return ctx.result({'success':not errors,'matched_count':len(matched),'unmatched_book':len(compared['unmatched_book']),'unmatched_extracted':len(compared['unmatched_extracted'])},pd.DataFrame(matched or extracted))


def vouchers(ctx):
    a=ctx.args;files=ctx.files(a['source_paths'],{'.pdf',*IMAGES});cfgpath=ctx.path(a['config_file'],True) if a.get('config_file') else ctx.root/'input'/'凭证检查配置表.xlsx'
    checkpoints=[]
    if cfgpath.exists():
        cfg=ctx.config(str(cfgpath));checkpoints=ctx.rows(cfg,'核查要点')
    if not checkpoints:raise ValueError('请先在 input/凭证检查配置表.xlsx 的核查要点中填写检查要求')
    prompt='按配置的科目、必要附件和核查要点检查凭证。返回 JSON，含科目、凭证号、日期、金额、每个要点的检查结果和证据。无法判断的要点填 待补证据；结果只作审计线索。核查配置:'+json.dumps(checkpoints,ensure_ascii=False)
    rows,errors=extract_files(ctx,files,prompt,3);df=pd.DataFrame(rows);ctx.report(df,'凭证检查结果.xlsx')
    return ctx.result({'success':not errors,'failed_files':len(errors)},df)


def formula_filler(ctx):
    cfg=ctx.config();jobs=ctx.rows(cfg);prompts={str(r['公式类型']):str(r.get('提示词') or '') for r in ctx.rows(cfg,'提示词')};rows=[]
    for task in jobs:
        if task.get('是否执行')!='是':continue
        p=ctx.path(task['文件路径'],True);dest=ctx.copy(p);wb=workbook(dest)
        target=wb[str(task['填充表'])];source=wb[str(task['链接表'])]
        def cells(ws):
            if ws.max_row*ws.max_column>50000:raise ValueError('公式填充单表超过 50000 个单元格，请缩小范围')
            return [{'cell':c.coordinate,'value':c.value} for row in ws for c in row if c.value is not None]
        answer=model(ctx,'根据两张工作表和任务要求生成待写入公式。返回 JSON {"formulas":[{"cell":"B2","formula":"=SUM(...)"}]}。只填写目标表指定数值单元格。仅引用当前工作簿工作表，禁止外部链接、网络函数、宏及命令。要求:'+prompts.get(str(task['公式类型']),'')+'\n目标表:'+json.dumps(cells(target),ensure_ascii=False,default=str)+'\n来源表名:'+source.title+'\n来源表:'+json.dumps(cells(source),ensure_ascii=False,default=str),schema=True)
        formulas=answer.get('formulas',[])
        if not formulas:raise ValueError('模型未生成可用公式')
        from openpyxl.utils import coordinate_to_tuple,column_index_from_string
        written=set();bounds=(target.max_row,target.max_column)
        for item in formulas:
            cell=str(item['cell']).upper();formula=str(item['formula'])
            if not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*',cell) or not formula.startswith('='):raise ValueError('模型返回的单元格或公式无效')
            row,col=coordinate_to_tuple(cell)
            if row>bounds[0] or col>bounds[1] or cell in written:raise ValueError('模型公式超出目标表范围或重复填写单元格: '+cell)
            if any(x in formula for x in ['[',']','|','http:','https:']) or re.search(r'\b(?:WEBSERVICE|HYPERLINK|CALL|EXEC|REGISTER|RTD|DDE)\s*\(',formula,re.I):raise ValueError('模型公式包含不允许的外部调用')
            tokenizer=Tokenizer(formula)
            for token in tokenizer.items:
                if token.type=='OPERAND' and token.subtype=='RANGE' and '!' in token.value:
                    sheet,reference=token.value.rsplit('!',1)
                    if sheet.strip("'").replace("''","'") not in wb.sheetnames:raise ValueError('模型公式引用了不存在的工作表: '+sheet)
            if '#REF!' in formula:raise ValueError('模型公式包含无效引用')
            target[cell]=formula
            written.add(cell)
        if wb.calculation:wb.calculation.fullCalcOnLoad=True
        wb.save(dest);wb.close();ctx.register(dest);rows.append({'文件':p.name,'填充表':target.title,'公式数量':len(formulas),'状态':'待 Excel 重算与人工复核'})
    if not rows:raise ValueError('配置表中没有标记 是 的执行任务')
    return ctx.result(frame=pd.DataFrame(rows))


def allocation(ctx):
    cfg=ctx.config();c=ctx.kv(cfg);standards=ctx.rows(cfg,'基础代码')+ctx.rows(cfg,'其它项');rows=[]
    for p in ctx.files(c['文件夹'],EXCEL,c.get('科目余额表文件名关键词') or ''):
        text,_=document(ctx,p)
        if len(text)>60000:raise ValueError(p.name+': 科目数据超过单次派工分析容量，请拆分工作表后处理')
        answer=model(ctx,'将数据中的科目映射到配置的报表科目，形成派工清单。只用候选科目编码。返回 JSON {"records":[{"科目编码":"...","科目名称":"...","科目级别":"...","是否需要派工":"是或否","依据":"..."}]}。候选配置:'+json.dumps(standards,ensure_ascii=False)+'\n数据:'+text[:60000],schema=True)
        allowed={str(r.get('科目编码')) for r in standards}
        for row in answer.get('records',[]):
            if str(row.get('科目编码')) not in allowed:raise ValueError('模型使用了配置之外的派工科目编码')
            rows.append({'来源文件':str(p.relative_to(ctx.root)),**row,'结论状态':'待复核'})
    if not rows:raise ValueError('没有生成派工记录')
    df=pd.DataFrame(rows);ctx.report(df,'科目分配派工结果.xlsx');return ctx.result(frame=df)


TASKS={'info_extract':info_extract,'info_extract_multi':info_multi,'file_classification':classification,'photo_rename':photo_rename,
       'text_match':text_match,'auxiliary_selection':auxiliary,'data_field_check':data_field_check,'paper_verify':paper_verify,
       'voucher_check':vouchers,'ai_formula_filler':formula_filler,'subject_allocation':allocation}
