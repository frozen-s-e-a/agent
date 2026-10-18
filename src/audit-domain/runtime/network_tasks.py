from datetime import datetime,timedelta
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
import json
import re
import time
import urllib.parse
import zipfile

import pandas as pd
from openpyxl import Workbook

from recovered.ipo import run_ipo
from .context import split, safe_name, json_value
from .transport import Session

CICPA='https://zsk-cmis.cicpa.org.cn'


def session(ctx):
    return ctx.session or Session()


def ipo(ctx,exchange):
    a=ctx.args
    # The recovered exchange routines publish in their supplied work directory;
    # a task-specific staging directory keeps outputs separate from prior runs.
    work=ctx.output/'IPO资料';work.mkdir()
    rows=run_ipo(exchange,session(ctx),work,a['start_date'],a['end_date'],a['update_mode'])
    for p in work.rglob('*'):
        if p.is_file():
            if p.suffix.lower()=='.pdf' and not p.read_bytes().startswith(b'%PDF'):
                raise RuntimeError('披露文件下载返回了非 PDF 内容: '+p.name)
            ctx.register(p)
    # Recovered return structure differs from the client representation.
    if isinstance(rows,dict):data=rows
    else:data={'downloaded':rows}
    return ctx.result(data)


class Tables(HTMLParser):
    def __init__(self):
        super().__init__();self.tables=[];self.table=None;self.row=None;self.cell=None;self.links=[]
    def handle_starttag(self,tag,attrs):
        if tag=='table':self.table=[]
        elif tag=='tr' and self.table is not None:self.row=[]
        elif tag in ['td','th'] and self.row is not None:self.cell=[]
        elif tag=='a':
            href=dict(attrs).get('href')
            if href:self.links.append(href)
    def handle_data(self,value):
        if self.cell is not None:self.cell.append(value)
    def handle_endtag(self,tag):
        if tag in ['td','th'] and self.cell is not None:
            self.row.append(''.join(self.cell).strip());self.cell=None
        elif tag=='tr' and self.row is not None:
            if self.row:self.table.append(self.row)
            self.row=None
        elif tag=='table' and self.table is not None:
            self.tables.append(self.table);self.table=None


def currency(ctx):
    s=session(ctx);url='https://www.safe.gov.cn/safe/2020/1218/17833.html'
    page=s.get(url,timeout=30);page.raise_for_status();parser=Tables();parser.feed(page.text)
    links=[urllib.parse.urljoin(url,h) for h in parser.links if re.search(r'\.xls[x]?(?:\?|$)',h,re.I)]
    if not links:raise RuntimeError('外汇管理局历史数据页面未找到 Excel 数据链接，需核对接口变化')
    reply=s.get(links[0],timeout=60);reply.raise_for_status()
    hist=pd.read_excel(BytesIO(reply.content))
    # Locate the date header without assuming a particular file preamble.
    if not any('日期' in str(c) for c in hist.columns):
        raw=pd.read_excel(BytesIO(reply.content),header=None)
        row=next((i for i,r in raw.iterrows() if any(str(v).strip()=='日期' for v in r)),None)
        if row is None:raise RuntimeError('历史汇率表缺少日期字段')
        hist=pd.read_excel(BytesIO(reply.content),header=row)
    date_col=next(c for c in hist if '日期' in str(c));hist=hist.rename(columns={date_col:'日期'})
    hist['日期']=pd.to_datetime(hist['日期'],errors='coerce');hist=hist.dropna(subset=['日期'])
    if hist.empty:raise RuntimeError('历史汇率表没有有效日期')
    start=(hist['日期'].max()+timedelta(days=1)).strftime('%Y-%m-%d');end=datetime.now().strftime('%Y-%m-%d')
    if start<=end:
        reply=s.post('https://www.safe.gov.cn/AppStructured/hlw/RMBQuery.do',data={'startDate':start,'endDate':end,'queryYN':'true'},timeout=30)
        reply.raise_for_status();p=Tables();p.feed(reply.text)
        for table in p.tables:
            header=[re.sub(r'\s+', '', h) for h in table[0]] if table else []
            if '日期' in header:
                recent=pd.DataFrame([r for r in table[1:] if len(r)==len(header)],columns=header);recent['日期']=pd.to_datetime(recent['日期'],errors='coerce')
                hist=pd.concat([hist,recent],ignore_index=True)
    hist=hist.dropna(subset=['日期']).drop_duplicates('日期',keep='last').sort_values('日期')
    for c in hist:
        if c!='日期':hist[c]=pd.to_numeric(hist[c],errors='coerce')
    ctx.report(hist,'外汇管理局汇率数据.xlsx')
    ctx.json({'history_source':links[0],'increment_source':'https://www.safe.gov.cn/AppStructured/hlw/RMBQuery.do','retrievedAt':datetime.now().isoformat()},'汇率数据来源.json')
    return ctx.result({'latest_date':str(hist['日期'].max().date())},hist)


def announcement(ctx):
    a=ctx.args;s=session(ctx);rows=[]
    market=a.get('market_type','京沪深')
    plate='hke' if '港' in market else 'third' if '三板' in market else 'szse'
    filename={'szse':'szse_stock','hke':'hke_stock','third':'gfzr_stock'}[plate]
    stock=s.get('https://www.cninfo.com.cn/new/data/'+filename+'.json',timeout=30).json()
    stockmap={str(r['code']):r for r in stock.get('stockList',[])}
    category={'年报':'category_ndbg_szsh','半年报':'category_bndbg_szsh','一季报':'category_yjdbg_szsh','三季报':'category_sjdbg_szsh',
              '业绩预告':'category_yjyg_szsh','中介报告':'category_zjbg_szsh','首发':'category_sf_szsh','增发':'category_zf_szsh'}.get(a.get('category_name'),'')
    if a.get('category_name') and not category:raise ValueError('尚未支持的公告类型，请使用标题关键词或留空类型')
    for code in a['stock_codes']:
        code=str(code).zfill(6);info=stockmap.get(code)
        if not info:raise ValueError('无法映射股票代码: '+code)
        page=1
        while True:
            reply=s.post('https://www.cninfo.com.cn/new/hisAnnouncement/query',data={'pageNum':page,'pageSize':30,'tabName':'fulltext','column':plate,
                         'stock':code+','+info['orgId'],'category':category,'searchkey':a.get('keyword',''),'seDate':a['start_date']+'~'+a['end_date'],'isHLtitle':'true'},timeout=30)
            data=reply.json();items=data.get('announcements') or []
            for item in items:
                title=re.sub('<[^>]+>','',item['announcementTitle']);date=datetime.fromtimestamp(item['announcementTime']/1000).strftime('%Y-%m-%d')
                dest=ctx.destination(code+'_'+date+'_'+safe_name(title)+'.pdf')
                url=urllib.parse.urljoin('https://static.cninfo.com.cn/',item['adjunctUrl'])
                response=s.get(url,timeout=60);response.raise_for_status()
                if not response.content.startswith(b'%PDF'):raise RuntimeError('公告下载返回非 PDF 内容: '+title)
                dest.write_bytes(response.content);ctx.register(dest);rows.append({'股票代码':code,'公司简称':item.get('secName',info.get('zwjc','')),'公告日期':date,'标题':title,'文件路径':str(dest.relative_to(ctx.root)),'来源网址':url})
            if not items or page*30>=int(data.get('totalAnnouncement',len(items))):break
            page+=1
            if page>500:raise RuntimeError('公告查询超过 500 页，请缩小日期范围')
    frame=pd.DataFrame(rows);ctx.report(frame,'上市公司公告清单.xlsx');return ctx.result({'download_count':len(rows)},frame)


def cicpa_session(ctx):
    raw=ctx.args.get('cookie_json') or ctx.settings.get('cicpaCookies')
    if not raw:raise ValueError('请先在运行设置中点击 登录注协；工商在线查询需要有效登录状态')
    cookies=json_value(raw,{})
    if isinstance(cookies,list):cookies={str(c['name']):str(c['value']) for c in cookies if c.get('name') and c.get('value')}
    if isinstance(cookies,dict) and 'cookies' in cookies:cookies=cookies['cookies']
    if isinstance(cookies,list):cookies={str(c['name']):str(c['value']) for c in cookies if c.get('name') and c.get('value')}
    if not isinstance(cookies,dict) or not cookies:raise ValueError('Cookie JSON 格式无效')
    cookies={str(k):str(v.get('value','') if isinstance(v,dict) else v) for k,v in cookies.items()}
    cookie='; '.join(str(k)+'='+str(v) for k,v in cookies.items())
    xsrf=urllib.parse.unquote(cookies.get('XSRF-TOKEN',''))
    headers={'Cookie':cookie,'X-Xsrf-Token':xsrf,'Referer':CICPA+'/companylibrarynew/'}
    return ctx.session or Session(headers),headers


def api_data(response):
    response.raise_for_status();data=response.json()
    if not isinstance(data,dict) or data.get('status_code') not in [0,'0']:
        raise RuntimeError('注协接口拒绝请求或授权失效，请更新登录授权')
    return data.get('data',{})


def cicpa_search(s,headers,name):
    return api_data(s.post(CICPA+'/open/industry_chain_api/v1/search/home_search',json={'condition':[],'query':name,'from':0,'page':1,'page_size':5,'size':5,'type':'company','ranges':[]},headers=headers,timeout=30)).get('list',[])


def cicpa_detail(s,headers,org):
    result={}
    for key,path in [('basic_info','/v3/find_company_basic_info?orgid='),('shareholders','/v4/stock_holder_newest?orgid='),('main_persons','/v3/main_person_list?orgid='),('equity','/v3/atlas/enterprise_equity?orgid=')]:
        suffix='&page=1&pagesize=50&pageSize=50&page_size=50' if key=='shareholders' else '&page=1&pageSize=50' if key=='main_persons' else ''
        result[key]=api_data(s.get(CICPA+'/open/enterprise_info_api'+path+urllib.parse.quote(str(org))+suffix,headers=headers,timeout=30))
    return result


def safe_zip(ctx,content):
    dest=ctx.output/'工商全维度数据';dest.mkdir(exist_ok=True)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        infos=archive.infolist()
        if len(infos)>10000 or sum(i.file_size for i in infos)>500_000_000:raise ValueError('导出压缩包超出文件数量或解压大小上限')
        for item in infos:
            target=(dest/item.filename).resolve()
            if not target.is_relative_to(dest.resolve()):raise ValueError('导出压缩包包含项目外路径')
            if (item.external_attr>>16)&0o170000==0o120000:raise ValueError('导出压缩包包含符号链接')
            if item.is_dir():target.mkdir(parents=True,exist_ok=True);continue
            target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(archive.read(item));ctx.register(target)
    return dest


def cicpa_export(ctx,s,headers,names):
    wb=Workbook();ws=wb.active;ws.append(['企业名称'])
    for name in names:ws.append([name])
    buf=BytesIO();wb.save(buf);buf.seek(0)
    upload=api_data(s.post(CICPA+'/open/industry_chain_api/v1/batch_search/upload_local',files={'file':('companies.xlsx',buf,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},headers=headers,timeout=60))
    batch=upload.get('batch_id')
    if not batch:raise RuntimeError('注协上传未返回批次编号')
    api_data(s.post(CICPA+'/open/industry_chain_api/v1/batch_search/enter_search',json={'batch_id':batch,'page':1,'page_size':10,'type':'credit'},headers=headers,timeout=60))
    dimensions=api_data(s.get(CICPA+'/open/industry_chain_api/v1/search/export/get_dimension_class',params={'internal':'false'},headers=headers,timeout=30))
    codes=[]
    def walk(items):
        for item in items:
            if item.get('code'):codes.append(item['code'])
            if item.get('children'):walk(item['children'])
    walk(dimensions if isinstance(dimensions,list) else dimensions.get('list',[]))
    if not codes:raise RuntimeError('注协服务未返回可导出维度')
    trigger_time=datetime.now().replace(second=0,microsecond=0)
    trigger=api_data(s.post(CICPA+'/open/industry_chain_api/v1/batch_search/enter_search_out_type',json={'batch_id':batch,'dimensions':codes},headers=headers,timeout=60))
    task_id=trigger.get('task_id') if isinstance(trigger,dict) else None
    for attempt in range(120):
        ctx.emit({'type':'progress','phase':'等待注协生成全维度导出文件','attempt':attempt+1})
        tasks=api_data(s.get(CICPA+'/open/industry_chain_api/v1/download/task_list',params={'page':1,'pageSize':20,'fresh':str(int(time.time()*1000))},headers=headers,timeout=30))
        if isinstance(tasks,dict):tasks=tasks.get('list',tasks.get('data',[]))
        selected=[]
        for task in tasks:
            if task.get('batch_id') is not None:
                belongs=str(task['batch_id'])==str(batch)
            elif task_id is not None:belongs=str(task.get('task_id') or task.get('id'))==str(task_id)
            else:
                date=pd.to_datetime(task.get('date'),errors='coerce')
                belongs=task.get('type')=='批量查询' and pd.notna(date) and date>=trigger_time
            if belongs and task.get('status') in [1,'1']:selected.append(task)
        if len(selected)>1:raise RuntimeError('发现同一时段多个注协导出任务，无法确认本次文件，请在下载中心核对批次后重试')
        for task in selected:
            url=task.get('download_url') or task.get('url') or task.get('file_url')
            if url:
                full=urllib.parse.urljoin(CICPA,url)
                if urllib.parse.urlparse(full).hostname not in ['zsk-cmis.cicpa.org.cn','cmis.cicpa.org.cn']:raise RuntimeError('注协返回了未知的下载域名')
                reply=s.get(full,headers=headers,timeout=120);reply.raise_for_status()
                if not zipfile.is_zipfile(BytesIO(reply.content)):raise RuntimeError('注协下载返回的内容不是有效 ZIP 文件，请检查授权和导出任务')
                path=ctx.destination('工商信息全维度.zip');path.write_bytes(reply.content);ctx.register(path)
                directory=safe_zip(ctx,reply.content)
                ctx.json({'batch_id':batch,'task_id':task.get('task_id') or task.get('id'),'requested_companies':names,'download_url':full,'retrievedAt':datetime.now().isoformat()},'工商导出来源.json')
                return {'batch_id':batch,'data_dir':str(directory.relative_to(ctx.root)),'dimensions':len(codes)}
        time.sleep(5)
    raise RuntimeError('注协导出等待超时；批次 '+str(batch)+' 可在注协下载中心查看后重试')


def cicpa(ctx):
    a=ctx.args;s,headers=cicpa_session(ctx);mode=a.get('query_mode','search')
    if mode=='check_cookies':
        api_data(s.get(CICPA+'/open/industry_chain_api/v1/search/get_user_index',headers=headers,timeout=30))
        return ctx.result({'authorized':True,'message':'注协授权检查通过'})
    names=split(a.get('company_name'))
    if mode=='export':
        if not names:raise ValueError('请输入需要导出的企业名称')
        return ctx.result(cicpa_export(ctx,s,headers,names))
    rows=[];details=[]
    if mode=='detail' and a.get('org_id'):details=[cicpa_detail(s,headers,a['org_id'])]
    else:
        if not names:raise ValueError('请输入企业名称或企业 org_id')
        for name in names:
            matches=cicpa_search(s,headers,name);rows.extend(matches)
            if mode in ['detail','subsidiary']:
                exact=[r for r in matches if str(r.get('company_name') or r.get('name') or '')==name]
                if len(exact)!=1:raise ValueError('企业名称未唯一匹配，请先搜索并确认 org_id: '+name)
                org=exact[0].get('org_id') or exact[0].get('orgid');details.append(cicpa_detail(s,headers,org))
    if mode=='subsidiary':
        rows=[]
        for item in details:
            for child in item.get('equity',{}).get('invests',{}).get('children',[]):
                value=str(child.get('czbl',child.get('ratio','0'))).rstrip('%')
                try:ratio=float(value)
                except ValueError:continue
                if ratio>=float(a.get('subsidiary_threshold',50)):rows.append(child)
    elif mode not in ['detail','search']:raise ValueError('不支持的注协查询模式')
    ctx.json({'search_results':rows,'details':details},'工商查询详情.json')
    frame=pd.DataFrame(rows);ctx.report(frame,'工商查询结果.xlsx');return ctx.result({'details':details},frame)


TASKS={**{n:(lambda c,e=e:ipo(c,e)) for n,e in [('ipo_sh','sh'),('ipo_sz','sz'),('ipo_bj','bj')]},
       'currency':currency,'announcement':announcement,'cicpa_query':cicpa}
