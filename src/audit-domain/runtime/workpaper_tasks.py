from copy import copy
import json
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

from recovered import workpapers
from recovered.accounting import age_count
from .context import split, json_value, number, EXCEL, safe_name
from .workbook_tasks import workbook


def config_dir(ctx):
    return ctx.path(Path(ctx.args['work_dir'])/'config')


def read_configs(ctx):
    directory=config_dir(ctx)
    if not directory.is_dir():raise ValueError('工作目录缺少 config 目录，请先保存明细底稿配置')
    result={}
    for p in directory.glob('*.json'):
        result[p.stem]=json.loads(ctx.path(p,True).read_text(encoding='utf-8-sig'))
    return normalized_configs(result)


def normalized_configs(configs):
    result=dict(configs)
    field=result.get('field_mapping',{})
    result['_field_options']={k:field.get(k,[]) for k in ['text_fields','uniqueness_fields']} if isinstance(field,dict) else {}
    result['field_mapping']=field.get('mapping',{k:v for k,v in field.items() if k not in ['text_fields','uniqueness_fields']}) if isinstance(field,dict) else field
    subject=result.get('subject_mapping',[])
    result['subject_mapping']=subject.get('mapping',[]) if isinstance(subject,dict) else subject
    template=result.get('template_assignment',{})
    if 'assignment' in template:
        result['template_assignment']={r['report_subject']:{k:v for k,v in r.items() if k!='report_subject'} for r in template['assignment']}
    return result


def save_config(ctx):
    kind=str(ctx.args['config_type'])
    if not re.fullmatch(r'[\w\u4e00-\u9fff-]{1,80}',kind):raise ValueError('配置类型不能含路径或特殊符号')
    data=json_value(ctx.args['config_data'])
    if kind=='subject_mapping' and not (isinstance(data,list) or isinstance(data,dict) and isinstance(data.get('mapping'),list)):raise ValueError('科目映射应为 JSON 列表或包含 mapping 列表的对象')
    if kind!='subject_mapping' and not isinstance(data,dict):raise ValueError('配置应为 JSON 对象')
    # Nested template and data paths are checked before any persistence.
    def check(value,key=''):
        if isinstance(value,dict):
            for k,v in value.items():
                if k=='sources' and isinstance(v,dict):
                    for p in v.values():ctx.path(p,True)
                else:check(v,str(k))
        elif isinstance(value,list):
            for item in value:check(item,key)
        elif isinstance(value,str) and value and (key.endswith(('_path','_dir')) or key in ['file_path','template_path']):ctx.path(value)
    check(data)
    directory=config_dir(ctx);directory.mkdir(parents=True,exist_ok=True)
    path=directory/(kind+'.json')
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    ctx.register(path)
    return ctx.result({'config_type':kind,'saved_to':str(path),'message':'配置已保存'})


def inspect_templates(ctx):
    ctx.path(ctx.args['template_dir'],True)
    result=workpapers.inspect_templates(ctx.path(ctx.args['template_dir'],True))
    ctx.json(result,'底稿模板批注.json')
    frame=pd.DataFrame([{'模板':v['file_name'],'工作表':s['name'],'批注字段':'、'.join(s['annotations'])} for v in result for s in v['sheets']])
    return ctx.result({'templates':result},frame)


def validate(ctx,configs=None):
    configs=configs or read_configs(ctx)
    fm=configs.get('field_mapping',{});sm=configs.get('subject_mapping',[]);ta=configs.get('template_assignment',{})
    converted={}
    for subject,item in ta.items():
        item=dict(item);item['template_path']=str(ctx.path(item['template_path'],True));converted[subject]=item
    p=ctx.args.get('file_path') or configs.get('paths',{}).get('file_path') or configs.get('paths',{}).get('source_data_path')
    if p:p=str(ctx.path(p,True))
    result=workpapers.validate_config(fm,[m for m in sm if enabled(m)],converted,None)
    if p:
        for path in ctx.files(p):
            columns=ctx.table(path,nrows=0).columns
            for standard,expression in fm.items():
                if not isinstance(expression,str):
                    result['errors'].append({'config':'field_mapping','field':standard,'error':'字段映射值必须是列名文本'});continue
                for side in expression.lstrip('$').split('/'):
                    terms=[side] if side in columns else side.split('-')
                    missing=[t for t in terms if t not in columns]
                    if missing:result['errors'].append({'config':'field_mapping','field':standard,'error':path.name+': 缺少列 '+','.join(missing)})
    for report,item in converted.items():
        wb=workbook(Path(item['template_path']))
        if item.get('sheet_name') not in wb:
            result['errors'].append({'config':'template_assignment','field':report,'error':'模板工作表不存在: '+str(item.get('sheet_name'))})
        wb.close()
    if result['errors']:result['status']='invalid';result['validation_passed']=False
    else:result['validation_passed']=True
    return result


def enabled(mapping):
    return mapping.get('generate',True) not in [False,'否',0]


def validate_task(ctx):
    result=validate(ctx);ctx.json(result,'底稿配置检查.json')
    return ctx.result(result)


def mapped_field(df,expression,standard,symbol,configs):
    lookup=expression.startswith('$');expression=expression.lstrip('$')
    if expression in df:terms=[expression]
    else:
        sides=expression.split('/')
        side=sides[1] if len(sides)>1 and symbol=='贷' else sides[0]
        terms=[side] if side in df else side.split('-')
    missing=[t for t in terms if t not in df]
    if missing:raise ValueError('字段映射引用不存在的列: '+','.join(missing))
    if standard in ['期初余额','期末余额','本期借方','本期贷方']:
        value=sum((number(df[t]) for t in terms),pd.Series(0.,index=df.index))
    else:
        value=pd.Series('',index=df.index,dtype=object)
        for term in terms:
            source=df[term].fillna('').astype(str)
            value=value.where(value.astype(str).str.strip()!='',source)
    if lookup:
        custom=configs.get(standard,configs.get({'披露':'disclosure','关联方':'related_party'}.get(standard,''),{}))
        mapping=custom.get('mapping',custom)
        if not isinstance(mapping,dict) or not mapping:raise ValueError('缺少 $ 字段的查找配置: '+standard)
        unmatched=set(value.astype(str))-{''}-set(map(str,mapping))
        if unmatched:raise ValueError(standard+'查找配置缺少值: '+','.join(sorted(unmatched)[:10]))
        value=value.astype(str).map(mapping).fillna('')
    return value


def template_info(ws):
    columns={};group=[];first=last=None;names=[]
    for row in ws:
        for cell in row:
            if not cell.comment:continue
            tag=cell.comment.text.strip();names.append(tag)
            if tag in ['表首','数据开始','first_row']:first=cell.row
            elif tag in ['表尾','数据结束','last_row']:last=cell.row
            elif tag=='尾行':last=cell.row-1
            elif tag=='项目名称':continue
            else:
                columns[cell.column]=tag.rstrip('#')
                if tag.endswith('#'):group.append(tag[:-1])
                if first is None:first=cell.row+1
    return columns,group,first,last,names


def prepared_data(ctx,configs,mapping,source,groups,needs_age):
    fm=configs['field_mapping'];symbol=mapping.get('symbol',mapping.get('direction','借'))
    if symbol not in ['借','贷']:raise ValueError('科目方向必须为 借 或 贷')
    files=ctx.files(source)
    if needs_age and len(files)>1:
        numbered={}
        for path in files:
            match=re.match(r'^(\d+)[.、_-]',path.name)
            if not match:raise ValueError('多期账龄文件必须由近及远编号，如 1.2026.xlsx、2.2025.xlsx: '+path.name)
            period=int(match[1])
            if period in numbered:raise ValueError('同一数据源的多期文件编号重复: '+str(period))
            numbered[period]=path
        if sorted(numbered)!=list(range(1,len(numbered)+1)):raise ValueError('多期账龄文件编号必须从 1 连续排列')
        periods=[[numbered[i]] for i in sorted(numbered)]
    else:periods=[files]
    numeric=['期初余额','期末余额','本期借方','本期贷方'];frames=[]
    for period,paths in enumerate(periods,1):
        parts=[]
        for path in paths:
            raw=ctx.table(path,dtype=str);data=raw.copy()
            for standard,expression in fm.items():data[standard]=mapped_field(raw,str(expression),standard,symbol,configs)
            data['科目名称']=data['科目名称'].astype(str).str.strip()
            subject=str(mapping.get('subject_name') or mapping.get('subject') or '').strip()
            data=data[data['科目名称']==subject]
            if mapping.get('company_name'):data=data[data['公司名称'].astype(str)==str(mapping['company_name'])]
            if mapping.get('sign')=='负':
                for field in ['期初余额','期末余额']:data[field]=-number(data[field])
            data['_来源文件']=str(path.relative_to(ctx.root));parts.append(data)
        joined=pd.concat(parts,ignore_index=True)
        keys=list(dict.fromkeys(['公司名称','科目名称']+groups+configs['_field_options'].get('uniqueness_fields',[])))
        ctx.require(joined,keys)
        if joined.empty:frames.append(joined);continue
        agg={c:'sum' if c in numeric else (lambda values:'；'.join(dict.fromkeys(values.fillna('').astype(str)))) if c=='_来源文件' else 'first' for c in joined if c not in keys}
        joined[keys]=joined[keys].fillna('');frames.append(joined.groupby(keys,dropna=False,sort=False).agg(agg).reset_index())
    keys=list(dict.fromkeys(['公司名称','科目名称']+groups+configs['_field_options'].get('uniqueness_fields',[])))
    all_keys=list(dict.fromkeys(tuple(row) for frame in frames for row in frame[keys].itertuples(index=False,name=None)))
    tables=[{tuple(row[k] for k in keys):row.to_dict() for _,row in f.iterrows()} for f in frames]
    rows=[]
    for key in all_keys:
        fallback=next(t[key] for t in tables if key in t);row=dict(fallback)
        current=tables[0].get(key,{})
        for field in numeric:row[field]=float(current.get(field,0))
        for i,table in enumerate(tables,1):
            for field,value in table.get(key,{}).items():row[field+':'+str(i)]=value
            for field in numeric:row.setdefault(field+':'+str(i),0.)
        row['_来源文件']='；'.join(dict.fromkeys(t[key].get('_来源文件','') for t in tables if key in t))
        if needs_age:
            info=age_count(float(row['期初余额:'+str(len(tables))]),[float(row['本期借方:'+str(i)]) for i in range(1,len(tables)+1)],[float(row['本期贷方:'+str(i)]) for i in range(1,len(tables)+1)],symbol)
            buckets=info[info['direction']]
            for i,value in buckets.items():row['账龄:'+str(i)]=value
            row['账龄核对差额']=round(sum(buckets.values())-row['期末余额'],2)
        rows.append(row)
    return pd.DataFrame(rows)


def expand_template(ws,first,last,rows):
    extra=rows-(last-first+1)
    if extra<=0:return last
    # Insert within the data range so totals spanning that range expand too.
    insertion=last
    prototype=[(copy(c._style),c.value,c.coordinate,c.data_type) for c in ws[first]]
    merged=list(ws.merged_cells.ranges)
    for region in merged:ws.unmerge_cells(str(region))
    ws.insert_rows(insertion,extra)
    def shift_reference(formula):
        tokenizer=Tokenizer(formula)
        for token in tokenizer.items:
            if token.type=='OPERAND' and token.subtype=='RANGE':
                prefix,mark,reference=token.value.rpartition('!')
                if mark and prefix.strip("'")!=ws.title:continue
                value=reference if mark else token.value
                region=re.fullmatch(r'(\$?[A-Z]{1,3}\$?)(\d+):(\$?[A-Z]{1,3}\$?)(\d+)',value)
                if region and int(region[2])<=first and int(region[4])>=last:
                    value=region[1]+region[2]+':'+region[3]+str(int(region[4])+extra)
                else:value=re.sub(r'(\$?[A-Z]{1,3}\$?)(\d+)',lambda m:m[1]+str(int(m[2])+extra if int(m[2])>=insertion else int(m[2])),value)
                token.value=prefix+'!'+value if mark else value
        return '='+''.join(t.value for t in tokenizer.items)
    for row in ws:
        for cell in row:
            if cell.data_type=='f':cell.value=shift_reference(cell.value)
    for r in range(insertion,insertion+extra):
        for c in range(1,ws.max_column+1):
            style,value,coordinate,kind=prototype[c-1];cell=ws.cell(r,c)
            cell._style=copy(style)
            if kind=='f':cell.value=Translator(value,origin=coordinate).translate_formula(cell.coordinate)
    for region in merged:
        minrow=region.min_row+extra if region.min_row>=insertion else region.min_row
        maxrow=region.max_row+extra if region.max_row>=insertion else region.max_row
        ws.merge_cells(start_row=minrow,end_row=maxrow,start_column=region.min_col,end_column=region.max_col)
    return last+extra


def generation(ctx):
    configs=read_configs(ctx);valid=validate(ctx,configs)
    if not valid['validation_passed']:raise ValueError('底稿配置未通过校验: '+json.dumps(valid['errors'],ensure_ascii=False))
    paths=configs.get('paths',{});sources=dict(paths.get('sources',{}))
    default=ctx.args.get('file_path') or paths.get('file_path') or paths.get('source_data_path')
    if default:sources['default']=default
    elif len(sources)==1:sources['default']=next(iter(sources.values()))
    if not sources:raise ValueError('请选择明细底稿数据源')
    ta=configs['template_assignment'];plans={};counts=[];age_reports={}
    for mapping in configs['subject_mapping']:
        if not enabled(mapping):continue
        report=mapping.get('report_subject');subject=mapping.get('subject_name') or mapping.get('subject')
        if not report or not subject:raise ValueError('每条科目映射必须填写 subject_name 和 report_subject')
        source=mapping.get('source','default')
        if source not in sources:raise ValueError('科目映射引用了未配置的数据源: '+str(source))
        info=ta[report];template=ctx.path(info['template_path'],True);book=workbook(template)
        columns,groups,first,last,tags=template_info(book[info['sheet_name']]);book.close()
        needs_age=any(re.fullmatch(r'账龄:\d+',t.rstrip('#')) for t in tags)
        selected=prepared_data(ctx,configs,mapping,sources[source],groups,needs_age)
        if selected.empty:
            ctx.warnings.append('数据源没有匹配科目: '+str(subject)+'（'+str(source)+'）')
            if paths.get('copy_empty_template') is True:
                companies=[]
                if mapping.get('company_name'):companies=[str(mapping['company_name'])]
                else:
                    for path in ctx.files(sources[source]):
                        raw=ctx.table(path,dtype=str)
                        values=mapped_field(raw,str(configs['field_mapping']['公司名称']),'公司名称',mapping.get('symbol',mapping.get('direction','借')),configs)
                        companies.extend(str(v) for v in values.dropna().unique() if str(v).strip())
                fields=list(dict.fromkeys(list(configs['field_mapping'])+list(columns.values())))
                for company in dict.fromkeys(companies):
                    plans.setdefault((company,template),{}).setdefault(info['sheet_name'],[]).append((report,pd.DataFrame(columns=fields)))
            continue
        if needs_age:
            age_reports[safe_name(report+'_'+subject)]=selected.copy()
            mismatch=selected['账龄核对差额'].abs()>=.01
            if mismatch.any():ctx.warnings.append(str(subject)+': '+str(int(mismatch.sum()))+' 行账龄合计与期末余额不一致，请核对多期发生额与期初余额。')
        for company,part in selected.groupby('公司名称',dropna=False):
            key=(str(company),template)
            plans.setdefault(key,{}).setdefault(info['sheet_name'],[]).append((report,part.copy()))
    if not plans:raise ValueError('数据源中没有匹配配置的科目或公司，请核对科目映射')
    for (company,template),sheet_plan in plans.items():
        dest=ctx.destination(safe_name(company)+'_'+template.stem+template.suffix);wb=workbook(template)
        if dest.exists():
            import hashlib
            dest=ctx.destination(safe_name(company)+'_'+template.stem+'_'+hashlib.sha256(str(template).encode()).hexdigest()[:8]+template.suffix)
        for name,parts in sheet_plan.items():
            ws=wb[name];data=pd.concat([p for _,p in parts],ignore_index=True)
            columns,groups,first,last,tags=template_info(ws)
            for row in ws:
                for cell in row:
                    if cell.comment and cell.comment.text.strip()=='项目名称':cell.value='；'.join(dict.fromkeys(r for r,_ in parts))
            if not columns:raise ValueError(template.name+'/'+name+': 模板没有与映射字段对应的批注')
            ctx.require(data,list(columns.values()))
            if first is None:raise ValueError(template.name+'/'+name+': 缺少数据起始位置')
            if last is None:raise ValueError(template.name+'/'+name+': 缺少表尾批注，无法确定写入区域')
            if last<first:raise ValueError('表尾必须在表首之后')
            last=expand_template(ws,first,last,len(data))
            for rownum in range(first,last+1):
                for col,field in columns.items():
                    cell=ws.cell(rownum,col)
                    if rownum>first and not cell.has_style:cell._style=copy(ws.cell(first,col)._style)
                    value=data.iloc[rownum-first][field] if rownum-first<len(data) else None
                    if pd.isna(value):value=None
                    if isinstance(value,str) and value.startswith(('=','+','@')):value="'"+value
                    cell.value=value
                    if field in configs['_field_options'].get('text_fields',[]):cell.number_format='@'
            counts.append({'公司名称':company,'底稿模板':template.name,'工作表':name,'报表科目':'；'.join(dict.fromkeys(r for r,_ in parts)),'写入行数':len(data),'是否空底稿':data.empty,'期初余额':data['期初余额'].sum(),'期末余额':data['期末余额'].sum()})
        if wb.calculation:wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
        wb.save(dest);wb.close();ctx.register(dest)
    if age_reports:ctx.report(age_reports,'账龄基础数据.xlsx')
    ctx.report(pd.DataFrame(counts),'明细底稿生成清单.xlsx')
    ctx.json(counts,'明细底稿任务计划.json')
    ctx.warnings.append('模板中的公式在 Excel 打开时重算；请复核模板公式和生成清单金额。')
    return ctx.result({'company_count':len({k[0] for k in plans})},pd.DataFrame(counts))


def subject_after(ctx):
    cfg=ctx.config();balance=ctx.kv(cfg,'余额配置');rules=ctx.rows(cfg,'辅助核算配置');outputs=[]
    for item in ctx.rows(cfg,'数据源'):
        for p in ctx.files(item['文件或文件夹路径'],EXCEL):
            df=ctx.table(p);ctx.require(df,['1级科目'])
            for rule in rules:
                mask=df['1级科目'].astype(str)==str(rule['1级科目'])
                for key,value in rule.items():
                    if not key.startswith('标准辅助核算'):continue
                    names=split(value,sep=r'-')
                    ctx.require(df,names)
                    if key not in df:df[key]=''
                    for col in names:
                        current=df.loc[mask,key].fillna('').astype(str)
                        source=df.loc[mask,col].fillna('').astype(str)
                        df.loc[mask,key]=current.where(current.str.strip()!='',source)
                if '标准方向' not in df:df['标准方向']=''
                df.loc[mask,'标准方向']=rule.get('方向') or ''
            for label in ['期初','期末']:
                debit,credit=balance.get(label+'借方余额'),balance.get(label+'贷方余额')
                value,direction=balance.get(label+'余额'),balance.get(label+'方向')
                if debit or credit:
                    ctx.require(df,[debit,credit]);signed=(number(df[debit]) if debit else 0)-(number(df[credit]) if credit else 0)
                elif value and direction:
                    ctx.require(df,[value,direction]);signed=number(df[value])*df[direction].map({'借':1,'贷':-1}).fillna(0)
                else:raise ValueError(label+'余额配置须指定借贷列，或余额列和方向列')
                factor=df['标准方向'].map({'借':1,'贷':-1})
                if factor.isna().any():raise ValueError('辅助核算配置未覆盖所有科目的标准方向')
                df[label+'标准余额']=signed*factor
            ctx.report(df,p.stem+'_标准余额.xlsx');outputs.append({'文件':p.name,'行数':len(df)})
    return ctx.result(frame=pd.DataFrame(outputs))


def xlsx_writer(ctx):
    cfg=ctx.config();paths=ctx.kv(cfg,'文件储存');folder=ctx.path(paths['数据源文件夹路径'],True)
    partitioned=paths.get('源文件是否分公司保存')=='是'
    maps={}
    for sheet in ['科目余额表映射','辅助科目余额表映射','序时账映射']:
        ws=cfg[sheet];values=list(ws.values)
        maps[sheet.removesuffix('映射')]=[(str(a),str(b)) for a,b in zip(values[0][1:],values[1][1:]) if a is not None and b is not None]
    companies=[p.name for p in folder.iterdir() if p.is_dir()] if partitioned else sorted({p.stem.split('-')[0] for p in ctx.files(folder,EXCEL)})
    rows=[]
    for company in companies:
        sheets={}
        for table,mapping in maps.items():
            files=ctx.files(folder/company,EXCEL) if partitioned else ctx.files(folder,EXCEL)
            matches=[p for p in files if table in p.stem and (partitioned or p.stem.startswith(company))]
            if not matches:continue
            parts=[]
            for p in matches:
                df=ctx.table(p);ctx.require(df,[k for k,v in mapping])
                parts.append(df[[k for k,v in mapping]].rename(columns=dict(mapping)))
            sheets[table]=pd.concat(parts,ignore_index=True)
        if not sheets:raise ValueError('没有与小账套映射匹配的来源文件: '+company)
        ctx.report(sheets,company+'_小账套.xlsx');rows.append({'公司名称':company,'生成表数':len(sheets)})
    return ctx.result(frame=pd.DataFrame(rows))


TASKS={'detailed_table_save_config':save_config,'detailed_table_inspect_template':inspect_templates,
       'detailed_table_validate':validate_task,'detailed_table_generate':generation,
       'subject_clean_after':subject_after,'xlsx_writer':xlsx_writer}
