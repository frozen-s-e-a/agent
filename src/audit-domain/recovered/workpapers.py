"""Recovered workpaper configuration, inspection and table-preparation rules."""
from pathlib import Path
import re
import pandas as pd
from openpyxl import load_workbook

STANDARD_FIELDS=['公司名称','科目名称','本期借方','本期贷方','期初余额','期末余额']

def first2second(df,columns=None):
    key=df.columns[0]
    columns=[df.columns[1]] if columns is None else columns
    result=df.fillna('').copy()
    result[key]=result[key].astype(str)
    if len(columns)==1:return result.set_index(key)[columns[0]].to_dict()
    return result.set_index(key)[columns].to_dict(orient='index')

def extract_field_mapping_from_df(df):
    return first2second(df)

def extract_template_assignment_from_df(df):
    selected=df.iloc[:,1:3].copy()
    selected.columns=['template_path','sheet_name']
    selected.index=df.iloc[:,0].fillna('').astype(str)
    return selected.to_dict(orient='index')

def computing_field_info(fill_column_name):
    # A slash separates debit/credit; a hyphen combines fields within a side.
    sides=re.split('/',fill_column_name)
    front=re.split('-',sides[0])
    back=re.split('-',sides[1]) if len(sides)>1 else front
    return {'借':front,'贷':back}

def force_numeric_columns(df,columns):
    columns=[columns] if isinstance(columns,str) else columns
    for column in columns:
        if column not in df:df[column]=0;continue
        values=df[column].astype(str).str.replace(',','',regex=False)
        df[column]=pd.to_numeric(values,errors='coerce').fillna(0)
    return df

def groupby(df,columns):
    agg={column:'sum' if pd.api.types.is_numeric_dtype(df[column]) else 'first' for column in df.columns if column not in columns}
    return df.groupby(columns).agg(agg).reset_index().reindex(columns=df.columns)

def aggregate_dataframe(df,columns):
    for column in columns:
        df[column]=df[column].fillna('')
    return groupby(df,columns)

def company2data(df,cname_company):
    return {company:df[df[cname_company]==company] for company in df[cname_company].dropna().unique()}

def check_dfs(dfs):return all(df.empty for df in dfs)
def calc_rows(dfs):return sum(len(df) for df in dfs)

def reverse_dataframe(dfs,filter_column,change_columns,values):
    for df in dfs:
        mask=df[filter_column].isin(values)
        for column in change_columns:df.loc[mask,column]=-df.loc[mask,column]

def excel_files(directory):
    return sorted(path for path in Path(directory).rglob('*') if path.is_file() and path.suffix.lower() in {'.xlsx','.xlsm','.xls'} and not path.name.startswith('~$'))

def inspect_file(file_path,sheet_name=None,preview_rows=5):
    path=Path(file_path)
    if not path.exists():return {'status':'error','message':'文件不存在: '+str(file_path)}
    if path.suffix.lower() not in {'.xlsx','.xlsm','.xls'}:
        return {'status':'error','message':'不支持的文件格式: '+path.suffix.lower()}
    try:df=pd.read_excel(path,sheet_name=0 if sheet_name is None else sheet_name)
    except Exception as error:return {'status':'error','message':str(error)}
    companies=[];subjects=[]
    for column in df:
        label=str(column).lower()
        if ('公司' in label or 'company' in label) and not companies:companies=df[column].dropna().unique().tolist()
        if ('科目' in label or 'subject' in label) and not subjects:subjects=df[column].dropna().unique().tolist()
    return {'status':'success','type':'file','file_path':str(file_path),'columns':df.columns.tolist(),
            'row_count':len(df),'companies':companies,'subjects':subjects,'preview':df.head(min(preview_rows,len(df))).fillna('').to_dict('records')}

def inspect_folder(folder_path,preview_rows=5):
    paths=excel_files(folder_path)
    if not paths:return {'status':'error','message':'文件夹中没有 Excel 文件'}
    first=inspect_file(str(paths[0]),preview_rows=preview_rows)
    if first['status']!='success':return first
    return {'status':'success','type':'folder','file_path':str(folder_path),'file_count':len(paths),
            'columns':first['columns'],'row_count':first['row_count'],'periods':len(paths),'preview':first['preview']}

def inspect_templates(template_dir):
    results=[]
    for path in excel_files(template_dir):
        wb=load_workbook(path,read_only=False,data_only=False,keep_vba=path.suffix.lower()=='.xlsm')
        sheets=[]
        for ws in wb.worksheets:
            annotations=[]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.comment:
                        text=cell.comment.text.strip()
                        if text and text not in annotations:annotations.append(text)
            sheets.append({'name':ws.title,'annotations':annotations})
        results.append({'file_name':path.name,'file_path':str(path),'sheets':sheets})
    return results

def validate_config(field_mapping,subject_mapping,template_assignment,data_source_path=None):
    errors=[]
    fm=field_mapping or {}; ta=template_assignment or {}; sm=subject_mapping or []
    for field in STANDARD_FIELDS:
        if not fm.get(field):errors.append({'config':'field_mapping','field':field,'error':'未映射'})
    for subject in sm:
        report=subject.get('report_subject')
        if report and report not in ta:
            errors.append({'config':'template_assignment','field':report,'error':f"报表科目 '{report}' 在模板匹配中无对应条目"})
    for report,info in ta.items():
        path=info.get('template_path','')
        if not path or not Path(path).is_file():
            errors.append({'config':'template_assignment','field':report,'error':'模板文件不存在: '+str(path)})
    if data_source_path and Path(data_source_path).is_file():
        columns=pd.read_excel(data_source_path,nrows=0).columns.tolist()
        for field,actual in fm.items():
            missing=[piece for piece in re.split(r'[/\-]',actual) if piece and piece not in columns]
            if missing:errors.append({'config':'field_mapping','field':field,'error':'数据源中缺少列: '+','.join(missing)})
    if errors:return {'status':'invalid','errors':errors,'summary':{'total_errors':len(errors)}}
    return {'status':'valid','errors':[],'summary':{'subjects':len(sm),'companies':0,
            'templates':len({item.get('template_path') for item in ta.values()}),'estimated_files':0}}

REGISTRY={
 ('detailed_table/_core','first2second'):first2second,
 ('detailed_table/_core','extract_field_mapping_from_df'):extract_field_mapping_from_df,
 ('detailed_table/_core','extract_template_assignment_from_df'):extract_template_assignment_from_df,
 ('detailed_table/_core','computing_field_info'):computing_field_info,
 ('detailed_table/_core','force_numeric_columns'):force_numeric_columns,
 ('detailed_table/_core','_aggregate_dataframe'):aggregate_dataframe,
 ('detailed_table/_core','_groupby'):groupby,
 ('detailed_table/_core','_company2data'):company2data,
 ('detailed_table/_core','_check_dfs'):check_dfs,
 ('detailed_table/_core','_calc_rows'):calc_rows,
 ('detailed_table/_core','reverse_dataframe'):reverse_dataframe,
 ('detailed_table/_core','inspect_file'):inspect_file,
 ('detailed_table/_core','inspect_folder'):inspect_folder,
 ('detailed_table/_core','inspect_templates'):inspect_templates,
 ('detailed_table/_core','validate_config'):validate_config,
}
