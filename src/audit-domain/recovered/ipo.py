"""Independently recovered IPO disclosure selection, naming and download flow.

An HTTP session is injected. Importing or testing this module makes no requests.
The original exchange protocol is preserved; current live API availability is
separate from compatibility against the captured synthetic responses.
"""
from pathlib import Path
from datetime import datetime
from urllib.parse import urljoin
import json
import re
import time
from openpyxl import Workbook
from openpyxl.styles import Font
from .kernels import sanitize_filename,process_jsonp

# Recovered from the native dictionary-construction instructions, not inferred
# from exchange terminology. See assembly/modules__ipo_{sh,bj}.annotated.asm.
SSE_STATUS={'1':'已受理','2':'已问询','3':'上市委审议','4':'提交注册','5':'注册结果',
            '7':'中止及财报更新','8':'中止','9':'上市委审议','10':'补充审核'}
BSE_STATUS={'P01':'已受理','P02':'已问询','P03':'通过','P04':'未通过','P05':'暂缓审议',
            'P06':'提交注册','P07':'注册','P08':'不予注册','P09':'中止','P10':'终止'}
NAMES={'bj':'IPO北交所','sh':'IPO上交所','sz':'IPO深交所'}
BSE_AGENT='Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/137.0.0.0 Safari/537.36'

def parse_iso(value):
    return datetime.strptime(str(value).split(' ')[0],'%Y-%m-%d')

def parse_compact(value):
    # The current SSE project endpoint includes HHMMSS after the date.
    return datetime.strptime(str(value).split(' ')[0][:8],'%Y%m%d')

def download_file(session,url,save_folder_path,file_name,exchange='sz'):
    path=Path(save_folder_path)/sanitize_filename(file_name)
    if path.exists():
        return str(path)
    kwargs={'timeout':60 if exchange=='sh' else 30}
    if exchange=='bj':kwargs['headers']={'User-Agent':BSE_AGENT}
    response=session.get(url,**kwargs)
    response.raise_for_status()
    with path.open('wb') as target:
        target.write(response.content)
    return str(path)

def extract_people(person_list,key):
    return ','.join(person.get('i_p_personName','') for person in person_list if person.get('i_p_jobTitle')==key)

def write_index(rows,path,exchange):
    if not rows:return
    headers=['文件本地路径','公司简称','项目状态','文件日期' if exchange=='sz' else '更新日期','文件标题','行业']
    if exchange=='sh':headers+=['保荐机构','签字会计师','律师事务所']
    wb=Workbook()
    ws=wb.active
    ws.title='IPO披露文件列表'
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
        cell=ws.cell(ws.max_row,1)
        absolute=Path(row['文件本地路径']).absolute()
        cell.hyperlink=absolute.as_uri()
        cell.font=Font(color='0000FF',underline='single')
    wb.save(path)

def run_ipo(exchange,session,work_dir,start_date,end_date,update_mode='项目',industry_map=None,stopped=None,clock=time.time):
    if exchange not in NAMES:raise ValueError('Unknown exchange: '+str(exchange))
    start,end=parse_iso(start_date),parse_iso(end_date)
    mode=update_mode if update_mode in ['项目','文件'] else '项目'
    base=Path(work_dir)/'output'/NAMES[exchange]
    base.mkdir(parents=True,exist_ok=True)
    index=base.parent/(NAMES[exchange]+'.xlsx')
    stopped=stopped or (lambda:False)
    rows=[]
    def add(company,nickname,status,industry,title,file_date,url,extra=None):
        if not url:return
        if mode=='文件':
            try:
                if not start<=parse_iso(file_date)<=end:return
            except (ValueError,TypeError):return
        clean_title=str(title).replace('.pdf','')
        filename=sanitize_filename(f'{nickname}_{status}_{file_date}_{clean_title}.pdf')
        folder=base/sanitize_filename(industry)
        folder.mkdir(parents=True,exist_ok=True)
        destination=download_file(session,url,str(folder),filename,exchange)
        row={'文件本地路径':destination,'公司简称':nickname,'项目状态':status,
             '文件日期' if exchange=='sz' else '更新日期':file_date,'文件标题':title,'行业':industry}
        row.update(extra or {})
        rows.append(row)

    if exchange=='sz':
        response=session.get('https://listing.szse.cn/api/ras/projectrends/query',
                             params={'bizType':'1','pageIndex':'0','pageSize':'10000'},timeout=60)
        response.raise_for_status()
        projects=response.json().get('data',[])
        projects=sorted(projects,key=lambda value:value.get('updtdt','1900-01-01'),reverse=True)
        for project in projects:
            if stopped():break
            try:pdate=parse_iso(project.get('updtdt',''))
            except (ValueError,TypeError):continue
            if pdate<start:break
            if mode=='项目' and pdate>end:continue
            company=project.get('cmpnm','未知公司')
            response=session.get('https://listing.szse.cn/api/ras/projectrends/details?id='+str(project.get('prjid','')),timeout=20)
            response.raise_for_status()
            detail=response.json().get('data',{})
            status,industry=detail.get('prjst','未知状态'),detail.get('csrcind','未知行业')
            files=[item for item in detail.get('disclosureMaterials',[]) if item.get('matnm')=='招股说明书']
            files+=detail.get('enquiryResponseAttachment',[])
            for item in files:
                add(company,company,status,industry,item.get('dfnm',''),item.get('ddt',''),
                    urljoin('https://reportdocs.static.szse.cn/',item.get('dfpth','')))
    elif exchange=='sh':
        page=1
        finish=False
        while not finish and not stopped():
            stamp=int(clock()*1000)
            params={'jsonCallBack':'jsonpCallback'+str(stamp),'sqlId':'SH_XM_LB','issueMarketType':'1,2',
                    'order':'updateDate|desc,stockAuditNum|desc','isPagination':'true','pageHelp.pageSize':'20',
                    'pageHelp.pageNo':str(page),'_':stamp}
            response=session.get('https://query.sse.com.cn/commonSoaQuery.do',params=params,headers={'Referer':'https://listing.sse.com.cn/'},timeout=30)
            response.raise_for_status()
            payload=process_jsonp(response.text)
            projects=(payload or {}).get('result',[])
            if not projects:break
            for project in projects:
                if stopped():finish=True;break
                try:pdate=parse_compact(project.get('updateDate',''))
                except (TypeError,ValueError):continue
                if pdate<start:finish=True;break
                if mode=='项目' and pdate>end:continue
                company=project.get('stockAuditName','未知公司')
                status_code=str(project.get('currStatus',''))
                status=SSE_STATUS.get(status_code,status_code)
                issuer=(project.get('stockIssuer') or [{}])[0]
                nickname=issuer.get('s_issueCompanyAbbrName',company)
                industry=issuer.get('s_csrcCodeDesc','未知行业')
                intermediary=project.get('intermediary',[])
                extra={'保荐机构':intermediary[0].get('i_intermediaryName','') if len(intermediary)>0 else '',
                       '签字会计师':extract_people(intermediary[1].get('i_person',[]),'签字会计师') if len(intermediary)>1 else '',
                       '律师事务所':intermediary[2].get('i_intermediaryName','') if len(intermediary)>2 else ''}
                stamp=int(clock()*1000)
                params={'jsonCallBack':'jsonpCallback'+str(stamp),'sqlId':'GP_COMMON_FILE_SEARCH',
                        'auditId':project.get('stockAuditNum'),'marketType':'1,2','isPagination':'false','_':stamp}
                response=session.get('https://query.sse.com.cn/commonSoaQuery.do',params=params,headers={'Referer':'https://listing.sse.com.cn/'},timeout=30)
                response.raise_for_status()
                payload=process_jsonp(response.text)
                for item in (payload or {}).get('result',[]):
                    if item.get('fileType')!=30 and item.get('fileTypeMap')!='I3010':continue
                    try:fdate=parse_compact(item.get('fileUpdTime','')).strftime('%Y-%m-%d')
                    except (TypeError,ValueError):continue
                    src=item.get('filePath','')
                    if not src:continue
                    url=src if src.startswith('http') else 'https://static.sse.com.cn/stock'+src
                    add(company,nickname,status,industry,item.get('fileTitle','').strip(),fdate,url,extra)
            page+=1
    else:
        page=0
        finish=False
        while not finish and not stopped():
            response=session.post('https://www.bse.cn/projectNewsController/infoResult.do?callback=jQuery_1',
                                  data={'page':str(page),'isNewThree':'1','sortfield':'updateDate','sorttype':'desc'},timeout=20)
            response.raise_for_status()
            payload=process_jsonp(response.text)
            projects=payload[0].get('listInfo',{}).get('content',[]) if payload else []
            if not projects:break
            for project in projects:
                if stopped():finish=True;break
                company=project.get('companyName','未知公司')
                response=session.post('https://www.bse.cn/projectNewsController/infoDetailResult.do',
                                      data={'id':project.get('id'),'callback':'jQuery_2'},timeout=20)
                response.raise_for_status()
                payload=process_jsonp(response.text)
                if not payload:continue
                detail=payload[0]
                basic=detail.get('projectNews',{})
                timestamp=basic.get('updateDate',{}).get('time')
                if timestamp is None:continue
                pdate=datetime.fromtimestamp(timestamp/1000)
                if pdate<start:finish=True;break
                if mode=='项目' and pdate>end:continue
                nickname=basic.get('stockName',company)
                status_code=basic.get('status','')
                status=BSE_STATUS.get(status_code,status_code)
                industry=(industry_map or {}).get(company,'未知行业')
                disclosed=detail.get('xxgkInfo',{}).get('GPFXSMS',{})
                files=[]
                for key in ['SYG','SBG','BHG']:files+=disclosed.get(key,[])
                files+=detail.get('wxhfhInfo',[])
                for item in files:
                    src=item.get('destFilePath','')
                    if not src:continue
                    add(company,nickname,status,industry,item.get('disclosureTitle',''),item.get('publishDate',''),
                        src if src.startswith('http') else 'https://www.bse.cn'+src)
            page+=1
    write_index(rows,index,exchange)
    return {'status':'success','msg':f'完成。下载了 {len(rows)} 个文件。','output_dir':str(base)}
