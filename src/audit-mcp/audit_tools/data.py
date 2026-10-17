"""资料识别、导入和查看（《接口约定》3.5、3.11，实现设计第六节）。

每个项目一个分析库。资料按模型确认的字段含义导入成 D- 表，每行带来源列（文件、工作表、行号）；
数值解析严格，无效值单独记录，不当作 0。
"""
import re
from decimal import Decimal

import openpyxl
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from .common import Result, ToolError, plain, require
from .files import read_csv_rows
from .numbers import BLANK, INVALID, PARSERS, kind_of
from .store import sql_name

TYPES = {'decimal': 'DECIMAL(38,6)', 'date': 'DATE', 'text': 'VARCHAR', 'integer': 'BIGINT'}
KINDS = {'trial_balance', 'journal', 'bank_statement', 'ledger_detail', 'aging', 'contract', 'other'}
MEANINGS = {
    'voucher_date': '凭证日期', 'voucher_no': '凭证号', 'account_code': '科目编码', 'account_name': '科目名称', 'summary': '摘要',
    'debit': '借方金额', 'credit': '贷方金额', 'opening': '期初余额', 'closing': '期末余额', 'direction': '余额方向',
    'counterparty': '往来单位', 'customer': '客户', 'supplier': '供应商', 'amount': '金额', 'balance': '余额', 'date': '日期',
    'invoice_no': '发票号', 'invoice_date': '发票日期', 'contract_no': '合同号', 'quantity': '数量', 'price': '单价',
    'inflow': '收入金额', 'outflow': '支出金额', 'currency': '币种', 'department': '部门', 'project': '项目',
}
TOTAL = re.compile(r'^(合\s*计|小\s*计|总\s*计|本页合计|累计)')
# 无效值超过这个比例或个数的列，导入不通过（实现设计 6.3）。
INVALID_RATIO, INVALID_COUNT = Decimal('0.001'), 10


def read_sheets(path):
    """→ [(工作表名, [[值…]…])]。"""
    if path.suffix.lower() == '.csv':
        return [('CSV', read_csv_rows(path))]
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ToolError('IO_UNREADABLE', '无法读取资料：%s' % error)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            continue
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        while rows and all(v is None or (isinstance(v, str) and not v.strip()) for v in rows[-1]):
            rows.pop()
        sheets.append((ws.title, rows))
    wb.close()
    return sheets


def blank_row(row):
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def header_score(row, following):
    cells = [v for v in row if v is not None and str(v).strip()]
    if len(cells) < 2:
        return 0
    texts = [v for v in cells if kind_of(v) == 'text']
    if len(texts) / len(cells) < 0.8:
        return 0
    data = [r for r in following if not blank_row(r)][:5]
    if not data:
        return 0
    typed = sum(1 for r in data for v in r if kind_of(v) in ('decimal', 'date'))
    return len(texts) + (2 if typed else 0)


def find_tables(rows):
    """一个工作表里找出若干张表：[(表头行号列表, 首行, 末行)]，行号从 1 起。"""
    tables, start = [], 0
    while start < len(rows):
        best, best_score = None, 0
        for i in range(start, min(len(rows), start + 30)):
            score = header_score(rows[i], rows[i + 1:i + 6])
            if score > best_score:
                best, best_score = i, score
        if best is None:
            break
        headers = [best]
        # 两层表头：下一行也全是文字，且再下一行才是数据。
        if best + 2 < len(rows) and header_score(rows[best + 1], rows[best + 2:best + 7]) and \
                sum(1 for v in rows[best] if v is None) > 0:
            headers.append(best + 1)
        first = headers[-1] + 1
        last, blanks = first, 0
        for j in range(first, len(rows)):
            if blank_row(rows[j]):
                blanks += 1
                if blanks >= 2:
                    break
            else:
                blanks, last = 0, j
        if last >= first:
            tables.append(([h + 1 for h in headers], first + 1, last + 1))
        start = last + 1
    return tables


def guess_kind(headers):
    text = ''.join(headers)
    if '科目' in text and ('期末' in text or '余额' in text) and '凭证' not in text:
        return 'trial_balance'
    if ('凭证' in text or '日期' in text) and '借方' in text and '贷方' in text:
        return 'journal'
    if ('收入' in text or '支出' in text) and '余额' in text:
        return 'bank_statement'
    return None


def inspect(ctx, args):
    fid, path, name = ctx.files.resolve(require(args, 'file'))
    found, candidates = [], ctx.store.index('candidates')
    for sheet, rows in read_sheets(path):
        for n, (header_rows, first, last) in enumerate(find_tables(rows), start=1):
            width = max(len(rows[r - 1]) for r in range(first, last + 1))
            columns, totals = [], []
            for r in range(first, last + 1):
                lead = [str(v).strip() for v in rows[r - 1][:3] if v is not None and str(v).strip()]
                if lead and TOTAL.match(lead[0]):
                    totals.append(r)
            for c in range(width):
                header = '/'.join(str(rows[h - 1][c]).strip() for h in header_rows if c < len(rows[h - 1]) and rows[h - 1][c] is not None and str(rows[h - 1][c]).strip())
                values = [rows[r - 1][c] if c < len(rows[r - 1]) else None for r in range(first, last + 1) if r not in totals]
                types, nulls, samples = {}, 0, []
                for v in values:
                    k = kind_of(v)
                    if k is None:
                        nulls += 1
                        continue
                    types[k] = types.get(k, 0) + 1
                    shown = plain(v) if not isinstance(v, str) else v.strip()
                    if len(samples) < 3 and shown not in samples:
                        samples.append(shown)
                if not header and not types:
                    continue
                columns.append({'col': get_column_letter(c + 1), 'header': header, 'types': types, 'nulls': nulls, 'samples': samples})
            key = '%s/%s/%d' % (fid, sheet, n)
            rows_count = last - first + 1 - len(totals)
            item = {'candidate': key, 'sheet': sheet, 'headerRows': header_rows, 'dataRange': '%s%d:%s%d' % ('A', first, get_column_letter(width), last),
                    'rows': rows_count, 'columns': columns, 'totalsRows': totals, 'kindGuess': guess_kind([c['header'] for c in columns])}
            found.append(item)
            candidates[key] = {'file': fid, 'name': name, 'sheet': sheet, 'headerRows': header_rows, 'first': first, 'last': last, 'totals': totals,
                               'headers': {c['col']: c['header'] for c in columns}}
    ctx.store.save_index('candidates', candidates)
    if not found:
        raise ToolError('DATA_NO_TABLE', '%s 里没有识别出带表头的表格' % name, hint='检查文件是否只有图片或说明文字')
    return Result({'file': fid, 'tables': found})


def quoted(name):
    return '"%s"' % name.replace('"', '""')


def import_table(ctx, args):
    key = require(args, 'candidate')
    candidates = ctx.store.index('candidates')
    if key not in candidates:
        raise ToolError('REF_NOT_FOUND', '找不到资料候选表 %s' % key, hint='先用 data_inspect 识别资料')
    cand = candidates[key]
    kind = args.get('kind') or 'other'
    if kind not in KINDS:
        raise ToolError('ARG_INVALID', '资料类型 %s 无效，应为 %s 之一' % (kind, '、'.join(sorted(KINDS))))
    fields = require(args, 'fields', list)
    if not fields:
        raise ToolError('ARG_INVALID', 'fields 不能为空')
    names, cols = set(), []
    for f in fields:
        col, name, typ = str(f.get('col', '')).upper(), str(f.get('name', '')).strip(), f.get('type', 'text')
        if col not in cand['headers']:
            raise ToolError('REF_NOT_FOUND', '列 %s 不在资料表中' % col, hint='可用的列：%s' % '、'.join('%s（%s）' % kv for kv in cand['headers'].items()))
        if not name or name.startswith('_') or name in names:
            raise ToolError('ARG_INVALID', '字段名 %r 为空、重复或以下划线开头' % name)
        if typ not in TYPES:
            raise ToolError('ARG_INVALID', '字段 %s 的类型 %s 无效，应为 decimal、date、text 或 integer' % (name, typ))
        names.add(name)
        cols.append((column_index_from_string(col) - 1, col, name, typ))
    exclude = bool(args.get('excludeInvalid'))
    _, path, _ = ctx.files.resolve(cand['file'])
    rows = dict(read_sheets(path)).get(cand['sheet'])
    if rows is None:
        raise ToolError('REF_NOT_FOUND', '资料里没有工作表 %s' % cand['sheet'])

    records, invalid, excluded = [], {}, []
    for r in range(cand['first'], cand['last'] + 1):
        if r in cand['totals']:
            excluded.append({'row': r, 'reason': '合计行'})
            continue
        raw = rows[r - 1] if r - 1 < len(rows) else []
        values, bad = [], []
        for index, col, name, typ in cols:
            value = raw[index] if index < len(raw) else None
            parsed = PARSERS[typ](value)
            if parsed is INVALID:
                bad.append((col, value))
                values.append(None)
            else:
                values.append(None if parsed is BLANK else parsed)
        if all(v is None for v in values) and not bad:
            continue
        for col, value in bad:
            entry = invalid.setdefault(col, {'col': col, 'count': 0, 'samples': []})
            entry['count'] += 1
            if len(entry['samples']) < 5:
                entry['samples'].append({'row': r, 'value': str(value)})
        if bad and exclude:
            excluded.append({'row': r, 'reason': '无效值', 'cols': [c for c, _ in bad]})
            continue
        records.append((values, r, bad))

    total = len(records) + sum(1 for e in excluded if e['reason'] == '无效值')
    if invalid and not exclude:
        worst = max(invalid.values(), key=lambda e: e['count'])
        if worst['count'] > INVALID_COUNT or Decimal(worst['count']) > INVALID_RATIO * max(total, 1):
            field = next(name for _, col, name, _ in cols if col == worst['col'])
            sample = worst['samples'][0]
            raise ToolError('DATA_INVALID_VALUES', '%s列有 %d 个值无法解析（如第 %d 行“%s”），超过阈值' % (field, worst['count'], sample['row'], sample['value']),
                            hint='可以排除这些行后继续导入，或修正资料后重新添加', detail={'invalid': list(invalid.values())})

    tables = ctx.store.index('tables')
    ref = tables.get(key) or ctx.store.allocate('D')
    version = (ctx.store.get(ref)['version'] + 1) if ctx.store.exists(ref) else 1
    name = sql_name(ref)
    with ctx.store.lock:
        cur = ctx.store.cursor()
        cur.execute('DROP TABLE IF EXISTS %s' % name)
        cur.execute('DROP TABLE IF EXISTS %s__invalid' % name)
        columns = ', '.join('%s %s' % (quoted(n), TYPES[t]) for _, _, n, t in cols)
        cur.execute('CREATE TABLE %s (%s, _source_file VARCHAR, _source_sheet VARCHAR, _source_row INTEGER)' % (name, columns))
        placeholders = ', '.join(['?'] * (len(cols) + 3))
        batch = []
        for values, r, _ in records:
            batch.append([str(v) if isinstance(v, Decimal) else v for v in values] + [cand['name'], cand['sheet'], r])
            if len(batch) >= 5000:
                cur.executemany('INSERT INTO %s VALUES (%s)' % (name, placeholders), batch)
                batch = []
        if batch:
            cur.executemany('INSERT INTO %s VALUES (%s)' % (name, placeholders), batch)
        cur.execute('CREATE TABLE %s__invalid (_source_row INTEGER, col VARCHAR, value VARCHAR)' % name)
        bad_cells = [[r, col, str(value)] for _, r, bad in records for col, value in bad]
        if bad_cells:
            cur.executemany('INSERT INTO %s__invalid VALUES (?, ?, ?)' % name, bad_cells)
        cur.close()
    obj = {'kind': 'table', 'id': ref, 'version': version, 'candidate': key, 'file': cand['file'], 'name': cand['name'], 'sheet': cand['sheet'],
           'type': kind, 'fields': [{'col': col, 'name': n, 'type': t, 'meaning': next((f.get('meaning') for f in fields if str(f.get('col', '')).upper() == col), None),
                                     'basis': next((f.get('basis') for f in fields if str(f.get('col', '')).upper() == col), None)} for _, col, n, t in cols],
           'rows': len(records), 'invalid': list(invalid.values()), 'excluded': excluded[:500], 'excludedCount': len(excluded)}
    ctx.store.put(ref, obj)
    tables[key] = ref
    ctx.store.save_index('tables', tables)
    data = {'table': ref, 'version': version, 'rows': len(records), 'invalid': [] if exclude else list(invalid.values()),
            'excluded': sum(1 for e in excluded if e['reason'] == '无效值'), 'totalsRows': cand['totals']}
    warnings = []
    if invalid and not exclude:
        warnings.append('有 %d 个值无法解析，已记为空值，不计入合计' % sum(e['count'] for e in invalid.values()))
    return Result(data, refs=[ref], warnings=warnings)


def table_obj(ctx, ref):
    if not isinstance(ref, str) or not re.fullmatch(r'(D|R)-\d{3,}', ref):
        raise ToolError('ARG_INVALID', '表编号应为 D- 或 R- 开头')
    return ctx.store.get(ref)


def columns_of(ctx, ref):
    """→ [(列名, 类型, 含义)]，不含来源列。"""
    obj = table_obj(ctx, ref)
    if obj['kind'] == 'table':
        return [(f['name'], f['type'], f.get('meaning')) for f in obj['fields']], obj
    return [(c['name'], c['type'], None) for c in obj['columns']], obj


def profile(ctx, args):
    ref = require(args, 'table')
    cols, obj = columns_of(ctx, ref)
    name = sql_name(ref)
    cur = ctx.store.cursor()
    out = []
    for col, typ, meaning in cols:
        q = quoted(col)
        nulls, distinct, low, high = cur.execute('SELECT COUNT(*) - COUNT(%s), COUNT(DISTINCT %s), MIN(%s), MAX(%s) FROM %s' % (q, q, q, q, name)).fetchone()
        item = {'name': col, 'type': typ, 'nulls': nulls, 'distinct': distinct, 'min': plain(low), 'max': plain(high)}
        if meaning:
            item['meaning'] = meaning
        if typ == 'decimal':
            item['sum'] = plain(cur.execute('SELECT SUM(%s) FROM %s' % (q, name)).fetchone()[0])
        item['samples'] = [plain(r[0]) for r in cur.execute('SELECT DISTINCT %s FROM %s WHERE %s IS NOT NULL LIMIT 3' % (q, name, q)).fetchall()]
        out.append(item)
    cur.close()
    data = {'table': ref, 'rows': obj['rows'], 'columns': out}
    if obj['kind'] == 'table' and obj.get('invalid'):
        data['invalid'] = [{'col': e['col'], 'count': e['count']} for e in obj['invalid']]
    return Result(data)


def meaning_column(obj, meaning):
    for f in obj.get('fields', []):
        if f.get('meaning') == meaning:
            return f['name']
    return None


def quality(ctx, args):
    """数据质量检查（实现设计 6.4）：能识别出常见类型时做对应的检查，识别不出只做通用检查。"""
    ref = require(args, 'table')
    obj = table_obj(ctx, ref)
    if obj['kind'] != 'table':
        raise ToolError('ARG_INVALID', 'data_quality 只检查资料表（D-）')
    name, cur, items = sql_name(ref), ctx.store.cursor(), []
    invalid = sum(e['count'] for e in obj.get('invalid', []))
    items.append({'check': 'invalid_values', 'passed': invalid == 0, 'severity': 'warning', 'detail': {'count': invalid, 'excludedRows': sum(1 for e in obj.get('excluded', []) if e['reason'] == '无效值')}})
    col = lambda m: meaning_column(obj, m)
    debit, credit = col('debit'), col('credit')
    if obj['type'] == 'journal' and debit and credit:
        d, c = cur.execute('SELECT COALESCE(SUM(%s),0), COALESCE(SUM(%s),0) FROM %s' % (quoted(debit), quoted(credit), name)).fetchone()
        items.append({'check': 'balanced', 'passed': d == c, 'severity': 'error', 'detail': {'debit': plain(d), 'credit': plain(c), 'diff': plain(d - c)}})
        voucher, date = col('voucher_no'), col('voucher_date') or col('date')
        if voucher:
            keys = ', '.join(quoted(k) for k in (date, voucher) if k)
            bad = cur.execute('SELECT %s, SUM(%s) - SUM(%s) AS diff FROM %s GROUP BY %s HAVING COALESCE(SUM(%s),0) <> COALESCE(SUM(%s),0) LIMIT 5'
                              % (keys, quoted(debit), quoted(credit), name, keys, quoted(debit), quoted(credit))).fetchall()
            items.append({'check': 'voucher_balanced', 'passed': not bad, 'severity': 'error', 'detail': {'samples': [[plain(v) for v in r] for r in bad]}})
    opening, closing = col('opening'), col('closing')
    if obj['type'] == 'trial_balance' and opening and closing and debit and credit:
        o, dd, cc, e = (quoted(x) for x in (opening, debit, credit, closing))
        bad = cur.execute('SELECT _source_row FROM %s WHERE NOT (COALESCE(%s,0) + COALESCE(%s,0) - COALESCE(%s,0) = COALESCE(%s,0) OR COALESCE(%s,0) - COALESCE(%s,0) + COALESCE(%s,0) = COALESCE(%s,0)) LIMIT 5'
                          % (name, o, dd, cc, e, o, dd, cc, e)).fetchall()
        items.append({'check': 'roll_forward', 'passed': not bad, 'severity': 'error', 'detail': {'rows': [r[0] for r in bad]}})
    for meaning in ('voucher_date', 'account_code', 'debit', 'credit', 'closing', 'date'):
        column = col(meaning)
        if column:
            n = cur.execute('SELECT COUNT(*) FROM %s WHERE %s IS NULL' % (name, quoted(column))).fetchone()[0]
            items.append({'check': 'nulls:%s' % column, 'passed': n == 0, 'severity': 'warning', 'detail': {'count': n}})
    cur.close()
    passed = all(i['passed'] for i in items if i['severity'] == 'error')
    return Result({'table': ref, 'passed': passed, 'items': items})


def table_view(ctx, args):
    ref = require(args, 'table')
    cols, obj = columns_of(ctx, ref)
    offset, limit = int(args.get('offset') or 0), min(int(args.get('limit') or 200), 1000)
    wanted = args.get('columns')
    if wanted:
        cols = [c for c in cols if c[0] in wanted]
    name = sql_name(ref)
    cur = ctx.store.cursor()
    total = cur.execute('SELECT COUNT(*) FROM %s' % name).fetchone()[0]
    is_data = obj['kind'] == 'table'
    order = '_source_row' if is_data else '_row'
    locate, highlight = args.get('locate'), None
    if locate is not None:
        if is_data:
            hit = cur.execute('SELECT COUNT(*) FROM %s WHERE _source_row < ?' % name, [int(locate)]).fetchone()[0]
        else:
            hit = int(locate) - 1
        offset, highlight = max(0, hit - hit % limit), int(locate)
    select = ', '.join(quoted(c[0]) for c in cols) + (', _source_file, _source_sheet, _source_row' if is_data else ', _row')
    rows = cur.execute('SELECT %s FROM %s ORDER BY %s LIMIT %d OFFSET %d' % (select, name, order, limit, offset)).fetchall()
    bad = {}
    if is_data:
        for row, col in cur.execute('SELECT _source_row, col FROM %s__invalid' % name).fetchall():
            bad.setdefault(row, set()).add(col)
    cur.close()
    letters = {f['name']: f['col'] for f in obj.get('fields', [])}
    columns = []
    for i, (col, typ, meaning) in enumerate(cols):
        item = {'key': 'c%d' % i, 'name': col, 'type': typ}
        if meaning:
            item['meaning'] = meaning
            if MEANINGS.get(meaning):
                item['meaningLabel'] = MEANINGS[meaning]
        columns.append(item)
    out = []
    for row in rows:
        values = [plain(v) for v in row[:len(cols)]]
        item = {'values': values}
        if is_data:
            item['source'] = {'file': obj['file'], 'sheet': row[len(cols) + 1], 'row': row[len(cols) + 2]}
            marks = bad.get(row[len(cols) + 2])
            if marks:
                item['invalid'] = [i for i, (col, _, _) in enumerate(cols) if letters.get(col) in marks]
        out.append(item)
    return Result({'table': ref, 'title': obj.get('name') or obj.get('purpose') or ref, 'columns': columns, 'rows': out, 'total': total, 'offset': offset, 'highlight': highlight},
                  truncated=offset + len(out) < total, next={'table': ref, 'offset': offset + len(out), 'limit': limit} if offset + len(out) < total else None)
