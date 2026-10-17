"""填写编译、写入回读、预览和单元格追溯（《接口约定》3.8～3.10，实现设计 8.1～8.5）。

写入只改计划里的单元格：在底稿快照的 OOXML 包里逐格打补丁，样式、合并、公式和其他包部件原样保留。
受影响的公式按 Excel 规则在本地重算并写回缓存；算不了的去掉旧缓存，并要求 Excel 打开时全部重算。
写好的副本重新打开逐项核对，结果写进核对表。
"""
import datetime
import re
import sys
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

from . import excel
from .common import Result, ToolError, money, plain, require
from .query import result_rows
from .workpaper import expand, region_of, split_ref

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'audit-worker'))
from workbook_plan import Book, FormulaError, NeedsInput, calculate_formula, formula_refs, patch_cell, set_full_calc  # noqa: E402

TASK = re.compile(r'^[A-Za-z0-9_-]{1,64}$')
CHECKS = ('readback', 'formulas', 'crossFoot', 'untouched', 'completeness', 'excel')
CHECK_LABELS = {'readback': '写入值回读一致', 'formulas': '受影响公式结果正确', 'crossFoot': '合计行与明细相符',
                'untouched': '计划外的单元格和样式未变', 'completeness': '待填位置已填或列为人工填写', 'excel': 'Excel 重算结果一致'}
PREVIEW_PAGE = 200


class Pending(Exception):
    """公式依赖一个本地算不了的公式。"""


def key_of(cell, sheet):
    s, c = split_ref(cell, sheet)
    if ':' in c:
        raise ToolError('ARG_INVALID', '写入位置应为单个单元格：%s' % cell)
    return s, c


def label(sheet, cell):
    return '%s!%s' % (sheet, cell)


def in_range(cell, span):
    c1, r1, c2, r2 = range_boundaries(span)
    row, col = coordinate_to_tuple(cell)
    return r1 <= row <= r2 and c1 <= col <= c2


def interiors(wp, sheet):
    """合并区域里除左上角以外的单元格。"""
    covered = set()
    for span in wp['merges'].get(sheet, []):
        cells = expand(span)
        covered.update(cells[1:])
    return covered


def number_text(value):
    """写进 <v> 的数字：去掉多余的零。"""
    value = value.normalize()
    return format(value, 'f') if value == value.to_integral_value() else format(value, 'f')


def shown(value):
    """给人看的金额：千分位、两位小数。"""
    try:
        d = Decimal(value)
    except (InvalidOperation, TypeError):
        return value
    return '{:,.2f}'.format(d.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) if d == d.quantize(Decimal('0.01')) else '{:,}'.format(d)


def output_path(task, name):
    return Path('输出') / task / ('%s-填写成果%s' % (Path(name).stem, Path(name).suffix.lower()))


def unique(root, rel):
    path = root / rel
    n = 2
    while path.exists():
        path = (root / rel).with_name('%s-%d%s' % (rel.stem, n, rel.suffix))
        n += 1
    return path


# ---------------------------------------------------------------------------
# 编译


def source_value(ctx, cache, source, where):
    """from → (值, 值类型, 来源说明)。不接受直接给的数值。"""
    if not isinstance(source, dict):
        raise ToolError('ARG_INVALID', '%s 的 from 应为对象' % where)
    if 'result' in source:
        ref = source['result']
        if ref not in cache:
            obj, rows = result_rows(ctx, ref)
            cache[ref] = (obj, rows)
        obj, rows = cache[ref]
        row, column = source.get('row'), source.get('column')
        if not isinstance(row, int) or not 1 <= row <= len(rows):
            raise ToolError('REF_ROW_OUT_OF_RANGE', '%s 引用 %s 第 %s 行，结果只有 %d 行' % (where, ref, row, len(rows)), where=where)
        names = [c['name'] for c in obj['columns']]
        if column not in names:
            raise ToolError('REF_UNKNOWN_COLUMN', '%s 引用的列 %s 不在 %s 里；结果列：%s' % (where, column, ref, '、'.join(names)), where=where)
        index = names.index(column)
        kind = obj['columns'][index]['type']
        value = rows[row - 1][index]
        return value, kind, '%s#%d' % (ref, row)
    if 'profile' in source:
        profile = ctx.store.index('profile')
        if source['profile'] not in profile:
            raise ToolError('REF_NOT_FOUND', '项目资料里没有 %s' % source['profile'], where=where, hint='先在项目资料里填写，或把这格列为人工填写')
        return profile[source['profile']], 'text', 'profile:%s' % source['profile']
    if 'text' in source:
        content = source.get('content')
        if not isinstance(source['text'], str) or not source['text'].startswith('C-') or not isinstance(content, str):
            raise ToolError('ARG_INVALID', '%s 的文字来源应为 { text: 卡片编号, content: 确认后的文字 }' % where, where=where)
        return content, 'text', source['text']
    raise ToolError('ARG_INVALID', '%s 的 from 只能是 result、profile 或 text，不能直接给数值' % where, where=where)


def scaled(value, kind, scale, rounding):
    if kind in ('decimal', 'integer'):
        d = value if isinstance(value, Decimal) else Decimal(str(value))
        if scale != 1:
            d = d / scale
        if rounding is not None:
            d = d.quantize(Decimal(1).scaleb(-rounding), rounding=ROUND_HALF_UP)
        elif scale != 1:
            d = d.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        return number_text(d), 'number'
    if kind == 'boolean':
        return bool(value), 'boolean'
    return plain(value), 'text'


def fill_compile(ctx, args):
    task = require(args, 'task', str)
    if not TASK.match(task):
        raise ToolError('ARG_INVALID', '任务编号只能含字母、数字、短横线和下划线')
    wp = ctx.store.get(require(args, 'workpaper'), 'workpaper')
    bindings = require(args, 'bindings', list)
    approved = set()
    for item in args.get('overwrite') or []:
        approved.add(label(*key_of(item, None)) if '!' in str(item) else str(item))
    errors, warnings, writes, manual, expand_rows, regions = [], [], {}, [], {}, []
    cache = {}
    for binding in bindings:
        region = binding.get('region') if isinstance(binding, dict) else None
        try:
            _, r = region_of(ctx, region)
            if region.split('#')[0] != wp['id']:
                raise ToolError('REF_NOT_FOUND', '区域 %s 不属于底稿 %s' % (region, wp['id']))
        except ToolError as e:
            errors.append({'code': e.code, 'message': e.message, 'region': region})
            continue
        sheet, span = r['sheet'], r['range'].split('!')[1]
        covered = interiors(wp, sheet)
        kinds = wp['kinds'][sheet]
        last_data = max([row for row in range(range_boundaries(span)[1], range_boundaries(span)[3] + 1) if row not in r['totals']] or [0])
        region_manual, bound, sources = [], set(), []
        for item in binding.get('manual') or []:
            try:
                s, c = split_ref(item, sheet)
            except ToolError as e:
                errors.append({'code': e.code, 'message': e.message, 'region': region, 'where': item})
                continue
            region_manual.extend(label(s, x) for x in (expand(c) if ':' in c else [c]))
        for spec in binding.get('cells') or []:
            where = spec.get('cell') if isinstance(spec, dict) else spec
            try:
                if not isinstance(spec, dict):
                    raise ToolError('ARG_INVALID', '每个写入位置应为 { cell, from }')
                s, c = key_of(spec.get('cell'), sheet)
                where = label(s, c)
                if s != sheet or not in_range(c, span):
                    if r['type'] == 'detail_table' and s == sheet and coordinate_to_tuple(c)[0] > range_boundaries(span)[3]:
                        expand_rows.setdefault(region, set()).add(coordinate_to_tuple(c)[0])
                        continue
                    raise ToolError('REF_OUTSIDE_REGION', '%s 不在区域 %s 内' % (where, r['range']), where=where)
                row = coordinate_to_tuple(c)[0]
                if r['type'] == 'detail_table' and row > last_data:
                    expand_rows.setdefault(region, set()).add(row)
                    continue
                if c in covered:
                    raise ToolError('REF_MERGED_CELL', '%s 在合并单元格内部，只能写合并区域的左上角' % where, where=where)
                if kinds.get(c) == 'formula' or c in r['formulas']:
                    raise ToolError('REF_FORMULA_CELL', '%s 是公式单元格，不能写入' % where, where=where)
                if where in writes:
                    raise ToolError('ARG_DUPLICATE_CELL', '%s 被绑定了两次' % where, where=where)
                value, kind, origin = source_value(ctx, cache, spec.get('from'), where)
                try:
                    scale = Decimal(str(spec.get('scale') or '1'))
                    rounding = spec.get('rounding')
                    rounding = int(rounding) if rounding not in (None, '') else None
                except (InvalidOperation, ValueError):
                    raise ToolError('ARG_INVALID', '%s 的 scale 或 rounding 不是数' % where, where=where)
                if scale <= 0:
                    raise ToolError('ARG_INVALID', '%s 的 scale 应大于 0' % where, where=where)
            except ToolError as e:
                errors.append({'code': e.code, 'message': e.message, 'region': region, **({'where': e.where} if e.where else {})})
                continue
            bound.add(c)
            if value is None or value == '':
                region_manual.append(where)
                warnings.append('%s 的来源值为空，已列为人工填写' % where)
                continue
            text, typ = scaled(value, kind, scale, rounding)
            old = wp['labels'][sheet].get(c) or wp['numbers'][sheet].get(c)
            existing = kinds.get(c) in ('label', 'number')
            writes[where] = {'cell': where, 'sheet': sheet, 'ref': c, 'value': text, 'valueType': typ, 'from': spec['from'], 'origin': origin,
                             'region': region, 'old': old if existing else None, 'overwrite': existing, 'approved': not existing or where in approved,
                             'scale': str(scale)}
            if origin.startswith('R-') and origin.split('#')[0] not in sources:
                sources.append(origin.split('#')[0])
        # 区域里没绑定、也没列为人工的待填格，列为人工填写，不留空白。
        listed = set(region_manual)
        missing = [label(sheet, c) for c in r['inputs'] if c not in bound and label(sheet, c) not in listed]
        if missing:
            warnings.append('%s 有 %d 个待填位置没有绑定来源，已列为人工填写：%s' % (region, len(missing), '、'.join(missing[:6])))
        manual.extend(x for x in region_manual + missing if x not in manual)
        regions.append({'region': region, 'cells': sum(1 for w in writes.values() if w['region'] == region), 'sources': sources})
    if errors:
        first = errors[0]
        raise ToolError(first['code'], first['message'], where=first.get('where'), detail={'errors': errors})
    fid, snapshot, name = ctx.files.resolve(wp['file'])
    plan = ctx.store.allocate('FP')
    output = output_path(task, name).as_posix()
    overwrites = [{'cell': w['cell'], 'old': w['old'], 'value': w['value']} for w in writes.values() if not w['approved']]
    expand_list = [{'region': region, 'rows': len(rows), 'after': max(rows) - len(rows)} for region, rows in expand_rows.items()]
    ctx.store.put(plan, {'kind': 'plan', 'id': plan, 'task': task, 'workpaper': wp['id'], 'file': fid, 'name': name, 'output': output,
                         'writes': list(writes.values()), 'manual': manual, 'overwrites': overwrites, 'expand': expand_list, 'regions': regions,
                         'warnings': warnings, 'createdAt': datetime.datetime.now().isoformat(timespec='seconds')})
    return Result({'plan': plan, 'cells': len(writes), 'manual': len(manual), 'manualCells': manual, 'overwrites': overwrites,
                   'expand': expand_list, 'output': output, 'regions': regions}, refs=[plan], warnings=warnings)


# ---------------------------------------------------------------------------
# 预览


def fill_preview(ctx, args):
    plan = ctx.store.get(require(args, 'plan'), 'plan')
    span = args.get('range')
    sheet, cells = (split_ref(span) if span else (None, None))
    offset = int(args.get('offset') or 0)

    def wanted(cell):
        if not span:
            return True
        s, c = split_ref(cell)
        return s == sheet and in_range(c, cells if ':' in cells else '%s:%s' % (cells, cells))

    items = [{'cell': w['cell'], 'value': w['value'], 'valueType': w['valueType'], 'from': w['origin'], 'old': w['old'], 'overwrite': w['overwrite']}
             for w in plan['writes'] if wanted(w['cell'])]
    items += [{'cell': c, 'manual': True} for c in plan['manual'] if wanted(c)]
    items.sort(key=lambda x: (x['cell'].rsplit('!', 1)[0], coordinate_to_tuple(x['cell'].rsplit('!', 1)[1])))
    page = items[offset:offset + PREVIEW_PAGE]
    more = offset + len(page) < len(items)
    nxt = {'plan': plan['id'], 'range': span, 'offset': offset + len(page)} if more else None
    return Result({'cells': page, 'next': nxt}, truncated=more, next=nxt)


# ---------------------------------------------------------------------------
# 写入


class Workbook:
    """底稿快照的单元格和公式依赖。"""

    def __init__(self, path):
        self.book = Book(path.parent, path.name)
        self.formulas, self.users, self.unknown = {}, {}, []
        for name, sheet in self.book.sheets.items():
            for ref, cell in sheet['cells'].items():
                if cell.get('formula') is None:
                    continue
                key = (name, ref)
                if cell.get('sharedFormula') or not cell['formula']:
                    self.unknown.append(key)
                    continue
                self.formulas[key] = cell['formula']
                try:
                    for target in formula_refs(cell['formula'], name):
                        self.users.setdefault(target, []).append(key)
                except Exception:
                    self.unknown.append(key)

    def stored(self, sheet, ref):
        return self.book.sheets[sheet]['cells'].get(ref, {}).get('value') if sheet in self.book.sheets else None

    def affected(self, written):
        found, queue = set(), list(written)
        while queue:
            for user in self.users.get(queue.pop(), []):
                if user not in found:
                    found.add(user)
                    queue.append(user)
        return found


def evaluate(wb, written, affected):
    computed, pending, active = {}, {}, set()

    def value(sheet, ref):
        key = (sheet, ref)
        if key in written:
            return written[key]
        if key in affected:
            return compute(key)
        return wb.stored(sheet, ref)

    def compute(key):
        if key in computed:
            return computed[key]
        if key in pending:
            raise Pending(key)
        if key in active:
            pending[key] = '循环引用'
            raise Pending(key)
        active.add(key)
        try:
            result = calculate_formula(wb.formulas[key], key[0], value)
        except Pending as e:
            pending[key] = '依赖的 %s!%s 需由 Excel 重算' % e.args[0] if e.args[0] != key else pending.get(key, '循环引用')
            raise Pending(key)
        except (NeedsInput, ValueError, FormulaError, RecursionError, KeyError) as e:
            pending[key] = str(e) or type(e).__name__
            raise Pending(key)
        finally:
            active.discard(key)
        computed[key] = result
        return result

    for key in sorted(affected):
        try:
            compute(key)
        except Pending:
            pass
    return computed, pending


def cache_type(v):
    if isinstance(v, dict):
        return 'error', v['error']
    if isinstance(v, bool):
        return 'boolean', v
    if isinstance(v, Decimal):
        return 'number', number_text(v)
    return 'text', '' if v is None else str(v)


def strip_cache(raw, cell):
    pat = r'(<(?:\w+:)?c\b[^>]*\br=["\']%s["\'][^>]*>)(.*?)(</(?:\w+:)?c>)' % cell
    return re.sub(pat, lambda m: re.sub(r'\s+t=["\'][^"\']*["\']', '', m.group(1)) + re.sub(r'<(?:\w+:)?v\b[^>]*(?:/>|>.*?</(?:\w+:)?v>)', '', m.group(2), flags=re.S) + m.group(3), raw, flags=re.S)


def same(actual, expected):
    if isinstance(expected, Decimal):
        try:
            return actual is not None and not isinstance(actual, (dict, bool)) and Decimal(str(actual)) == expected
        except InvalidOperation:
            return False
    if isinstance(expected, dict):
        return isinstance(actual, dict) and actual.get('error') == expected.get('error')
    return actual == expected


def typed(write):
    if write['valueType'] == 'number':
        return Decimal(write['value'])
    return write['value']


def cross_foot(ctx, plan, wb, written, computed, pending):
    """合计行的 SUM 公式和明细对比；明细整列来自一张结果表时，和查询算出的合计对比。"""
    items = []
    by_cell = {(w['sheet'], w['ref']): w for w in plan['writes']}
    wp = ctx.store.get(plan['workpaper'], 'workpaper')
    for info in plan['regions']:
        r = wp['regions'][info['region']]
        sheet = r['sheet']
        for cell, formula in r['formulas'].items():
            row = coordinate_to_tuple(cell)[0]
            m = re.fullmatch(r'\s*SUM\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*', formula, re.I)
            if row not in r['totals'] or not m or m.group(1) != m.group(3):
                continue
            key = (sheet, cell)
            if key in pending:
                continue
            cells = [(sheet, '%s%d' % (m.group(1), n)) for n in range(int(m.group(2)), int(m.group(4)) + 1)]
            if not any(c in written for c in cells):
                continue
            actual = computed.get(key, wb.stored(sheet, cell))
            parts = [by_cell[c] for c in cells if c in by_cell]
            origins = {p['origin'].split('#')[0] for p in parts}
            columns = {p['from'].get('column') for p in parts}
            expected, how = None, '明细合计'
            if len(origins) == 1 and len(columns) == 1 and len(parts) == sum(1 for c in cells if wb.stored(*c) is not None or c in written):
                ref = origins.pop()
                if ref.startswith('R-'):
                    obj = ctx.store.get(ref, 'result')
                    column = columns.pop()
                    if len(parts) == obj['rows'] and column in obj['sums']:
                        scale = Decimal(parts[0]['scale'])
                        expected, how = Decimal(obj['sums'][column]) / scale, '%s 的 %s 合计' % (ref, column)
            if expected is None:
                total = Decimal(0)
                for c in cells:
                    v = written.get(c, computed.get(c, wb.stored(*c)))
                    if isinstance(v, Decimal):
                        total += v
                expected = total
            tolerance = Decimal('0.005') * len(parts) if any(Decimal(p['scale']) != 1 for p in parts) else Decimal(0)
            ok = isinstance(actual, Decimal) and abs(actual - expected) <= tolerance
            items.append({'cell': label(sheet, cell), 'passed': ok, 'formula': plain(actual) if isinstance(actual, Decimal) else str(actual), 'expected': plain(expected), 'by': how})
    return items


def write_report(path, plan, checks, details, recalc):
    import openpyxl
    from openpyxl.styles import Font
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = '写入明细'
    ws.append(['单元格', '写入值', '来源', '区域', '原有内容', '覆盖'])
    for w in plan['writes']:
        ws.append([w['cell'], w['value'], w['origin'], w['region'], w['old'], '是' if w['overwrite'] else ''])
    ws = book.create_sheet('核对')
    ws.append(['核对项', '结果', '说明'])
    for name in CHECKS:
        ws.append([CHECK_LABELS[name], {'passed': '通过', 'failed': '不通过', 'skipped': '未执行'}[checks[name]], '；'.join(details.get(name, []))[:30000]])
    ws = book.create_sheet('待人工填写')
    ws.append(['单元格'])
    for cell in plan['manual']:
        ws.append([cell])
    ws = book.create_sheet('待重算公式')
    ws.append(['单元格', '公式', '原因'])
    for item in recalc:
        ws.append([item['cell'], item['formula'], item['reason']])
    for sheet in book.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.column_dimensions['A'].width = 22
        for col in 'BCDE':
            sheet.column_dimensions[col].width = 26
    book.save(path)


def fill_apply(ctx, args):
    task = require(args, 'task', str)
    plan = ctx.store.get(require(args, 'plan'), 'plan')
    if plan['task'] != task:
        raise ToolError('ARG_TASK_MISMATCH', '写入计划属于任务 %s，不是 %s' % (plan['task'], task))
    card = args.get('approvedCard')
    if not isinstance(card, str) or not card.startswith('C-'):
        raise ToolError('NEEDS_INPUT', '写入前需要用户在写入确认卡上确认', hint='先发写入确认卡，带上卡片编号再调用')
    if plan['expand']:
        raise ToolError('NEEDS_INPUT', '明细表行数不够：%s' % '；'.join('%s 需要再加 %d 行' % (e['region'], e['rows']) for e in plan['expand']),
                        hint='在底稿里插入行后重新解析，或把超出的行列为人工填写')
    if plan['overwrites']:
        raise ToolError('NEEDS_INPUT', '有 %d 格会覆盖原有内容，尚未确认' % len(plan['overwrites']), hint='确认后用 overwrite 重新编译计划')
    _, snapshot, _ = ctx.files.resolve(plan['file'])
    try:
        wb = Workbook(snapshot)
    except (ValueError, NeedsInput, OSError, zipfile.BadZipFile) as e:
        raise ToolError('CHECK_WRITE_FAILED', '无法打开底稿快照：%s' % e)
    written = {(w['sheet'], w['ref']): typed(w) for w in plan['writes']}
    affected = wb.affected(written) - set(written)
    computed, pending = evaluate(wb, written, affected)
    book = wb.book
    changes = {}
    try:
        for w in plan['writes']:
            part = book.sheet(w['sheet'])['part']
            changes[part] = patch_cell(changes.get(part, book.sheet(w['sheet'])['raw']), w['ref'], w['value'], w['valueType'])
        for (sheet, cell), v in computed.items():
            part = book.sheet(sheet)['part']
            typ, text = cache_type(v)
            changes[part] = patch_cell(changes.get(part, book.sheet(sheet)['raw']), cell, text, typ, True)
        for sheet, cell in pending:
            part = book.sheet(sheet)['part']
            changes[part] = strip_cache(changes.get(part, book.sheet(sheet)['raw']), cell)
        if pending or (wb.unknown and written):
            with zipfile.ZipFile(book.path) as z:
                changes['xl/workbook.xml'] = set_full_calc(z.read('xl/workbook.xml').decode('utf-8'))
        target = unique(ctx.store.root, Path(plan['output']))
        target.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(book.path) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as destination:
            for item in source.infolist():
                destination.writestr(item, changes[item.filename].encode('utf-8') if item.filename in changes else source.read(item.filename))
        actual = Book(target.parent, target.name)
    except (ValueError, KeyError, OSError, zipfile.BadZipFile) as e:
        raise ToolError('CHECK_WRITE_FAILED', '写入底稿副本失败：%s' % e, hint='底稿结构可能有本工具不支持的部分，请把这份底稿发给维护人员')

    checks, details = {}, {}

    def result(name, problems, skipped=False):
        checks[name] = 'skipped' if skipped else 'failed' if problems else 'passed'
        if problems:
            details[name] = problems

    def now(sheet, cell):
        return actual.sheets[sheet]['cells'].get(cell, {}).get('value')

    result('readback', ['%s 写入 %s，读回 %s' % (label(*k), plain(v), plain(now(*k))) for k, v in written.items() if not same(now(*k), v)])

    problems = ['%s 缓存 %s，应为 %s' % (label(*k), now(*k), plain(v) if isinstance(v, Decimal) else v) for k, v in computed.items() if not same(now(*k), v)]
    problems += ['%s 的结果是 %s' % (label(*k), v['error']) for k, v in computed.items() if isinstance(v, dict)]
    result('formulas', problems, skipped=bool(pending) and not problems)
    if pending:
        details.setdefault('formulas', []).append('%d 个公式需由 Excel 打开时重算' % len(pending))

    feet = cross_foot(ctx, plan, wb, written, computed, pending)
    result('crossFoot', ['%s 合计 %s，%s %s' % (i['cell'], i['formula'], i['by'], i['expected']) for i in feet if not i['passed']], skipped=not feet)

    problems = []
    with zipfile.ZipFile(book.path) as source, zipfile.ZipFile(target) as out:
        if source.namelist() != out.namelist():
            problems.append('包部件清单发生变化')
        for part in source.namelist():
            if part not in changes and part in out.namelist() and source.read(part) != out.read(part):
                problems.append('包部件 %s 被改变' % part)
        if 'xl/workbook.xml' in changes:
            strip = lambda raw: re.sub(r'<(?:\w+:)?calcPr\b[^>]*>', '', raw)
            if strip(source.read('xl/workbook.xml').decode('utf-8')) != strip(out.read('xl/workbook.xml').decode('utf-8')):
                problems.append('工作簿设置被改变')
    touched = set(written) | set(computed) | set(pending)
    for name, sheet in book.sheets.items():
        if sheet['merges'] != actual.sheets[name]['merges']:
            problems.append('%s 的合并区域发生变化' % name)
        for cell, old in sheet['cells'].items():
            new = actual.sheets[name]['cells'].get(cell, {})
            if old.get('style') != new.get('style') or old.get('formula') != new.get('formula'):
                problems.append('%s 的样式或公式被改变' % label(name, cell))
            elif (name, cell) not in touched and old.get('value') != new.get('value'):
                problems.append('%s 不在计划内却被改变' % label(name, cell))
    result('untouched', problems[:50])

    wp = ctx.store.get(plan['workpaper'], 'workpaper')
    manual = set(plan['manual'])
    problems = []
    for info in plan['regions']:
        r = wp['regions'][info['region']]
        for cell in r['inputs']:
            v = now(r['sheet'], cell)
            if label(r['sheet'], cell) not in manual and (v is None or (isinstance(v, str) and not v.strip())):
                problems.append('%s 没有填写' % label(r['sheet'], cell))
    result('completeness', problems[:50])

    recalc = [{'cell': label(*k), 'formula': '=' + wb.formulas.get(k, ''), 'reason': reason} for k, reason in sorted(pending.items())]
    targets = sorted(set(computed) | set(pending))[:500]
    values = excel.recalc(target, targets) if args.get('excelRecalc', True) and targets else None
    if values is None:
        result('excel', [], skipped=True)
        details['excel'] = ['本机没有可用的 Excel，或没有需要重算的公式']
    else:
        problems = []
        for k in targets:
            got = values.get(label(*k))
            if k in pending:
                for item in recalc:
                    if item['cell'] == label(*k):
                        item['excel'] = got
                continue
            want = computed[k]
            if isinstance(want, dict):
                continue
            if isinstance(want, Decimal):
                try:
                    ok = got is not None and abs(Decimal(got) - want) <= Decimal('1e-9') * max(Decimal(1), abs(want))
                except InvalidOperation:
                    ok = False
            else:
                ok = (got or '') == ('' if want is None else str(want if not isinstance(want, bool) else str(want).upper()))
            if not ok:
                problems.append('%s 本地算得 %s，Excel 算得 %s' % (label(*k), plain(want), got))
        result('excel', problems)

    report = unique(ctx.store.root, Path('输出') / task / '填写核对.xlsx')
    write_report(report, plan, checks, details, recalc)
    root = ctx.store.root
    output, report_rel = target.relative_to(root).as_posix(), report.relative_to(root).as_posix()
    outputs = ctx.store.index('outputs')
    outputs[output] = plan['id']
    ctx.store.save_index('outputs', outputs)
    plan['applied'] = {'output': output, 'report': report_rel, 'card': card, 'checks': checks, 'details': details,
                       'at': datetime.datetime.now().isoformat(timespec='seconds')}
    ctx.store.put(plan['id'], plan)
    from .checks import record
    for name in CHECKS:
        if checks[name] != 'skipped':
            record(ctx, 'fill_' + name, None, plan['id'], checks[name] == 'passed', {'problems': details.get(name, [])[:10]})
    warnings = list(plan['warnings']) + ['%s 需在 Excel 打开时重算：%s' % (i['cell'], i['reason']) for i in recalc[:10]]
    return Result({'output': output, 'checks': checks, 'manualCells': plan['manual'], 'report': report_rel,
                   'recalc': [i['cell'] for i in recalc], 'problems': {k: v[:5] for k, v in details.items() if checks.get(k) == 'failed'}},
                  warnings=warnings)


# ---------------------------------------------------------------------------
# 追溯

BASIS_ITEMS = [('period', '期间', ('value',)), ('scope', '范围', ('include',)), ('amount', '金额与方向', ('column', 'direction')),
               ('grouping', '分组与对应', ('description',)), ('unit', '单位与精度', ('scale', 'rounding')), ('control', '控制数', ('label',))]
CHECK_NAMES = {'control': '控制数核对', 'compare': '双路复算', 'conservation': '守恒检查', 'basis_complete': '口径要素齐全'}


def basis_lines(basis):
    lines = []
    for key, name, fields in BASIS_ITEMS:
        item = basis.get(key)
        if not isinstance(item, dict):
            continue
        value = '，'.join(str(item[f]) for f in fields if item.get(f) not in (None, ''))
        if key == 'scope' and item.get('exclude'):
            value += '；排除 ' + '、'.join(str(x.get('condition')) for x in item['exclude'] if isinstance(x, dict))
        if value:
            lines.append({'label': name, 'value': value, 'basis': item.get('basis', 'inferred')})
    return lines


def trace_cell(ctx, args):
    output = require(args, 'output', str).replace('\\', '/')
    outputs = ctx.store.index('outputs')
    plan_id = outputs.get(output) or next((p for o, p in outputs.items() if output.endswith(o) or o.endswith('/' + output)), None)
    if not plan_id:
        raise ToolError('REF_NOT_FOUND', '%s 不是本项目填写生成的文件' % output)
    plan = ctx.store.get(plan_id, 'plan')
    sheet, cell = key_of(require(args, 'cell', str), None)
    if sheet is None:
        raise ToolError('ARG_INVALID', '单元格应写作 工作表!A1')
    where = label(sheet, cell)
    write = next((w for w in plan['writes'] if w['cell'] == where), None)
    if write is None:
        return Result({'cell': where, 'from': None, **({'manual': True} if where in plan['manual'] else {})})
    data = {'cell': where, 'value': write['value'], 'text': shown(write['value']) if write['valueType'] == 'number' else write['value'], 'from': write['from']}
    ref = write['origin'].split('#')[0]
    log = ctx.store.index('checks')
    entries = []
    if ref.startswith('R-'):
        obj = ctx.store.get(ref, 'result')
        tables = [i.split('@')[0] for i in obj['inputs']]
        if tables:
            table = ctx.store.get(tables[0], 'table')
            data['source'] = {'table': tables[0], 'file': table.get('file'), 'sheet': table.get('sheet'), 'rowCount': table.get('rows'),
                              **({'tables': tables} if len(tables) > 1 else {})}
        entries += log.get(ref, [])
    basis = ctx.store.index('basis').get(write['region'])
    data['basis'] = basis_lines(basis) if basis else []
    entries += [e for e in log.get(write['region'], []) if e not in entries]
    latest = {}
    for e in entries:
        latest[e['check']] = e
    data['checks'] = [{'check': e['check'], 'label': (e.get('detail') or {}).get('label') or CHECK_NAMES.get(e['check'], e['check']), 'passed': e['passed']}
                      for e in latest.values()]
    return Result(data)
