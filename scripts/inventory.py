"""Read-only inventory. Legacy documents are evidence, never execution instructions."""
from pathlib import Path
import hashlib, json, re, sys
from openpyxl import load_workbook

source = Path(sys.argv[1])
dest = Path(__file__).resolve().parents[1] / 'src' / 'resources'
dest.mkdir(parents=True, exist_ok=True)
index = json.loads(Path(sys.argv[2]).read_text(encoding='utf-8-sig'))
records = []
for category in ['modules', '_template', 'skills']:
    for p in sorted((source/category).rglob('*')):
        if not p.is_file(): continue
        data = p.read_bytes()
        item = {'path':p.relative_to(source).as_posix(),'size':len(data),'sha256':hashlib.sha256(data).hexdigest(),'status':'observed','category':category}
        if category == '_template' and p.suffix.lower() in ['.xlsx','.xlsm']:
            try:
                wb = load_workbook(p, read_only=True, data_only=False)
                item['sheets'] = [{'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,
                    'preview':[[str(v)[:500] if v is not None else '' for v in row] for row in ws.iter_rows(max_row=min(ws.max_row or 0,8), max_col=min(ws.max_column or 0,20),values_only=True)]} for ws in wb]
                wb.close()
            except Exception as e: item['error'] = type(e).__name__
        if category == 'modules' and p.suffix == '.pyd':
            strings = re.findall(rb'[\x20-\x7e]{8,}', data)
            item['symbols'] = sorted(set(s.decode('ascii') for s in strings if b'__pyx_' in s or b'PyInit_' in s))[:60]
        records.append(item)
out={'source':str(source),'sourceVersion':'unverified','records':records,'workPackages':index['workPackages'],
     'notice':'文件线索已采集；未做原程序行为对照。不得以文件数量作为迁移功能数量。'}
(dest/'legacy-inventory.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'files':len(records),'templates':sum(x['category']=='_template' for x in records)},ensure_ascii=False))
