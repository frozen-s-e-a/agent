"""Derive delivery documentation from current contracts and passing evidence."""
from pathlib import Path
import hashlib,json,sys
from openpyxl import load_workbook
from runtime.registry import TASKS

ROOT=Path(__file__).resolve().parents[1]
def read(p):return json.loads((ROOT/p).read_text(encoding='utf8'))
def digest(p):return hashlib.sha256((ROOT/p).read_bytes()).hexdigest()
def cell(value):return str(value).replace('|','／').replace('\n',' ').replace('\r',' ')

business_path,pipeline_path,public_path=sys.argv[1:4]
business,pipeline,public=read(business_path),read(pipeline_path),read(public_path)
assert business['cases']==business['passedCases'] and business['tasksCovered']==63
assert pipeline['passed']==pipeline['total']
assert public['passed']==public['total']==5
ui_path='artifacts/workbench-preview-20261008/verification.json'
ui=read(ui_path);assert not ui['errors'] and ui['checks'].get('all79Searchable')
model_path='artifacts/independent-live-validation/模型实测记录.json'
model_records=[]
for identifier,kind in [('53fb1de8-7200-461e-9744-b5ca24defecd','text'),('7a05aae9-6b42-42eb-9da1-cb7ca4f34580','vision')]:
    p='artifacts/independent-live-validation/project/outputs/'+identifier+'/阶段结果.json'
    result=read(p);value=result.get('structuredContent') or json.loads(result['content'][0]['text'])
    assert value['success'] is True
    model_records.append({'kind':kind,'taskId':identifier,'passed':True,'syntheticInputOnly':True,'resultSha256':digest(p)})
(ROOT/model_path).write_text(json.dumps({'passed':2,'total':2,'records':model_records},ensure_ascii=False,indent=2),encoding='utf8')
evidence={name:{'path':p,'sha256':digest(p)} for name,p in [('本地业务验收',business_path),('客户端入口验收',pipeline_path),('公开服务验收',public_path),('界面操作验收',ui_path),('模型实测',model_path)]}
verification={'date':'2026-10-10','localBusiness':{k:business[k] for k in ['cases','passedCases','tasksCovered']},'clientPipeline':{k:pipeline[k] for k in ['passed','total']},'publicServices':{'passed':5,'total':5,'tasks':[r['task'] for r in public['records']]},'liveModels':{'passed':2,'total':2,'syntheticInputOnly':True},'cicpa':{'status':'skipped-by-user','protocolFixturesPassed':True,'liveVerified':False},'ui':ui['checks'],'allRemoteServicesVerified':False,'evidence':evidence}
(ROOT/'src/resources/verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
(ROOT/'src/resources/mcp-migration.json').write_text(json.dumps({'runtime':'independent-python','taskCount':63,'originalExeRequired':False,'originalBusinessPydRequired':False,'implemented':sorted(TASKS),'verification':verification},ensure_ascii=False,indent=2),encoding='utf8')

catalog=read('build/task-catalog.json');assert len(catalog)==79
observations={name:[r for r in business['results'] if r['task']==name] for name in TASKS}
lines=['# 任务与文件格式对照','', '版本 0.1.0-internal.29，2026-10-10。63 项审计任务及 16 项通用操作。以下内容直接从客户端当前参数合同、独立处理器注册表、模板和验收成果生成。', '', '每项任务在页面填写的参数都由后台校验。模型、注协登录及成果目录由客户端统一管理，无需填写任务密钥或输出路径。表格处理另存成果；明细配置保存写入明确指定的 config 目录。', '', '## 63 项审计任务', '', '|操作与任务 ID|输入格式|配置模板|阶段成果格式|独立实现|', '|---|---|---|---|---|']
for t in catalog:
    if t['engine']!='mcp':continue
    name=t['mcpName'];outputs=sorted({Path(x['name']).suffix.upper() for r in observations[name] for x in r['outputs'] if x['name']!='manifest.json'})
    lines.append('|'+ '|'.join(map(cell,[t['name']+' / '+name,t['contract']['format'],'；'.join(t['contract'].get('templates',[])) or '页面参数或项目内已保存配置',' / '.join(outputs) or 'JSON 结果',TASKS[name].__module__.removeprefix('runtime.')+'.py']))+'|')
lines+=['','## 逐项参数与成果示例','', '必填状态以客户端表单为准；可留空、已保存路径及默认值会在页面说明。成果示例来自合成资料验收，不包含客户数据。','']
for t in catalog:
    if t['engine']!='mcp':continue
    name=t['mcpName'];lines+=['### '+t['name']+'（'+name+'）','',t['contract']['format'],'', '|参数|含义|填写要求|','|---|---|---|']
    for f in t['fields']:
        rule='必填' if f.get('required') and not f.get('allowBlank') else '可选／可留空'
        if f.get('default') is not None:rule+='；默认 '+str(f['default'])
        if f.get('enum'):rule+='；取值 '+ '、'.join(map(str,f['enum']))
        lines.append('|'+ '|'.join(map(cell,[f['key'],f.get('help') or f['label'],rule]))+'|')
    samples=list(dict.fromkeys(x['name'] for r in observations[name] for x in r['outputs'] if x['name']!='manifest.json'))[:6]
    lines+=['','验收成果示例：'+'、'.join(samples)+'。','']
    if t['contract'].get('dependency'):lines+=[t['contract']['dependency']+'。','']
lines+=['## 配置模板与工作表','', '复制模板后填写业务参数，保留工作表名称及列标题。来源路径限当前项目。模板内使用说明页不属于待处理数据。','', '|配置／格式文件|业务工作表|','|---|---|']
template_names=sorted({n for t in catalog for n in t['contract'].get('templates',[])})
for name in template_names:
    wb=load_workbook(ROOT/'src/resources/templates'/name,read_only=True,keep_vba=name.endswith('.xlsm'))
    lines.append('|'+cell(name)+'|'+cell('、'.join(n for n in wb.sheetnames if n!='使用说明'))+'|');wb.close()
lines+=['','## 16 项通用资料操作','','|操作|输入格式|','|---|---|']
for t in catalog:
    if t['engine']=='builtin':lines.append('|'+cell(t['name']+' / '+t['id'])+'|'+cell(t['contract']['format'])+'|')
(ROOT/'docs/任务与文件格式对照-20261010.md').write_text('\n'.join(lines)+'\n',encoding='utf8')

report=f'''# 独立客户端验收与使用

版本：0.1.0-internal.29。日期：2026-10-10。

## 已交付的运行链路

63 项任务全部注册到独立 Python 处理器，通过客户端后台、项目资料范围校验、工作进程和成果记录执行。另有 16 项通用资料操作。程序包包含 Electron、Node、Python、依赖和 {len(template_names)} 份配置／格式文件，不调用原工具箱 EXE 或原业务 PYD，不要求安装原工具箱。

界面使用资料工作台、搜索／推荐操作、工作记录、成果文件、可选流程模板及运行设置。任务按资料与工作上下文衔接，细碎工具不强制分成独立业务模块。每项任务的输入、参数、成果和源码对应见《任务与文件格式对照》。

## 验收结果与证据

|验证层次|结果|验证范围|
|---|---|---|
|独立业务处理器|{business['passedCases']}/{business['cases']}，覆盖 63 项|合成资料，断言金额、匹配、字段、账龄、公式、输出哈希及输入未被覆盖；模型及部分 HTTP 使用声明的夹具|
|客户端实际任务入口|{pipeline['passed']}/{pipeline['total']}|后台 → 参数合同 → 项目范围 → Python → 成果文件 → 持久化；模型接口使用本地夹具|
|真实公开网络服务|5/5|沪、深、北三所 IPO 披露 PDF、公告 PDF、汇率数据；通过实际客户端入口执行|
|真实文本／视觉模型|2/2|用户已保存的模型配置，仅发送合成文字／图片，核对提取内容|
|界面操作|通过|79 操作可搜索、IPO 参数、收藏、预览、真实通用操作、成果衔接、工作恢复、窄窗口、运行设置|
|注协在线授权与查询|按用户要求跳过|工具箱没有可复用授权，不要求用户继续登录；搜索、详情、子公司及导出使用合成协议响应测试|

本次额外核对银行对手拆分、比例分摊、公司／账户隔离、多对多汇总和跨月匹配；明细底稿的标准配置包装、多期账龄、分组、文本编号、查找映射、模板扩行、合计公式与按公司空底稿。毛利率参数审阅与确认继续经过客户端入口复测。

验收资料和结果摘要位于随包“验收记录”目录。源工作区中的完整案例和成果路径记录在 src/resources/verification.json；文件哈希确保可核对。规则恢复研究另保留 2237 个基础函数案例、12 个表格流程、28 个 IPO 模拟流程、195 个关联维度案例及 JET 行为对比记录；这些研究记录不作为本次真实网络服务已验证的证明。

## 使用与可复核范围

解压完整程序目录后双击 AI助手.exe；无需 pnpm 或开发环境。程序默认使用本机现有 AuditAssistant 配置，因此已保存模型可以继续使用。注协授权仅在需要在线工商任务时登录。

验收覆盖上述具体场景，不等于穷尽原程序所有参数组合与真实项目数据。注协实际服务未验证；AI 提取和审计线索仍需项目复核。关联方已有资料分析与在线工商授权分开，缺失工商信息会展示覆盖提示。模板中的公式在 Excel 打开时重算；按数据重建的部分操作不复制 VBA 或复杂原版式。

本包定位为可独立运行的内部测试版，具体构建版本、运行库、来源提交、文件哈希和验收范围见 build-manifest.json、package-files.json 及 requirements.txt。
'''
(ROOT/'docs/独立客户端验收与使用-20261010.md').write_text(report,encoding='utf8')
deps=read('build/python-packages/independent-dependencies.json')
notices=['# 第三方组件说明','','程序内保留随运行库及依赖分发的许可证和来源元数据。Python 包的确切许可证文件路径及 SHA-256 位于 resources/app/runtime/python-packages/independent-dependencies.json。','', '|组件|固定版本|许可证／说明|','|---|---|---|','|Electron|38.0.0|MIT；附带 Chromium 声明见程序目录 LICENSES.chromium.html|','|Node.js|24.19.0|MIT 及内置依赖，runtime/NODE_LICENSE|','|Python|3.12.14|PSF；runtime/python/LICENSE.txt|','|React / React DOM|19.1.1|MIT|','|Lucide React|0.468.0|ISC|','|react-markdown / remark-gfm|10.1.0 / 4.0.1|MIT；依赖许可证位于 runtime/ui-licenses|']
for dep in deps['distributions']:
    license=str(dep.get('license') or '见随包许可证文件及来源元数据').splitlines()[0][:140]
    notices.append('|'+cell(dep['name'])+'|'+cell(dep['version'])+'|'+cell(license)+'|')
notices+=['','PyMuPDF 的 COPYING 与来源链接随发行包保留，使用其 AGPL／商业许可条件。构建工具为 Vite 7.1.3、TypeScript 5.9.2、Playwright 1.55.0；Node 直接及间接依赖由 pnpm-lock.yaml 固定。','','旧工具箱仅作规则与公开空白模板结构的只读研究，程序包不包含旧应用 EXE、业务 PYD、用户凭据或真实底稿。公开格式模板及 JET 规则来源保留在源码资源中。','','模型与 HTTP 测试区分夹具和真实接口：业务回归使用声明夹具，另完成真实公开服务与合成内容的真实模型实测。']
(ROOT/'docs/THIRD_PARTY_NOTICES.md').write_text('\n'.join(notices)+'\n',encoding='utf8')
print(json.dumps({'tasks':len(TASKS),'templates':len(template_names),'evidence':len(evidence),'verification':'src/resources/verification.json'},ensure_ascii=False))
