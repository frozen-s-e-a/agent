"""Business-level checks for ordinary, unannotated templates and saved outputs."""
import json, sys, tempfile, unittest, zipfile, copy
from pathlib import Path
from decimal import Decimal
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'src/audit-worker'))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
from workbook_plan import validate_plan, fill_workbook, read_workbook, NeedsInput, sha

def fixture(folder):
    wb=Workbook();wb.remove(wb.active)
    for name in ['Lead Schedule','Disclosure','BKD24']:
        s=wb.create_sheet(name);s['A1']='投资收益';s.merge_cells('A1:F1');s['B2']='项目';s['C2']=2025;s['D2']=2024
        s['B3']='经核对明细';s['E3']='=C3-D3';s['F3']='=E3/D3';s['C3'].fill=PatternFill('solid',fgColor='FFEEDD');s['C3'].number_format='#,##0.00';s['H8']='模板尾部'
    s=wb.create_sheet('Sheet1');s['A1']='来源明细';s.append(['科目','部门','2025','2024'])
    for n in range(1,121):s.append(['投资收益','本部',n,5])
    file=folder/'普通模板.xlsx';wb.save(file)
    with zipfile.ZipFile(file,'a') as z:z.writestr('customXml/template-marker.xml','<original>preserve unrelated objects</original>')
    header={'file_path':file.name,'sheet':'Sheet1','row':2,'columns':{'A':'科目','B':'部门','C':'2025','D':'2024'}}
    writes=[];required=[];checks=[]
    for name in ['Lead Schedule','Disclosure','BKD24']:
        required.append({'sheet':name,'range':'C3:F3'})
        for target,column,year in [('C3','C',2025),('D3','D',2024)]:
            writes.append({'sheet':name,'cell':target,'reason':f'按实际 {year} 表头汇总全部投资收益明细','expr':{'op':'sum','file_path':file.name,'sheet':'Sheet1','headerRow':2,'column':column,'endRow':'auto','filters':[{'column':'A','value':'投资收益'}]}})
            checks.append({'source':{'file_path':file.name,'sheet':'Sheet1','cell':column+'2'},'target':{'file_path':file.name,'sheet':name,'cell':target[0]+'2'}})
    return file,{'file_path':file.name,'targetSheets':['Lead Schedule','Disclosure','BKD24'],'headers':[header],'required':required,'periodChecks':checks,'writes':writes}

ITEMS=[('611101','成本法核算的长期股权投资收益'),('611102','权益法核算的长期股权投资收益'),('611103','处置长期股权投资产生的投资收益'),('611105','交易性金融资产持有期间取得的投资收益'),('611106','处置交易性金融资产取得的投资收益')]
def excel_saved(file):
    """Rewrite formulas the way Excel saves them: cached results, and D/F filled down as shared formulas."""
    import re,shutil
    tmp=file.with_suffix('.tmp');shutil.move(file,tmp)
    def cache(m):
        cell=m.group(0).replace('<v />','').replace('<v/>','');col=re.search(r'r="([A-Z]+)\d+"',cell).group(1)
        if col=='G':return cell.replace('<c ','<c t="e" ',1).replace('</f>','</f><v>#DIV/0!</v>')
        return cell.replace('</f>','</f><v>0</v>')
    with zipfile.ZipFile(tmp) as src,zipfile.ZipFile(file,'w',zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet'):
                xml=re.sub(r'<c [^>]*>(?:(?!</c>).)*<f>(?:(?!</c>).)*</c>',cache,data.decode())
                if item.filename=='xl/worksheets/sheet1.xml':
                    for si,col in enumerate('DF'):
                        for row in range(4,9):
                            m=re.search(rf'<c r="{col}{row}"[^>]*><f>([^<]*)</f>',xml)
                            f=f'<f t="shared" ref="{col}4:{col}8" si="{si}">{m.group(1)}</f>' if row==4 else f'<f t="shared" si="{si}"/>'
                            xml=xml[:m.start()]+m.group(0).replace(f'<f>{m.group(1)}</f>',f)+xml[m.end():]
                data=xml.encode()
            dst.writestr(item,data)
    tmp.unlink()

def ledger_fixture(folder):
    """Shaped like the reported workbook: Sheet1 is a trial balance with no year, blank zeros and subtotal/auxiliary rows."""
    wb=Workbook();wb.remove(wb.active)
    ls=wb.create_sheet('Lead Schedule');ls['A1']='投资收益 Lead Schedule';ls.merge_cells('A1:G1');ls['E2']='单位：元'
    ls.append(['项目','2025年未审数','审计调整','2025年审定数','2024年审定数','变动额','变动率'])
    for i,(c,n) in enumerate(ITEMS,start=4):ls.append([n,None,None,f'=B{i}+C{i}',None,f'=D{i}-E{i}',f'=F{i}/E{i}'])
    ls.append(['合计','=SUM(B4:B8)','=SUM(C4:C8)','=SUM(D4:D8)','=SUM(E4:E8)','=D9-E9','=F9/E9'])
    d=wb.create_sheet('Disclosure');d['A1']='投资收益';d.append(['项目','本期发生额','上期发生额'])
    for c,n in ITEMS:d.append([n,None,None])
    d.append(['合计','=SUM(B3:B7)','=SUM(C3:C7)'])
    b=wb.create_sheet('BKD24');b['A1']='投资收益明细表';b.append(['科目编码','科目','科目名称','部门','类别','项目辅助核算名称','项目','2023','2024','2025'])
    for row in [('611101','成本法核算的长期股权投资收益','被投资单位A',500000,800000),('611102','权益法核算的长期股权投资收益','联营企业B',1600000,1900000),('611102','权益法核算的长期股权投资收益','合营企业D',500000,700000),('611103','处置长期股权投资产生的投资收益','被投资单位C',None,None),('611105','交易性金融资产持有期间取得的投资收益',None,40000,48000),('611106','处置交易性金融资产取得的投资收益',None,-5000,9000)]:
        b.append([row[0],'投资收益',row[1],'本部','损益',row[2],None,row[3],row[4],None])
    s=wb.create_sheet('Sheet1');s['A1']='科目余额表';s['A2']='编制单位：示例股份有限公司    金额单位：元'
    s.append(['科目编码','科目名称','辅助核算','期初借方','期初贷方','本期借方','本期贷方','期末借方','期末贷方'])
    for n in range(40):s.append([str(1001+n),'资产科目'+str(n),None,1000*n or None,None,100*n or None,None,None,None])
    for r in [('6111','投资收益',None,162000,4706000),('611101','投资收益-成本法核算的长期股权投资收益',None,None,1200000),('611101','投资收益-成本法核算的长期股权投资收益','被投资单位A',None,1200000),('611102','投资收益-权益法核算的长期股权投资收益',None,150000,3450000),('611102','投资收益-权益法核算的长期股权投资收益','联营企业B',150000,2650000),('611102','投资收益-权益法核算的长期股权投资收益','合营企业D',None,800000),('611103','投资收益-处置长期股权投资产生的投资收益',None,None,None),('611105','投资收益-交易性金融资产持有期间取得的投资收益',None,None,56000),('611106','投资收益-处置交易性金融资产取得的投资收益',None,12000,None)]:
        s.append([r[0],r[1],r[2],None,None,r[3],r[4],None,None])
    file=folder/'投资收益底稿.xlsx';wb.save(file);excel_saved(file)
    F=file.name
    def net(code,aux=None):
        match=lambda value:{'column':'C','op':'equals','value':''} if value is None else value
        filters=lambda: [{'column':'A','op':'equals','value':code},match(aux) if isinstance(aux,dict) else {'column':'C','op':'equals','value':''}]
        side=lambda col:{'op':'sum','file_path':F,'sheet':'Sheet1','headerRow':3,'column':col,'endRow':'auto','filters':filters()}
        return {'op':'subtract','args':[side('本期贷方'),side('F')]}
    writes=[]
    for i,(code,name) in enumerate(ITEMS):
        writes.append({'sheet':'Lead Schedule','cell':f'B{4+i}','expr':net(code),'reason':f'Sheet1 {code} 汇总行 本期贷方-本期借方'})
        writes.append({'sheet':'Lead Schedule','cell':f'E{4+i}','expr':{'op':'sum','file_path':F,'sheet':'BKD24','headerRow':2,'column':'I','endRow':'auto','filters':[{'column':'A','value':code}]},'reason':'BKD24 2024 年合计'})
        writes.append({'sheet':'Disclosure','cell':f'B{3+i}','expr':{'op':'cell','file_path':F,'sheet':'Lead Schedule','cell':f'D{4+i}'},'reason':'取 Lead Schedule 本年审定数'})
        writes.append({'sheet':'Disclosure','cell':f'C{3+i}','expr':{'op':'cell','file_path':F,'sheet':'Lead Schedule','cell':f'E{4+i}'},'reason':'取 Lead Schedule 上年审定数'})
    key=lambda col:{'op':'sum','file_path':F,'sheet':'Sheet1','headerRow':3,'column':col,'endRow':'auto','allowEmpty':True,'filters':[{'column':'A','valueExpr':{'op':'cell','file_path':F,'sheet':'BKD24','cell':'A{row}'}},{'column':'C','valueExpr':{'op':'cell','file_path':F,'sheet':'BKD24','cell':'F{row}'}}]}
    writes.append({'sheet':'BKD24','range':'J3:J8','expr':{'op':'subtract','args':[key('G'),key('F')]},'reason':'按科目编码和被投资单位匹配 Sheet1，本期贷方-本期借方'})
    year={'file_path':F,'sheet':'Sheet1','confirmedYear':'2025'}
    plan={'file_path':F,'targetSheets':['Lead Schedule','Disclosure','BKD24'],
          'required':[{'sheet':'Lead Schedule','range':'B4:B8'},{'sheet':'Lead Schedule','range':'D4:G8'},{'sheet':'Disclosure','range':'B3:C7'},{'sheet':'BKD24','range':'J3:J8'}],
          'headers':[{'file_path':F,'sheet':'Sheet1','row':3,'columns':{'A':'科目编码','C':'辅助核算','F':'本期借方','G':'本期贷方'}},{'file_path':F,'sheet':'BKD24','row':2,'columns':{'A':'科目编码','F':'项目辅助核算名称','I':'2024'}}],
          'periodChecks':[{'source':year,'target':{'file_path':F,'sheet':'Lead Schedule','cell':'B3'}},{'source':{'file_path':F,'sheet':'BKD24','cell':'I2'},'target':{'file_path':F,'sheet':'Lead Schedule','cell':'E3'}},{'source':year,'target':{'file_path':F,'sheet':'BKD24','cell':'J2'}}],
          'writes':writes}
    return file,plan

class WorkbookPlanTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name).resolve();self.file,self.plan=fixture(self.root);self.files=[self.file];self.original=sha(self.file)
    def tearDown(self):self.tmp.cleanup()
    def validate(self,plan=None):return validate_plan(self.root,self.files,{'plan':plan or self.plan})
    def test_full_population_and_actual_reread_preserve_template(self):
        prepared=self.validate();self.assertTrue(prepared['ok']);self.assertEqual(prepared['writeCount'],6)
        stage=self.root/'stage';stage.mkdir();result=fill_workbook(self.root,self.files,{'plan':self.plan,'expectedPlanHash':prepared['planHash']},stage)
        self.assertEqual(sha(self.file),self.original);actual=load_workbook(stage/'填写成果.xlsx',data_only=True)
        for name in self.plan['targetSheets']:
            self.assertEqual(actual[name]['C3'].value,7260);self.assertEqual(actual[name]['D3'].value,600)
            self.assertEqual(actual[name]['E3'].value,6660);self.assertEqual(actual[name]['F3'].value,11.1)
            self.assertEqual(actual[name]['C3'].number_format,'#,##0.00')
        actual.close();receipt=json.loads((stage/'填写核对.json').read_text(encoding='utf-8'))
        self.assertEqual(receipt['verifiedCount'],12);self.assertEqual(receipt['formulaCount'],6)
        with zipfile.ZipFile(self.file) as original,zipfile.ZipFile(stage/'填写成果.xlsx') as output:
            for part in original.namelist():
                if not part.startswith('xl/worksheets/sheet') or part=='xl/worksheets/sheet4.xml':self.assertEqual(original.read(part),output.read(part),part)
    def test_range_beyond_preview_and_pagination(self):
        result=read_workbook(self.root,self.files,{'sheet':'Sheet1','range':'A100:D122','offset':80,'limit':12})
        self.assertFalse(result['sampleOnly']);self.assertEqual(result['pages'][0]['nextOffset'],None)
        self.assertEqual(result['pages'][0]['rows'][-1]['row'],122);self.assertEqual(result['pages'][0]['rows'][-1]['C'],'120')
        self.assertEqual(result['pages'][0]['headerCandidates'][0]['row'],2)
    def test_blank_amount_counts_as_zero_like_excel_and_is_reported(self):
        w=load_workbook(self.file);w['Sheet1']['C100']=None;w.save(self.file);w.close()
        result=self.validate();self.assertTrue(result['ok'],result);self.assertEqual(result['preview'][0]['expected'],'7162');self.assertIn('空白金额',result['warnings'][0])
    def test_text_in_amount_column_goes_back_to_the_model(self):
        w=load_workbook(self.file);w['Sheet1']['C100']='待补';w.save(self.file);w.close()
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('不是金额',result['issue']);self.assertIn('startRow',result['hint'])
    def test_period_conflict_goes_back_to_the_model(self):
        self.plan['periodChecks'][0]['target']['cell']='D2';result=self.validate();self.assertFalse(result['ok']);self.assertEqual(result['status'],'plan_revision');self.assertIn('期间',result['issue'])
    def test_duplicate_lookup_cannot_pick_first(self):
        self.plan['writes'][0]['expr']['op']='lookup';result=self.validate();self.assertFalse(result['ok']);self.assertEqual(result['status'],'plan_revision');self.assertIn('120 行',result['issue']);self.assertIn('sum',result['hint'])
    def test_missing_required_mapping_cannot_pass(self):
        self.plan['writes'].pop();result=self.validate();self.assertFalse(result['ok']);self.assertEqual(result['status'],'plan_revision');self.assertIn('缺少映射',result['issue']);self.assertTrue(result['hint'])
    def test_numeric_data_row_cannot_be_header(self):
        self.plan['headers'][0]={'file_path':self.file.name,'sheet':'Sheet1','row':3,'columns':{'C':'1'}}
        result=self.validate();self.assertFalse(result['ok']);self.assertEqual(result['status'],'plan_revision');self.assertIn('不能当作表头',result['issue']);self.assertTrue(result['hint'])
    def test_duplicate_and_merged_targets_rejected(self):
        bad=copy.deepcopy(self.plan);bad['writes'].append(copy.deepcopy(bad['writes'][0]));self.assertRaisesRegex(ValueError,'重复写入',self.validate,bad)
        bad=copy.deepcopy(self.plan);bad['required'].append({'sheet':'Lead Schedule','range':'C1'});bad['writes'][0]['cell']='C1';self.assertRaisesRegex(ValueError,'合并区域内部',self.validate,bad)
    def test_no_arbitrary_code_or_business_literals(self):
        self.plan['writes'][0]['expr']={'op':'literal','value':999};self.assertRaisesRegex(ValueError,'猜测',self.validate)
    def test_changed_input_invalidates_saved_plan(self):
        prepared=self.validate();w=load_workbook(self.file);w['Sheet1']['C122']=121;w.save(self.file);w.close();stage=self.root/'stage';stage.mkdir()
        self.assertRaisesRegex(NeedsInput,'版本已变化',fill_workbook,self.root,self.files,{'plan':self.plan,'expectedPlanHash':prepared['planHash']},stage)
        self.assertFalse((stage/'填写成果.xlsx').exists())
    def test_template_formula_reading_planned_writes_is_recomputed_not_its_cache(self):
        # Disclosure!C3 links to Lead Schedule!E3 (=C3-D3), whose inputs this plan writes.
        link=next(w for w in self.plan['writes'] if w['sheet']=='Disclosure' and w['cell']=='C3')
        link['expr']={'op':'cell','file_path':self.file.name,'sheet':'Lead Schedule','cell':'E3'}
        prepared=self.validate();self.assertTrue(prepared['ok'],prepared)
        self.assertEqual(next(item['expected'] for item in prepared['preview'] if item['sheet']=='Disclosure' and item['cell']=='C3'),'6660')
        stage=self.root/'stage';stage.mkdir();fill_workbook(self.root,self.files,{'plan':self.plan,'expectedPlanHash':prepared['planHash']},stage)
        actual=load_workbook(stage/'填写成果.xlsx',data_only=True);self.assertEqual(actual['Disclosure']['C3'].value,6660);self.assertEqual(actual['Disclosure']['E3'].value,6060);actual.close()
    def test_a_write_that_reads_itself_is_rejected(self):
        self.plan['writes'][0]['expr']={'op':'cell','file_path':self.file.name,'sheet':'Lead Schedule','cell':'E3'}
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('自身',result['issue'])
    def test_grouped_fill_uses_each_template_row_key_without_large_model_output(self):
        w=load_workbook(self.file)
        for n in range(1,121):w['Sheet1'].cell(n+2,1,'组'+str(n%3))
        for row,name in [(3,'组0'),(4,'组1'),(5,'组2')]:w['Lead Schedule'].cell(row,2,name)
        w.save(self.file);w.close()
        for item in self.plan['writes']:item['expr']['filters']=[{'column':'B','value':'本部'}]
        first=self.plan['writes'][0];first.pop('cell');first['range']='C3:C5'
        first['expr']['filters']=[{'column':'A','valueExpr':{'op':'cell','file_path':self.file.name,'sheet':'Lead Schedule','cell':'B{row}'}}]
        self.plan['required'].append({'sheet':'Lead Schedule','range':'C4:C5'})
        prepared=self.validate();self.assertTrue(prepared['ok']);stage=self.root/'stage';stage.mkdir()
        fill_workbook(self.root,self.files,{'plan':self.plan,'expectedPlanHash':prepared['planHash']},stage)
        actual=load_workbook(stage/'填写成果.xlsx',data_only=True)
        self.assertEqual([actual['Lead Schedule'][f'C{r}'].value for r in range(3,6)],[2460,2380,2420]);actual.close()


class LedgerTemplateTests(unittest.TestCase):
    """The reported task: fill Lead Schedule, Disclosure and BKD24 from a trial balance in Sheet1."""
    ANSWER='补充说明：\n1. Sheet1 是哪一年的数据？\n回答：2025年'
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name).resolve();self.file,self.plan=ledger_fixture(self.root);self.files=[self.file]
        self.plan['_userStatements']=['用sheet1表格填写前三个sheet表',self.ANSWER]
    def tearDown(self):self.tmp.cleanup()
    def validate(self,plan=None):return validate_plan(self.root,self.files,{'plan':plan or self.plan})
    def fill(self):
        prepared=self.validate();self.assertTrue(prepared['ok'],prepared);stage=self.root/'stage';stage.mkdir()
        result=fill_workbook(self.root,self.files,{'plan':self.plan,'expectedPlanHash':prepared['planHash']},stage)
        return result,stage
    def test_year_without_a_label_needs_the_users_word(self):
        self.plan['_userStatements']=['用sheet1表格填写前三个sheet表']
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('confirmedYear',result['issue']);self.assertIn('ask_user_questions',result['hint'])
        self.plan['periodChecks'][0]['source']={'file_path':self.file.name,'sheet':'Sheet1','cell':'A1'}
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('无法唯一确定年份',result['issue']);self.assertIn('confirmedYear',result['hint'])
    def test_answered_year_fills_all_three_sheets_with_excel_results(self):
        result,stage=self.fill();receipt=result['workbookValidation']
        self.assertEqual((receipt['requiredCount'],receipt['verifiedCount'],receipt['recalcCount']),(41,41,0))
        book=load_workbook(stage/'填写成果.xlsx',data_only=True);lead,disclosure,bkd=book['Lead Schedule'],book['Disclosure'],book['BKD24']
        # Net credit of each subtotal row; blank debits count as 0.
        self.assertEqual([lead[f'B{r}'].value for r in range(4,9)],[1200000,3300000,0,56000,-12000])
        self.assertEqual([lead[f'E{r}'].value for r in range(4,9)],[800000,2600000,0,48000,9000])
        # Shared formulas D5:D8 / F5:F8 are recomputed, not left at the cached 0.
        self.assertEqual([lead[f'D{r}'].value for r in range(4,9)],[1200000,3300000,0,56000,-12000]);self.assertEqual(lead['F5'].value,700000)
        self.assertEqual(lead['G6'].value,'#DIV/0!');self.assertEqual(lead['D9'].value,4544000)
        # Disclosure links to the Lead Schedule results computed in the same fill.
        self.assertEqual([disclosure[f'B{r}'].value for r in range(3,8)],[1200000,3300000,0,56000,-12000]);self.assertEqual(disclosure['C8'].value,3457000)
        self.assertEqual([bkd[f'J{r}'].value for r in range(3,9)],[1200000,2500000,800000,0,56000,-12000]);book.close()
        self.assertTrue(any('#DIV/0!' in item for item in result['warnings']));self.assertTrue(any('BKD24!J6' in item for item in result['warnings']))
        formulas=load_workbook(stage/'填写成果.xlsx')['Lead Schedule'];self.assertEqual(formulas['D5'].value,'=B5+C5')
    def test_no_matching_source_row_needs_allow_empty(self):
        for arg in self.plan['writes'][-1]['expr']['args']:arg.pop('allowEmpty')
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('BKD24!J6',result['issue']);self.assertIn('被投资单位C',result['issue']);self.assertIn('allowEmpty',result['hint'])
    def test_unknown_column_lists_declared_headers(self):
        self.plan['writes'][0]['expr']['args'][0]['column']='期末贷方'
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('本期贷方',result['hint'])
    def test_dash_placeholder_in_a_target_is_filled(self):
        w=load_workbook(self.file);w['BKD24']['J3']='-';w.save(self.file);w.close()
        self.assertTrue(self.validate()['ok'])
        w=load_workbook(self.file);w['BKD24']['J3']='旧数';w.save(self.file);w.close()
        result=self.validate();self.assertEqual(result['status'],'plan_revision');self.assertIn('allowOverwrite',result['hint'])
    def test_formula_this_engine_cannot_evaluate_is_left_for_excel(self):
        w=load_workbook(self.file);w['Lead Schedule']['G4']='=SUMPRODUCT(F4:F4,1/E4:E4)';w.save(self.file);w.close()
        # Excel does not save fullCalcOnLoad; openpyxl does, so take it out to see the fill add it.
        packed=self.file.with_suffix('.tmp');self.file.rename(packed)
        with zipfile.ZipFile(packed) as src,zipfile.ZipFile(self.file,'w') as dst:
            for item in src.infolist():
                data=src.read(item.filename)
                dst.writestr(item,data.replace(b' fullCalcOnLoad="1"',b'') if item.filename=='xl/workbook.xml' else data)
        packed.unlink()
        with zipfile.ZipFile(self.file) as z:self.assertNotIn('fullCalcOnLoad',z.read('xl/workbook.xml').decode())
        result,stage=self.fill();receipt=result['workbookValidation']
        self.assertEqual(receipt['recalcCount'],1);self.assertEqual(receipt['verifiedCount']+receipt['recalcCount'],receipt['requiredCount'])
        with zipfile.ZipFile(stage/'填写成果.xlsx') as z:self.assertIn('fullCalcOnLoad="1"',z.read('xl/workbook.xml').decode())
        self.assertIsNone(load_workbook(stage/'填写成果.xlsx',data_only=True)['Lead Schedule']['G4'].value)


class TemplateFormulaTests(unittest.TestCase):
    """Template formulas recompute with Excel semantics: blanks are 0, IF/IFERROR choose a branch."""
    def calc(self,formula,cells):
        from workbook_plan import calculate_formula
        return calculate_formula(formula,'S',lambda sheet,cell:cells.get(cell))
    def test_blank_adjustment_counts_as_zero(self):
        self.assertEqual(self.calc('B4+C4',{'B4':Decimal(5)}),Decimal(5))
    def test_if_skips_the_division_by_zero_branch(self):
        self.assertEqual(self.calc('IF(E4=0,"",F4/E4)',{'F4':Decimal(3)}),'')
        self.assertEqual(self.calc('IF(E4=0,"",F4/E4)',{'E4':Decimal(2),'F4':Decimal(3)}),Decimal('1.5'))
        self.assertEqual(self.calc('IFERROR(F4/E4,0)',{'F4':Decimal(3)}),Decimal(0))
    def test_unhandled_division_by_zero_is_the_excel_error_value(self):
        self.assertEqual(self.calc('F4/E4',{'F4':Decimal(3)}),{'error':'#DIV/0!'})
        self.assertEqual(self.calc('IFERROR(F4/E4,0)',{'F4':Decimal(3)}),Decimal(0))
        self.assertEqual(self.calc('E4+1',{'E4':{'error':'#N/A'}}),{'error':'#N/A'})

if __name__=='__main__':unittest.main()
