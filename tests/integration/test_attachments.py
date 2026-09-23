import unittest,tempfile,sys,zipfile,json,subprocess,os
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'build/python-packages'))
from openpyxl import Workbook
from pypdf import PdfWriter
class AttachmentTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(dir=ROOT/'artifacts/test-results');self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def extract(self,name,tool='__attachment',parameters=None):
  request={'protocolVersion':1,'taskId':'attachment-test','root':str(self.root),'tool':tool,'files':[name],'mode':'auto','parameters':parameters or {},'limits':{'memoryMiB':512,'threads':1}}
  env={**os.environ,'PYTHONPATH':str(ROOT/'build/python-packages'),'PYTHONIOENCODING':'utf-8','PYTHONUTF8':'1'}
  r=subprocess.run([sys.executable,'-B',str(ROOT/'src/audit-worker/worker.py')],input=json.dumps(request),text=True,encoding='utf-8',capture_output=True,env=env,timeout=20)
  events=[json.loads(x) for x in r.stdout.splitlines()];result=next((e['result'] for e in events if e['type']=='result'),None);self.assertIsNotNone(result,events);return result
 def test_excel_excerpt_limits_and_formulas(self):
  w=Workbook();w.active.title='收入明细';w.active.append(['客户','收入','公式']);w.active.append(['甲公司',12345,'=B2*2'])
  for i in range(60):w.active.append(['第'+str(i)+'行',i])
  w.save(self.root/'资料.xlsx');r=self.extract('资料.xlsx');self.assertIn('甲公司',r['text']);self.assertIn('=B2*2',r['text']);self.assertTrue(r['truncated']);self.assertNotIn('第59行',r['text'])
 def test_docx_text_and_table(self):
  with zipfile.ZipFile(self.root/'资料.docx','w') as z:z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>审计说明</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>收入12345</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
  r=self.extract('资料.docx');self.assertIn('审计说明',r['text']);self.assertIn('收入12345',r['text'])
 def test_scanned_pdf_explicitly_reports_no_ocr(self):
  w=PdfWriter();w.add_blank_page(width=100,height=100)
  with (self.root/'扫描.pdf').open('wb') as f:w.write(f)
  r=self.extract('扫描.pdf');self.assertIn('OCR',r['text']);self.assertIn('未提取到文本',r['text'])
 def test_workbook_structure_hidden_sheets_merged_cells_and_header(self):
  w=Workbook();w.active.title='明细';w.active.append(['审计说明']);w.active.merge_cells('A1:C1');w.active.append(['客户','金额','日期']);w.active.append(['甲公司',12345,'2026-09-09']);w.active['B3'].number_format='#,##0.00';hidden=w.create_sheet('隐藏数据');hidden.sheet_state='hidden';hidden.append(['隐藏字段']);w.save(self.root/'结构.xlsx');r=self.extract('结构.xlsx','__structure',{'headerRow':2,'sheets':'明细'});self.assertEqual(r['sheetCount'],2);self.assertEqual(r['sheets'][1]['state'],'hidden');self.assertEqual(r['previews'][0]['columns'][:3],['客户','金额','日期']);self.assertIn('A1:C1',r['previews'][0]['mergedRanges']);self.assertEqual(r['previews'][0]['preview'][0][1]['format'],'#,##0.00');self.assertTrue(r['sampleOnly'])
if __name__=='__main__':unittest.main()
