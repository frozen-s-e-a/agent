"""Business acceptance cases using only the app's independent runtime.

All inputs are synthetic. Model and web fixtures are explicitly simulated;
passing these cases does not establish availability of a live remote service.
Outputs are retained for inspection under artifacts/independent-validation.
"""
from pathlib import Path
from io import BytesIO
import hashlib
import json
import sys
import traceback
import uuid
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from runtime.context import Context
from runtime.registry import dispatch, verify_registry
from runtime.transport import Response

REPO = Path(__file__).resolve().parents[2]
ROOT = REPO / 'artifacts' / 'independent-validation' / ('run-' + uuid.uuid4().hex[:10])
ROOT.mkdir(parents=True)
CASES = []
CHECKS = []


def write(name, sheets):
    p = ROOT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(sheets, list):
        sheets = {'数据': sheets}
    wb = Workbook(); wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows: ws.append(row)
    wb.save(p); wb.close()
    return str(p.relative_to(ROOT))


def kv(name, values, extra=None):
    return write(name, {'配置表': [[k,v] for k,v in values.items()], **(extra or {})})


def equal(actual, expected):
    assert actual == expected, f'{actual!r} != {expected!r}'


def sheet(result, name=None):
    p = next(ROOT / x['path'] for x in result['outputs'] if x['path'].endswith(('.xlsx','.xlsm')))
    return pd.read_excel(p, sheet_name=name or 0)


def output_cell(result, address, tab=None):
    p = next(ROOT / x['path'] for x in result['outputs'] if x['path'].endswith(('.xlsx','.xlsm')))
    wb=load_workbook(p);ws=wb[tab] if tab else wb.active;value=ws[address].value;wb.close();return value


class FixtureSession:
    def __init__(self, answer=None, route=None):
        self.answer, self.route, self.calls = answer, route, []
    def request(self, method, url, **kw):
        self.calls.append((method,url,kw))
        if self.route: return self.route(method,url,kw)
        assert url.endswith('/chat/completions'), url
        return reply({'choices':[{'message':{'content':json.dumps(self.answer,ensure_ascii=False)}}]})
    def get(self, url, **kw): return self.request('GET',url,**kw)
    def post(self, url, **kw): return self.request('POST',url,**kw)


def reply(data=None, content=None):
    if content is None: content=json.dumps(data,ensure_ascii=False).encode()
    return Response(content,200,{})


def run(name, args, check, answer=None, route=None, label=None):
    identifier=name+'-'+uuid.uuid4().hex[:8]
    source_hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.rglob('*') if p.is_file() and 'outputs' not in p.parts}
    record={'task':name,'case':label or name,'mode':'fixture-model' if answer is not None else 'fixture-http' if route else 'local', 'arguments':args}
    if answer is not None:record['fixtureAnswer']=answer
    try:
        session=FixtureSession(answer,route) if answer is not None or route else None
        ctx=Context({'root':str(ROOT),'taskId':identifier,'arguments':args,'settings':{'baseUrl':'https://fixture.invalid/v1','model':'fixture-text','visionModel':'fixture-vision','apiKey':'synthetic','cicpaCookies':'{"fixture":"synthetic"}'}}, session=session)
        result=dispatch(ctx,name)
        equal(result['success'],True)
        check(result)
        for item in result['outputs']:
            p=ROOT/item['path']; assert p.is_file();equal(hashlib.sha256(p.read_bytes()).hexdigest(),item['sha256'])
        if answer is not None: assert session.calls, 'AI processing made no model call'
        # Saving explicit configuration is the only task that writes beyond outputs.
        for p,digest in source_hashes.items(): equal(hashlib.sha256(Path(p).read_bytes()).hexdigest(),digest)
        record.update(passed=True,outputs=result['outputs'],warnings=result.get('warnings',[]))
        print('PASS',record['case'],flush=True)
    except Exception as error:
        record.update(passed=False,error=str(error),traceback=traceback.format_exc())
        print('FAIL',record['case'],str(error),flush=True)
    CASES.append(record)
    return record


def local_cases():
    # Each edit runs against its own copy. Source hashes are checked above.
    a=write('table/duplicate.xlsx',[['名称','项目','项目'],['甲','A','B'],['乙','C','D']])
    run('add_column',dict(file_path=a,keyword='',column_name='项目',delimiter='-',max_column_num=2),lambda r:equal(output_cell(r,'D2'),'A-B'))
    a=write('table/aux-data.xlsx',[['名称','辅助'],['甲','[客户:甲][部门:销售]']])
    run('asist_split',dict(file_path=a,keyword='',col_name='辅助',kv_sep=':',item_sep='[]'),lambda r:equal(output_cell(r,'C2'),'甲'))
    a=write('table/fill.xlsx',[['公司','金额'],['甲',10],[None,20],['乙',30]])
    run('fill_column',dict(file_path=a,keyword='',column_names='公司'),lambda r:equal(output_cell(r,'A3'),'甲'))
    a=write('table/header.xlsx',[['科目','金额','金额'],['名称','借方','贷方'],['现金',50,10]])
    run('headers_process',dict(file_path=a,keyword='',header_rows='1,2'),lambda r:equal(output_cell(r,'B1'),'金额借方'))
    a=write('table/merge.xlsx',[['公司','部门'],['甲','销售']])
    run('merge_column',dict(file_path=a,keyword='',column_names='公司,部门',delimiter='-',insert_index='3',insert_name='组合'),lambda r:equal(output_cell(r,'C2'),'甲-销售'))
    a=write('table/summary.xlsx',[['科目','金额'],['银行',10],['合计',10],['现金',20],['现金',20]])
    run('dropsummary',dict(file_path=a,keywords='合计',column_name='科目'),lambda r:equal(len(sheet(r)),3))
    run('insert_column',dict(file_path=a,insert_position='B',column_name='备注',fill_value='待复核',sheet_name='数据'),lambda r:equal(output_cell(r,'B2'),'待复核'))
    a=write('selection/a.xlsx',{'收入表':[['科目','金额'],['收入',20]],'其它':[['科目','金额'],['现金',99]]})
    run('select_column',dict(folder_path='selection',should_merge='是',sheet_keywords='收入',columns_to_keep='金额'),lambda r:(equal(len(sheet(r)),1),equal(sheet(r)['金额'].sum(),20)))
    run('pickup_value',dict(folder_path='selection',field_mappings=json.dumps({'金额':'B2'},ensure_ascii=False)),lambda r:equal(sorted(sheet(r)['金额'].tolist()),[20,99]))
    a=write('subjects.xlsx',[['代码','科目','辅助','期初','公司'],['1001','现金','',50,'甲'],['100101','人民币','[币种:CNY]',50,'甲'],['2001','应付','[供应商:乙]',80,'甲']])
    for task in ['tool_preview','monthly_preview','detailed_table_inspect']:
        run(task,dict(file_path=a),lambda r:equal(r['row_count'],3))
    run('tool_execute',dict(file_path=a,code_column='代码',code_name_column='科目',header_row=1,asist_column='辅助',asist_item_wrapper='[]',asist_name_split=':',grouping_column='公司'),lambda r:(equal(r['rowCount'],2),equal(r['preview'][0]['1级科目'],'现金')))
    mapping=dict(date='日期',subject_code='代码',subject_name='科目',debit='借',credit='贷',company='公司')
    this=write('monthly/this.xlsx',[['日期','代码','科目','借','贷','公司'],['2026-01-01','6001','收入',0,1200,'甲'],['2026-02-01','6001','收入',0,400,'甲']])
    last=write('monthly/last.xlsx',[['日期','代码','科目','借','贷','公司'],['2025-01-01','6001','收入',0,1000,'甲']])
    run('monthly_detect_structure',dict(file_path=this,column_mapping=json.dumps(mapping)),lambda r:equal(r['subjects'][0]['code'],'6001'))
    monthly=run('monthly_execute',dict(this_year_file=this,last_year_file=last,subjects='[{"code":"6001","direction":"贷"}]',column_mapping=json.dumps(mapping),threshold_amt=100,threshold_rate=.1),lambda r:(equal(sheet(r,'科目汇总')['本期金额'].sum(),1600),equal(sheet(r,'科目汇总')['变动额'].sum(),600)))
    if monthly['passed']:
        rpath=monthly['outputs'][0]['path']
        run('monthly_ai_analysis',dict(result_file=rpath,this_year_file=this,column_mapping=json.dumps(mapping),threshold_amt=100,threshold_rate=.1),lambda r:equal(r['rowCount'],1),answer={'变动原因线索':'一月增加200，二月新增400','待补证据':'销售明细','审计应对建议':'检查截止'})
    cur=write('margin/this.xlsx',[['产品','营业收入','营业成本','数量'],['A',1200,700,100]])
    prev=write('margin/last.xlsx',[['产品','营业收入','营业成本','数量'],['A',1000,600,100]])
    run('gross_margin_analysis',dict(this_data_path=cur,last_data_path=prev,confirmed=True),lambda r:(equal(sheet(r)['本期毛利'].sum(),500),equal(sheet(r)['毛利变动'].sum(),100)))
    run('gross_margin_analysis',dict(this_data_path=cur),lambda r:equal(r['status'],'needs_confirmation'),label='gross_margin_confirmation')
    gl=write('jet.xlsx',[['金额','摘要'],[100,'普通'],[200,'调整'],[300,'调整']])
    run('jet_test_inspect',dict(file_path=gl),lambda r:(equal(r['row_count'],3),equal(len(r['rules']),33)))
    run('jet_test_execute',dict(file_path=gl,features='[]',rules='[{"name":"调整凭证","conditions":{"摘要":"调整"}}]'),lambda r:equal(r['rule_summary'][0]['命中行数'],2))
    data=write('cutoff/source.xlsx',[['日期','科目','凭证','借','贷'],['2025-12-30','费用','1',300,0],['2026-01-02','费用','2',200,0],['2026-01-02','现金','2',0,200],['2026-01-03','费用','3',30,0]])
    cfg=kv('cutoff/config.xlsx',{'文件或文件夹路径':data,'日期':'日期','1级科目':'科目','凭证号':'凭证','借方金额':'借','贷方金额':'贷','开始日期':'2025-12-28','结束日期':'2026-01-05','模式':'金额','金额':100})
    run('cutoff_sampling',dict(config_file=cfg),lambda r:(equal(r['sample_count'],2),equal(r['population_count'],3)))
    ledger=write('ledger/甲.xlsx',{'序时账':[['科目代码','借方金额','贷方金额'],['1001',100,0],['2001',0,100]],'科目余额表':[['科目代码','科目名称','期初余额','期末余额','借方金额','贷方金额','方\n向','方\n向.1'],['1001','现金',0,100,100,0,'借','借'],['2001','应付',0,100,0,100,'贷','贷']]})
    run('ledger_check',dict(directory='ledger'),lambda r:(equal(sheet(r,'账套检查')['借贷平衡'].iloc[0],True),equal(sheet(r,'科目差异')['账表贷方差异'].sum(),0)))
    bank=write('bank/bank.xlsx',[['日期','对手','收','付','余额'],['2026-01-05','甲',100,0,100],['2026-01-07','乙',0,60,40],['2026-01-07','乙',0,40,0]])
    ledger=write('bank/gl.xlsx',[['日期','凭证','科目','摘要','对手','借','贷'],['2026-01-05','1','银行','收到甲货款','甲',100,0],['2026-01-07','2','银行','支付乙','乙',0,100]])
    run('bankflowmerge',dict(file_paths=[bank],header_start_cols=['A'],header_start_rows=[1],bank_names=['合成银行'],account_names=['合成公司'],account_numbers=['synthetic'],transaction_date_cols=['日期'],transaction_time_cols=[''],counterparty_name_cols=['对手'],counterparty_account_cols=[''],summary_cols=[''],income_cols=['收'],expense_cols=['付'],balance_cols=['余额']),lambda r:(equal(sheet(r)['收入金额'].sum(),100),equal(sheet(r)['支出金额'].sum(),100)))
    run('bank_flow',dict(gl_path=ledger,gl_date_col='日期',gl_voucher_col='凭证',gl_subject_col='科目',gl_bank_subject='银行',gl_summary_col='摘要',gl_partner_direct_col='对手',gl_partner_col='',gl_debit_col='借',gl_credit_col='贷',bank_path=bank,bank_date_col='日期',bank_income_col='收',bank_expense_col='付',bank_partner_col='对手'),lambda r:(equal(r['matched_groups'],2),equal(sheet(r,'核对汇总')['未匹配笔数'].sum(),0),equal(sheet(r,'匹配明细')['匹配组差额'].sum(),0)))
    for variant,bank_lines,counter_lines in [('split',[[100,0]],[[0,60,'甲'],[0,40,'乙']]),('proportion',[[100,0]],[[0,90,'甲'],[0,60,'乙']]),('matching_lines',[[60,0],[40,0]],[[0,60,'甲'],[0,40,'乙']])]:
        gl2=write('bank/'+variant+'-gl.xlsx',[['日期','凭证','科目','摘要','对手','借','贷'],*[['2026-01-05','1','银行','合成收款','',d,c] for d,c in bank_lines],*[['2026-01-05','1','应收','合成收款',p,d,c] for d,c,p in counter_lines]])
        b2=write('bank/'+variant+'-bank.xlsx',[['日期','对手','收','付'],['2026-01-05','甲',60,0],['2026-01-05','乙',40,0]])
        run('bank_flow',dict(gl_path=gl2,gl_date_col='日期',gl_voucher_col='凭证',gl_subject_col='科目',gl_bank_subject='银行',gl_summary_col='摘要',gl_partner_direct_col='',gl_partner_col='对手',gl_debit_col='借',gl_credit_col='贷',bank_path=b2,bank_date_col='日期',bank_income_col='收',bank_expense_col='付',bank_partner_col='对手'),lambda r:(equal(r['matched_groups'],2),equal(sheet(r,'核对汇总')['未匹配笔数'].sum(),0),equal(sheet(r,'匹配明细')['序时账_借'].sum(),100),equal(sheet(r,'匹配明细')['序时账核对金额'].tolist(),[60,40])),label='bank_'+variant)
    common=dict(gl_date_col='日期',gl_voucher_col='凭证',gl_subject_col='科目',gl_bank_subject='银行',gl_summary_col='摘要',gl_partner_direct_col='对手',gl_partner_col='',gl_debit_col='借',gl_credit_col='贷',bank_date_col='日期',bank_income_col='收',bank_expense_col='付',bank_partner_col='对手')
    g=write('bank/scoped-gl.xlsx',[['日期','凭证','科目','摘要','对手','借','贷','银行账号','公司名称'],['2026-01-05','1','银行','收款','客户',100,0,'001','甲公司'],['2026-01-05','1','银行','收款','客户',100,0,'001','乙公司']])
    b=write('bank/scoped-bank.xlsx',[['日期','对手','收','付','银行账号','公司名称'],['2026-01-05','客户',100,0,'001','乙公司'],['2026-01-05','客户',100,0,'001','甲公司']])
    def scoped_check(r):
        df=sheet(r,'匹配明细');equal((df['流水_公司名称']==df['序时账_公司名称']).all(),True);equal(r['matched_groups'],2)
    run('bank_flow',{**common,'gl_path':g,'bank_path':b},scoped_check,label='bank_company_account_scope')
    for variant,glrows,brows,expected in [
      ('many_to_many',[['2026-01-05',40],['2026-01-07',60]],[['2026-01-06',30],['2026-01-08',70]],'L4'),
      ('cross_month',[['2026-01-31',100]],[['2026-02-01',60],['2026-02-02',40]],'L9'),
      ('date_formats',[['20260105',100]],[['2026年1月5日',100]],'L1')]:
        g=write('bank/'+variant+'-gl.xlsx',[['日期','凭证','科目','摘要','对手','借','贷'],*[[d,str(i),'银行','收款','甲',v,0] for i,(d,v) in enumerate(glrows)]])
        b=write('bank/'+variant+'-bank.xlsx',[['日期','对手','收','付'],*[[d,'甲',v,0] for d,v in brows]])
        run('bank_flow',{**common,'gl_path':g,'bank_path':b},lambda r,expected=expected:(equal(r['matched_groups'],1),equal(sheet(r,'核对汇总')['未匹配笔数'].sum(),0),equal(sheet(r,'匹配明细')['匹配规则'].str.startswith(expected).all(),True)),label='bank_'+variant)
    cfg=write('bank/config.xlsx',{'主信息-基础信息':[['名称','公司'],['TEMPORARYCODE','SCCBACONFIRMATION_ORGNAME'],['001','甲']], '主信息-固定信息':[['名称','年份'],['TEMPORARYCODE','SCCBACONFIRMATION_YEAR'],[None,'2026']], '函证信息-存款':[['编号','金额'],['TEMPORARYCODE','GROUP1ROW0COL5'],['001',100],['001',200]],'函证信息-借款':[['编号','金额'],['TEMPORARYCODE','GROUP2ROW0COL3'],['001',80]]})
    for task in ['bank_digit2','bank_paper1','bank_paper2']:
        def check_confirmation(r):
            equal(r['confirmation_count'],1)
            output=next(ROOT/x['path'] for x in r['outputs'] if x['name'].startswith('银行函证-'))
            frame=pd.read_excel(output,sheet_name='函证信息',header=3)
            equal(frame['GROUP1ROW0COL5'].dropna().tolist(),[100,200])
            equal(frame['GROUP2ROW0COL3'].dropna().tolist(),[80])
        run(task,dict(config_file=cfg),check_confirmation)


def workbook_cases():
    src=write('books/source.xlsx',{'数据':[['名称','金额'],['甲',100],['乙',200]]})
    target=write('books/target.xlsx',{'目标':[['名称','金额','参考','填充'],['甲',10,1,'=B2*2'],['乙',20,1,None]]})
    wb=load_workbook(ROOT/target);wb.active['B4']="='[old.xlsx]数据'!A1";wb.active['A2'].hyperlink='https://example.invalid/old.xlsx';wb.save(ROOT/target);wb.close()
    cfg=write('books/extract.xlsx',{'配置表':[['文件或文件夹路径'],[target]]})
    run('link_extractor',dict(config_file=cfg),lambda r:equal(r['rowCount'],2))
    cfg=write('books/replace-text.xlsx',{'配置表':[['文件夹',target,'关键词',''],['旧文本','新文本'],['old.xlsx','new.xlsx']]})
    run('link_replace_text',dict(config_file=cfg),lambda r:equal(output_cell(r,'B4','目标'),"='[new.xlsx]数据'!A1"))
    cfg=write('books/replace-workbook.xlsx',{'配置表':[['工作簿路径','老链接文本','新链接文本'],[target,'old.xlsx','new.xlsx']]})
    run('link_replace_workbook',dict(config_file=cfg),lambda r:equal(output_cell(r,'B4','目标'),"='[new.xlsx]数据'!A1"))
    cfg=write('books/range.xlsx',{'配置表':[['现工作簿路径','原工作簿路径','现工作表名','原工作表名','现单元格区域','原单元格区域'],[target,src,'目标','数据','B2:B3','B2:B3']]})
    run('link_replace_range',dict(config_file=cfg),lambda r:equal(output_cell(r,'B3','目标').split('!')[-1],'B3'))
    run('link_ref_replace',dict(target_path=target,sheet_keywords='目标',ref_axis='C',target_axis='D',example_cell='D2'),lambda r:equal(output_cell(r,'D3','目标'),'=B3*2'))
    for mode in ['file','folder']:
        cfg=write('books/collect-'+mode+'.xlsx',{'路径匹配':[['目标工作簿','来源工作簿','来源文件夹'],[target,src,'books/source-dir']], '表匹配关系':[['来源表名','目标表名','来源工作簿文件关键词'],['数据','收入','source']]})
        write('books/source-dir/source.xlsx',{'数据':[['名称','金额'],['甲',100],['乙',200]]})
        run('multi_collect_'+mode,dict(config_file=cfg),lambda r:equal(output_cell(r,'B3','收入'),200))
    notes=write('notes.xlsx',[['科目','金额'],['现金',20],['银行',30]])
    cfg=write('comments.xlsx',{'配置表':[['文件夹或文件路径','工作表','开始行号','结束行号','是否清除所有批注'],[notes,'数据',1,3,'是']]})
    run('insert_footnote_comments',dict(config_file=cfg),lambda r:equal(r['preview'][0]['表首表尾数量'],1))
    wb=load_workbook(ROOT/notes);wb.active['A1'].comment=Comment('表首','test');wb.active['A3'].comment=Comment('表尾','test');wb.save(ROOT/notes);wb.close()
    cfg=kv('scope.xlsx',{'文件路径':notes,'表名':'数据','表头批注关键词':'表首','表尾批注关键词':'表尾'})
    run('extract_disclosure_scope',dict(config_file=cfg),lambda r:equal(sheet(r)['披露口径'].tolist(),['现金','银行']))
    cfg=kv('disclosure.xlsx',{'文件夹或文件路径':notes,'表首关键词':'表首','表尾关键词':'表尾'})
    run('extract_disclosure_tables',dict(config_file=cfg),lambda r:equal(r['table_count'],1))


def workpaper_cases():
    columns=['公司名称','科目名称','本期借方','本期贷方','期初余额','期末余额']
    src=write('papers/source.xlsx',[columns,['甲','应收',100,10,20,110],['乙','应收',200,30,40,210]])
    template=write('papers/templates/应收.xlsx',{'底稿':[columns,[None]*6,[None]*6]})
    wb=load_workbook(ROOT/template);ws=wb.active
    for i,col in enumerate(columns,1):ws.cell(1,i).comment=Comment(col,'test')
    ws['A2'].comment=Comment('表首','test');ws['A3'].comment=Comment('表尾','test');wb.save(ROOT/template);wb.close()
    configs={'field_mapping':dict(zip(columns,columns)),'subject_mapping':[{'subject_name':'应收','report_subject':'应收账款'}],'template_assignment':{'应收账款':{'template_path':template,'sheet_name':'底稿'}}}
    # Public save endpoint; the later generation consumes these exact saved files.
    for kind,data in configs.items():
        run('detailed_table_save_config',dict(work_dir='papers',config_type=kind,config_data=json.dumps(data,ensure_ascii=False)),lambda r:equal(json.loads((ROOT/r['saved_to']).read_text(encoding='utf-8')),data),label='save_'+kind)
    run('detailed_table_inspect_template',dict(template_dir='papers/templates'),lambda r:equal(set(r['templates'][0]['sheets'][0]['annotations']),set(columns+['表首','表尾'])))
    run('detailed_table_validate',dict(work_dir='papers',file_path=src),lambda r:equal(r['validation_passed'],True))
    run('detailed_table_generate',dict(work_dir='papers',file_path=src,confirmed=True),lambda r:(equal(r['company_count'],2),equal(sum(pd.read_excel(ROOT/x['path'])['期末余额'].dropna().sum() for x in r['outputs'] if x['name'].endswith('应收.xlsx')),320)))
    extra=['客户编码','客户-名称','供应商-名称','分类']
    current=write('aging/periods/1.2026.xlsx',[columns+extra,['甲',' 应收 ',15,5,40,50,'001','客户A','','押金'],['甲','应收',15,5,40,50,'001','客户A','','押金'],['甲','应收',20,0,10,30,'002','客户B','','贸易'],['甲','应收',5,0,0,5,'003','客户C','','贸易']])
    write('aging/periods/2.2025.xlsx',[columns+extra,['甲','应收',40,10,50,80,'001','客户A','','押金'],['甲','应收',10,0,0,10,'002','客户B','','贸易'],['甲','应收',0,0,0,0,'003','客户C','','贸易']])
    fields=columns+['客户编号','客户名称','披露','关联方','账龄:1','账龄:2','账龄:3']
    tpl=write('aging/template.xlsx',{'底稿':[fields+['核对'],[None]*len(fields)+['=F2'],['合计']+[None]*(len(fields)-1)+['=SUM(F2:F2)']]})
    w=load_workbook(ROOT/tpl);ws=w.active
    for i,field in enumerate(fields,1):ws.cell(1,i).comment=Comment(field+'#' if field=='客户名称' else field,'test')
    ws['A2'].comment=Comment('表首','test');ws['A3'].comment=Comment('尾行','test');w.save(ROOT/tpl);w.close()
    wrapped={
      'field_mapping':{'mapping':{**dict(zip(columns,columns)),'客户编号':'客户编码','客户名称':'客户-名称/供应商-名称','披露':'$分类','关联方':'$客户-名称'},'text_fields':['客户编号'],'uniqueness_fields':['客户编号']},
      'subject_mapping':{'mapping':[{'subject':'应收','report_subject':'应收账款','symbol':'借','sign':'正','source':'default','generate':True},{'subject':'不存在','report_subject':'未启用','generate':False}]},
      'template_assignment':{'assignment':[{'report_subject':'应收账款','template_path':tpl,'sheet_name':'底稿'}]},
      '披露':{'mapping':{'押金':'保证金','贸易':'应收货款'}},'关联方':{'mapping':{'客户A':'子公司','客户B':'无关联','客户C':'无关联'}}}
    for kind,data in wrapped.items():run('detailed_table_save_config',dict(work_dir='aging',config_type=kind,config_data=json.dumps(data,ensure_ascii=False)),lambda r:equal(Path(r['saved_to']).is_file(),True),label='wrapped_'+kind)
    def check_aging(r):
        df=sheet(r,'底稿').iloc[:3]
        equal(df['期末余额'].sum(),135)
        equal(df['账龄:1'].tolist(),[30,20,5]);equal(df['账龄:2'].tolist(),[40,10,0]);equal(df['账龄:3'].tolist(),[30,0,0])
        equal(df['披露'].tolist(),['保证金','应收货款','应收货款'])
        equal(output_cell(r,'G2','底稿'),'001');equal(output_cell(r,'N4','底稿'),'=F4');equal(output_cell(r,'N5','底稿'),'=SUM(F2:F4)')
        equal(sum('账龄合计' in warning for warning in r['warnings']),0)
    run('detailed_table_generate',dict(work_dir='aging',file_path='aging/periods',confirmed=True),check_aging,label='multi_period_aging_group_mapping_formula')
    empty_configs={**configs,'paths':{'sources':{'default':'papers/source.xlsx'},'copy_empty_template':True},'subject_mapping':[{'subject_name':'应付','report_subject':'应收账款'}]}
    directory=ROOT/'empty-papers/config';directory.mkdir(parents=True)
    for kind,data in empty_configs.items():(directory/(kind+'.json')).write_text(json.dumps(data,ensure_ascii=False),encoding='utf8')
    def check_empty(r):
        equal(r['company_count'],2)
        equal(sum(x['name'].endswith('.xlsx') for x in r['outputs']),3)
        equal(output_cell(r,'F2','底稿'),None)
        summary=next(ROOT/x['path'] for x in r['outputs'] if x['name']=='明细底稿生成清单.xlsx')
        equal(pd.read_excel(summary)['写入行数'].tolist(),[0,0])
    run('detailed_table_generate',dict(work_dir='empty-papers'),check_empty,label='empty_template_per_company')
    src=write('cleanup.xlsx',[['1级科目','客户','项目','借初','贷初','借末','贷末'],['应收','甲','A',20,0,110,0],['应付','乙','B',0,30,0,70]])
    cfg=write('cleanup-config.xlsx',{'余额配置':[['期初借方余额','借初'],['期初贷方余额','贷初'],['期末借方余额','借末'],['期末贷方余额','贷末']], '辅助核算配置':[['1级科目','方向','标准辅助核算1'],['应收','借','客户-项目'],['应付','贷','客户']], '数据源':[['文件或文件夹路径'],[src]]})
    run('subject_clean_after',dict(config_file=cfg),lambda r:(equal(sheet(r)['期末标准余额'].tolist(),[110,70]),equal(sheet(r)['标准辅助核算1'].tolist(),['甲','乙'])))
    write('mini/source/甲-科目余额表.xlsx',[['编码','名称','金额'],['1001','现金',50]])
    write('mini/source/甲-序时账.xlsx',[['编码','名称','金额'],['1001','现金',50]])
    cfg=write('mini/config.xlsx',{'文件储存':[['数据源文件夹路径','mini/source'],['源文件是否分公司保存','否']], **{k+'映射':[['列','编码','名称','金额'],['映射','科目代码','科目名称','借方金额']] for k in ['科目余额表','辅助科目余额表','序时账']}})
    run('xlsx_writer',dict(config_file=cfg),lambda r:equal(output_cell(r,'C2','科目余额表'),50))


def document_cases():
    (ROOT/'docs/company/contracts').mkdir(parents=True)
    (ROOT/'docs/company/contracts/合同.txt').write_text('合成合同，编号 HT001，合同金额100元。',encoding='utf-8')
    doc='docs/company/contracts'
    answer={'records':[{'编号':'HT001','金额':100}]}
    run('info_extract',dict(folder=doc,role_prompt='提取编号和金额',max_workers=2),lambda r:equal(sheet(r)['金额'].sum(),100),answer=answer)
    cfg=write('multi.xlsx',{'基础配置':[['处理文件夹','docs']], '文件夹配置':[['文件夹名称','匹配模式','提示词','是否合并'],['contracts','完全','编号与金额','是']]})
    run('info_extract_multi',dict(config_file=cfg),lambda r:equal(sheet(r)['公司名称'].iloc[0],'company'),answer=answer)
    cfg=write('classify.xlsx',{'基础配置':[['处理文件夹',doc],['输出模式','公司']], '类型配置':[['文件类别','说明'],['合同','双方合同']]})
    run('file_classification',dict(config_file=cfg),lambda r:(equal(r['preview'][0]['类别'],'合同'),equal((ROOT/r['preview'][0]['成果文件']).read_text(encoding='utf-8'),'合成合同，编号 HT001，合同金额100元。')),answer={'类别':'合同','公司名称':'合成公司','理由':'合同标题'})
    from PIL import Image
    (ROOT/'images').mkdir();Image.new('RGB',(100,50),'white').save(ROOT/'images/合成凭证.png')
    run('photo_rename',dict(directory='images',user_prompt='按凭证编号命名'),lambda r:equal(r['preview'][0]['新文件名'],'合成凭证HT001.png'),answer={'文件名':'合成凭证HT001'})
    cfg=kv('match.xlsx',{'任务模式':'分类','任务指令及结果要求':'分类合同','批处理数量':20},{'待匹配数据':[['正文'],['合同 HT001']], '分类项':[['分类名称'],['合同']]})
    run('text_match',dict(config_file=cfg),lambda r:equal(sheet(r)['匹配结果'].iloc[0],'合同'),answer={'records':[{'记录编号':0,'匹配结果':'合同'}]})
    cfg=write('aux-choice.xlsx',{'辅助核算配置':[['1级科目','辅助核算','示例'],['应收','客户','甲'],['应收','项目','A']]})
    run('auxiliary_selection',dict(config_file=cfg),lambda r:equal(sheet(r)['标准辅助核算1'].iloc[0],'客户-项目'),answer={'d':'借','s':'客户-项目'})
    write('field-data/source.xlsx',[['编号','金额'],['HT001',100]])
    cfg=write('field-check.xlsx',{'配置表':[['扫描行数',10],['其它',None],['资料类别','资料目录'],['合同','field-data']], '合同':[['合同编号','合同金额']]})
    run('data_field_check',dict(config_file=cfg),lambda r:equal(sheet(r)['检查结果'].tolist(),['已找到','已找到']),answer={'mapping':{'合同编号':'编号','合同金额':'金额'}})
    cfg=write('paper-verify.xlsx',{'账面数据':[['编号','金额'],['HT001',100]]})
    run('paper_verify',dict(file_path=doc,config_file=cfg,match_keys='编号',max_workers=1,extra_prompt=''),lambda r:equal(r['matched_count'],1),answer=answer)
    cfg=write('voucher-check.xlsx',{'核查要点':[['科目','核查要点','必要附件'],['应收','核查金额与审批','合同']]})
    run('voucher_check',dict(source_paths=['images'],config_file=cfg),lambda r:equal(sheet(r)['凭证号'].iloc[0],'HT001'),answer={'records':[{'科目':'应收','凭证号':'HT001','金额':100,'检查结果':'待补证据'}]})
    book=write('formulas.xlsx',{'目标':[['项目','金额'],['甲',None]], '来源':[['项目','金额'],['甲',100]]})
    cfg=write('formula-config.xlsx',{'配置表':[['是否执行','文件路径','填充表','链接表','公式类型'],['是',book,'目标','来源','直接引用']], '提示词':[['公式类型','提示词'],['直接引用','链接来源金额']]})
    run('ai_formula_filler',dict(config_file=cfg),lambda r:equal(output_cell(r,'B2','目标'),"='来源'!B2"),answer={'formulas':[{'cell':'B2','formula':"='来源'!B2"}]})
    cfg=kv('allocation.xlsx',{'文件夹':'field-data','科目余额表文件名关键词':''},{'基础代码':[['科目编码','科目名称'],['1001','现金']],'其它项':[['科目编码','科目名称']]})
    run('subject_allocation',dict(config_file=cfg),lambda r:equal(sheet(r)['科目编码'].astype(str).iloc[0],'1001'),answer={'records':[{'科目编码':'1001','科目名称':'现金','科目级别':'1','是否需要派工':'是','依据':'合成配置'}]})


def company_cases():
    run('address_split',dict(address_list='浙江省杭州市西湖区文一路',address_column=''),lambda r:equal(sheet(r)['county'].iloc[0],'西湖区'))
    run('id_card',dict(id_list='11010519491231002X',id_column=''),lambda r:equal(sheet(r)['gender'].iloc[0],'女'))
    write('companies/基础工商信息.xlsx',[['企业名称','法定代表人','登记状态','企业地址','电话','邮箱','参保人数'],['合成审计公司','张甲','存续','测试路1号','01011112222','a@test.invalid',10],['合成客户公司','张乙','注销','测试路2号','01011112222','b@other.invalid',1]])
    write('companies/股东信息.xlsx',[['公司名称','股东','出资比例'],['合成审计公司','合成投资公司','60%'],['合成客户公司','合成投资公司','80%']])
    write('input/关联方核查配置表.xlsx',{'客户名单':[['公司名称'],['合成客户公司']],'供应商名单':[['公司名称']]})
    run('related_party',dict(audit_target='合成审计公司',data_dir='companies',equity_penetration='否',threshold=50),lambda r:(equal(r['missing_companies'],[]),equal(set(sheet(r)['维度']),{'D02-共同股东','D08-联系方式共用','D11-参保人数异常','D12-交易对手状态异常'})))
    w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    root=ET.Element('{'+w+'}document');body=ET.SubElement(root,'{'+w+'}body')
    p=ET.SubElement(body,'{'+w+'}p');r=ET.SubElement(p,'{'+w+'}r');ET.SubElement(r,'{'+w+'}t').text='1. 基本信息'
    table=ET.SubElement(body,'{'+w+'}tbl')
    for values in [['字段','内容'],['企业名称','合成公司']]:
        tr=ET.SubElement(table,'{'+w+'}tr')
        for v in values:
            tc=ET.SubElement(tr,'{'+w+'}tc');r=ET.SubElement(ET.SubElement(tc,'{'+w+'}p'),'{'+w+'}r');ET.SubElement(r,'{'+w+'}t').text=v
    (ROOT/'qcc').mkdir()
    with zipfile.ZipFile(ROOT/'qcc/合成公司.docx','w') as z:z.writestr('word/document.xml',ET.tostring(root))
    run('qcc_processor',dict(folder_path='qcc'),lambda r:(equal(sheet(r)['内容'].iloc[0],'合成公司'),equal(pd.read_excel(ROOT/r['data_dir']/'基础工商信息.xlsx')['企业名称'].tolist(),['合成公司'])))


def network_cases():
    from datetime import datetime
    archive=BytesIO()
    with zipfile.ZipFile(archive,'w') as z:
        data=BytesIO();pd.DataFrame({'企业名称':['合成公司']}).to_excel(data,index=False)
        z.writestr('基础工商信息.xlsx',data.getvalue())
    def route(method,url,kw):
        if 'szse_stock.json' in url:return reply({'stockList':[{'code':'000001','orgId':'synthetic','zwjc':'合成公司'}]})
        if 'hisAnnouncement/query' in url:return reply({'totalAnnouncement':1,'announcements':[{'announcementTitle':'合成年报','announcementTime':1767225600000,'adjunctUrl':'synthetic.pdf','secName':'合成公司'}]})
        if url.endswith('.pdf'):return reply(content=b'%PDF-1.4\n% Synthetic fixture\n%%EOF')
        if url.endswith('17833.html'):return reply(content=b'<a href="history.xlsx">history</a>')
        if url.endswith('history.xlsx'):
            buf=BytesIO();pd.DataFrame({'日期':['2026-01-01'],'美元':[700]}).to_excel(buf,index=False);return reply(content=buf.getvalue())
        if 'RMBQuery' in url:return reply(content='<table><tr><th>日期</th><th>美元</th></tr><tr><td>2026-01-02</td><td>701</td></tr></table>'.encode())
        if 'home_search' in url:return reply({'status_code':0,'data':{'list':[{'company_name':'合成公司','org_id':'synthetic'}]}})
        if 'upload_local' in url:return reply({'status_code':0,'data':{'batch_id':'synthetic-batch','hit_num':1,'all_num':1}})
        if 'enter_search' in url:return reply({'status_code':0,'data':{'list':[{'company_name':'合成公司'}]}})
        if 'get_dimension_class' in url:return reply({'status_code':0,'data':[{'code':'D01','children':[{'code':'D02'}]}]})
        if 'task_list' in url:return reply({'status_code':0,'data':{'list':[{'type':'批量查询','status':1,'date':datetime.now().strftime('%Y-%m-%d %H:%M:%S'),'task_id':'synthetic-task','url':'/synthetic-export.zip'}]}})
        if url.endswith('synthetic-export.zip'):return reply(content=archive.getvalue())
        if 'find_company_basic_info' in url:return reply({'status_code':0,'data':{'company_name':'合成公司'}})
        if 'enterprise_equity' in url:return reply({'status_code':0,'data':{'holders':{'children':[]},'invests':{'children':[{'name':'合成子公司','czbl':'60%'}]}}})
        if 'stock_holder_newest' in url or 'main_person_list' in url:return reply({'status_code':0,'data':{'list':[]}})
        if 'listing.szse.cn/api/ras/projectrends/query' in url:return reply({'data':[{'prjid':'synthetic','updtdt':'2026-01-01','cmpnm':'合成公司'}]})
        if 'listing.szse.cn/api/ras/projectrends/details' in url:return reply({'data':{'prjst':'已问询','csrcind':'合成行业','disclosureMaterials':[{'matnm':'招股说明书','dfnm':'合成招股书','ddt':'2026-01-01','dfpth':'synthetic.pdf'}]}})
        if 'query.sse.com.cn' in url:
            params=kw.get('params',{})
            data=[]
            if params.get('sqlId')=='SH_XM_LB' and params.get('pageHelp.pageNo')=='1':data=[{'updateDate':'20260101','stockAuditName':'合成公司','stockAuditNum':'synthetic','currStatus':'2','stockIssuer':[{'s_issueCompanyAbbrName':'合成','s_csrcCodeDesc':'合成行业'}]}]
            if params.get('sqlId')=='GP_COMMON_FILE_SEARCH':data=[{'fileType':30,'fileUpdTime':'20260101','filePath':'/synthetic.pdf','fileTitle':'合成招股书'}]
            return reply(content=('fixture('+json.dumps({'result':data},ensure_ascii=False)+')').encode())
        if 'infoResult.do' in url:
            items=[{'companyName':'合成公司','id':'synthetic'}] if kw['data']['page']=='0' else []
            return reply(content=('fixture('+json.dumps([{'listInfo':{'content':items}}],ensure_ascii=False)+')').encode())
        if 'infoDetailResult.do' in url:return reply(content=('fixture('+json.dumps([{'projectNews':{'stockName':'合成','status':'P02','updateDate':{'time':1767225600000}},'xxgkInfo':{'GPFXSMS':{'SYG':[{'destFilePath':'/synthetic.pdf','disclosureTitle':'合成招股书','publishDate':'2026-01-01'}]}}}],ensure_ascii=False)+')').encode())
        raise AssertionError('Unexpected fixture request '+url)
    run('announcement',dict(stock_codes=['000001'],start_date='2026-01-01',end_date='2026-01-02',market_type='京沪深',category_name='年报',keyword=''),lambda r:equal(r['download_count'],1),route=route)
    run('currency',{},lambda r:equal(sheet(r)['美元'].tolist(),[700,701]),route=route)
    run('cicpa_query',dict(query_mode='search',company_name='合成公司'),lambda r:equal(sheet(r)['org_id'].iloc[0],'synthetic'),route=route)
    run('cicpa_query',dict(query_mode='detail',org_id='synthetic'),lambda r:equal(r['details'][0]['basic_info']['company_name'],'合成公司'),route=route,label='cicpa_detail')
    run('cicpa_query',dict(query_mode='subsidiary',company_name='合成公司',subsidiary_threshold=50),lambda r:equal(sheet(r)['name'].iloc[0],'合成子公司'),route=route,label='cicpa_subsidiary')
    run('cicpa_query',dict(query_mode='export',company_name='合成公司'),lambda r:(equal(r['dimensions'],2),equal(pd.read_excel(ROOT/r['data_dir']/'基础工商信息.xlsx')['企业名称'].tolist(),['合成公司'])),route=route,label='cicpa_export_protocol')
    for task in ['ipo_sh','ipo_sz','ipo_bj']:
        run(task,dict(start_date='2026-01-01',end_date='2026-01-02',update_mode='项目'),lambda r:equal(sum(x['name'].endswith('.pdf') for x in r['outputs']),1),route=route)


def main():
    equal(verify_registry(),63)
    for group in [local_cases,workbook_cases,workpaper_cases,document_cases,company_cases,network_cases]:
        try:group()
        except Exception: print(traceback.format_exc(),flush=True)
    passed={x['task'] for x in CASES if x.get('passed')}
    from runtime.registry import TASKS
    report={'project':str(ROOT),'liveServicesVerified':False,'cases':len(CASES),'passedCases':sum(x.get('passed',False) for x in CASES),'tasksCovered':len(passed),'missingTasks':sorted(set(TASKS)-passed),'results':CASES}
    (ROOT/'验收记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='results'},ensure_ascii=False),flush=True)
    return 0 if not report['missingTasks'] and report['cases']==report['passedCases'] else 1


if __name__=='__main__':sys.exit(main())
