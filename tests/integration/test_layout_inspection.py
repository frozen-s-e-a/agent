import unittest,tempfile,sys,json,subprocess,os
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'build/python-packages'))
from openpyxl import Workbook
class LayoutTests(unittest.TestCase):
 def test_blank_first_row_ragged_title_hidden_and_late_columns(self):
  with tempfile.TemporaryDirectory(dir=ROOT/'artifacts/test-results') as d:
   root=Path(d);w=Workbook();s=w.active;s.title='问题';s['A2']='问题';s['B2']='说明';s['H3']='右侧数据';title=w.create_sheet('底稿');title['A1']='底稿标题';title.merge_cells('A1:D1');title.append(['编号','金额']);title.append(['001',100]);hidden=w.create_sheet('隐藏说明');hidden.sheet_state='hidden';hidden['D5']='注释';w.save(root/'布局.xlsx')
   req={'protocolVersion':1,'taskId':'layout','root':str(root),'files':['布局.xlsx'],'tool':'inspect','mode':'auto','parameters':{},'limits':{'memoryMiB':512,'threads':1}}
   r=subprocess.run([sys.executable,'-B',str(ROOT/'src/audit-worker/worker.py')],input=json.dumps(req),text=True,encoding='utf-8',capture_output=True,env={**os.environ,'PYTHONPATH':str(ROOT/'build/python-packages'),'PYTHONIOENCODING':'utf-8','PYTHONUTF8':'1'},timeout=30)
   events=[json.loads(x) for x in r.stdout.splitlines()];result=next((e['result'] for e in events if e['type']=='result'),None);self.assertIsNotNone(result,events);self.assertTrue(result['structureOnly']);self.assertEqual(result['inputRows'],3);report=json.loads((Path(result['outputDir'])/'结构概览.json').read_text(encoding='utf-8'))['files'][0];self.assertEqual(report['sheetCount'],3);self.assertEqual(report['sheets'][2]['state'],'hidden');self.assertIn('A1:D1',report['previews'][1]['mergedRanges']);cells=report['previews'][0]['rawRows'][2]['cells'];self.assertEqual(cells[7]['value'],'右侧数据');self.assertFalse(report['previews'][0]['headerConfirmed']);self.assertEqual(result['completeness'],'structure_sample')
if __name__=='__main__':unittest.main()
