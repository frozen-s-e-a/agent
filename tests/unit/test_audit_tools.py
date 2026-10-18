"""底稿任务新工具（audit_tools）的测试：合成的应收账款底稿、余额表和序时账，走一遍账龄和前五名的填写。"""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src' / 'audit-mcp'))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

import audit_tools  # noqa: E402
import audit_tools_fixtures as fx  # noqa: E402
from audit_tools import numbers, query  # noqa: E402

TASK = 'T-0c21'
JOURNAL_FIELDS = [
    {'col': 'A', 'name': '日期', 'type': 'date', 'meaning': 'voucher_date', 'basis': 'inferred'},
    {'col': 'B', 'name': '凭证号', 'type': 'text', 'meaning': 'voucher_no', 'basis': 'inferred'},
    {'col': 'C', 'name': '科目编码', 'type': 'text', 'meaning': 'account_code', 'basis': 'inferred'},
    {'col': 'D', 'name': '摘要', 'type': 'text', 'meaning': 'summary', 'basis': 'inferred'},
    {'col': 'E', 'name': '客户', 'type': 'text', 'meaning': 'counterparty', 'basis': 'inferred'},
    {'col': 'F', 'name': '借方金额', 'type': 'decimal', 'meaning': 'debit', 'basis': 'inferred'},
    {'col': 'G', 'name': '贷方金额', 'type': 'decimal', 'meaning': 'credit', 'basis': 'inferred'},
]
BALANCE_FIELDS = [
    {'col': 'A', 'name': '科目编码', 'type': 'text', 'meaning': 'account_code', 'basis': 'inferred'},
    {'col': 'B', 'name': '科目名称', 'type': 'text', 'meaning': 'account_name', 'basis': 'inferred'},
    {'col': 'C', 'name': '期初余额', 'type': 'decimal', 'meaning': 'opening', 'basis': 'inferred'},
    {'col': 'D', 'name': '本期借方', 'type': 'decimal', 'meaning': 'debit', 'basis': 'inferred'},
    {'col': 'E', 'name': '本期贷方', 'type': 'decimal', 'meaning': 'credit', 'basis': 'inferred'},
    {'col': 'F', 'name': '期末余额', 'type': 'decimal', 'meaning': 'closing', 'basis': 'inferred'},
]
AGING_SQL = '''WITH t AS (SELECT CASE WHEN aging_days("日期", DATE '2025-12-31') <= 365 THEN '1年以内'
  WHEN aging_days("日期", DATE '2025-12-31') <= 730 THEN '1-2年' WHEN aging_days("日期", DATE '2025-12-31') <= 1095 THEN '2-3年'
  ELSE '3年以上' END AS 账龄, "借方金额" - "贷方金额" AS 金额 FROM D002 WHERE "科目编码" = '1122')
SELECT 账龄, SUM(金额) AS 金额 FROM t GROUP BY 账龄 ORDER BY CASE 账龄 WHEN '1年以内' THEN 1 WHEN '1-2年' THEN 2 WHEN '2-3年' THEN 3 ELSE 4 END'''
TOP_SQL = '''SELECT "客户" AS 客户, SUM("借方金额" - "贷方金额") AS 余额 FROM D002 WHERE "科目编码" = '1122'
GROUP BY "客户" HAVING SUM("借方金额" - "贷方金额") <> 0 ORDER BY 余额 DESC LIMIT 5'''


class Project:
    """一个临时项目：状态目录单独放，和宿主的做法一样。"""

    def __init__(self, invalid=False):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / 'project'
        self.root.mkdir()
        self.state = Path(self.tmp.name) / 'state'
        fx.build(self.root, invalid=invalid)

    def call(self, name, args):
        previous = os.environ.get('AUDIT_STATE_DIR')
        os.environ['AUDIT_STATE_DIR'] = str(self.state)
        try:
            return audit_tools.call(self.root, name, args)
        finally:
            if previous is None:
                os.environ.pop('AUDIT_STATE_DIR', None)
            else:
                os.environ['AUDIT_STATE_DIR'] = previous

    def ok(self, name, args):
        envelope = self.call(name, args)
        assert envelope['ok'], (name, envelope)
        return envelope['data']

    def restart(self):
        audit_tools._contexts.pop(str(self.root), None)

    def close(self):
        self.restart()
        self.tmp.cleanup()

    def imported(self):
        """解析底稿、导入两张资料表：→ (W-, 余额表 D-, 序时账 D-)。"""
        wp = self.ok('workpaper_parse', {'file': '应收账款底稿.xlsx', 'sheets': None})
        balance = self.ok('data_inspect', {'file': '科目余额表2025.xlsx'})['tables'][0]
        journal = self.ok('data_inspect', {'file': '序时账2025.xlsx'})['tables'][0]
        d1 = self.ok('data_import', {'candidate': balance['candidate'], 'kind': 'trial_balance', 'fields': BALANCE_FIELDS})['table']
        d2 = self.ok('data_import', {'candidate': journal['candidate'], 'kind': 'journal', 'fields': JOURNAL_FIELDS, 'excludeInvalid': True})['table']
        return wp['workpaper'], d1, d2


class NumbersTests(unittest.TestCase):
    def test_amounts_are_parsed_strictly(self):
        self.assertEqual(numbers.parse_decimal('1,234.50'), Decimal('1234.50'))
        self.assertEqual(numbers.parse_decimal('（1,000）'), Decimal('-1000'))
        self.assertEqual(numbers.parse_decimal('-'), Decimal(0))
        self.assertIs(numbers.parse_decimal('12万'), numbers.INVALID)
        self.assertIs(numbers.parse_decimal(''), numbers.BLANK)

    def test_dates(self):
        self.assertEqual(str(numbers.parse_date('2025年3月1日')), '2025-03-01')
        self.assertEqual(str(numbers.parse_date('20250301')), '2025-03-01')
        self.assertIs(numbers.parse_date('2025-13-02'), numbers.INVALID)

    def test_functions(self):
        self.assertEqual(query.name_norm('甲公司（北京）有限公司'), query.name_norm('甲公司(北京)有限公司'))
        self.assertGreater(query.name_sim('上海甲科技有限公司', '上海甲科技公司'), 0.7)
        self.assertEqual(query.first_digit('-0.0345'), 3)


class WorkpaperTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p = Project()
        cls.wp = cls.p.ok('workpaper_parse', {'file': '应收账款底稿.xlsx', 'sheets': None})
        cls.W = cls.wp['workpaper']

    @classmethod
    def tearDownClass(cls):
        cls.p.close()

    def test_regions_meta_and_inputs(self):
        sheet = self.wp['sheets'][0]
        regions = {r['range']: r for r in sheet['regions']}
        aging, top, notes = regions['应收账款!A5:E11'], regions['应收账款!A14:D20'], regions['应收账款!A22:E25']
        self.assertEqual((aging['type'], aging['headerRows']), ('fixed_table', [5, 6]))
        self.assertEqual([r['text'] for r in aging['rowLabels']], ['1年以内', '1-2年', '2-3年', '3年以上', '合计'])
        self.assertEqual(aging['columns'][0]['path'], ['期末余额', '账面余额'])
        self.assertEqual((top['type'], top['caption']), ('detail_table', '期末余额前五名'))
        self.assertEqual(notes['type'], 'text')
        kinds = {m['kind'] for m in sheet['meta']}
        self.assertTrue({'entity', 'unit', 'period'} <= kinds)
        self.assertTrue(self.wp['fingerprint'].startswith('TPL-'))

    def test_region_view_shows_formulas_and_inputs(self):
        text = self.p.ok('workpaper_region_view', {'region': self.W + '#R1'})['text']
        self.assertIn('B7 [待填]', text)
        self.assertIn('B11 =SUM(B7:B10)', text)

    def test_semantics_check_catches_wrong_header_and_unassigned(self):
        good = [{'region': self.W + '#R1', 'columns': [{'col': 'B', 'measure': '期末账面余额', 'basisRef': '应收账款!B6', 'basis': 'template', 'unit': '元'}],
                 'rows': [{'row': 7, 'text': '1年以内'}], 'manual': ['C7:C10', 'E7:E10']}]
        self.assertEqual(self.p.ok('workpaper_semantics_check', {'workpaper': self.W, 'regions': good})['errors'], [])
        bad = [{'region': self.W + '#R1', 'columns': [{'col': 'B', 'measure': '期初余额', 'basisRef': '应收账款!B6', 'basis': 'template', 'unit': '万元'}],
                'rows': [{'row': 8, 'text': '1年以内'}]}]
        codes = {e['code'] for e in self.p.ok('workpaper_semantics_check', {'workpaper': self.W, 'regions': bad})['errors']}
        self.assertTrue({'HEADER_MISMATCH', 'HEADER_UNIT_MISMATCH', 'REF_TEXT_MISMATCH', 'HEADER_UNASSIGNED'} <= codes, codes)

    def test_template_save_and_match(self):
        self.assertFalse(self.p.ok('template_match', {'fingerprint': self.wp['fingerprint']})['known'])
        semantics = [{'region': self.W + '#R2', 'type': 'detail_table', 'columns': [{'col': 'B', 'measure': '客户名称'}, {'col': 'C', 'measure': '期末余额'}]}]
        self.p.ok('template_save', {'workpaper': self.W, 'regions': semantics})
        match = self.p.ok('template_match', {'fingerprint': self.wp['fingerprint']})
        self.assertTrue(match['known'])
        self.assertEqual(match['regions'][0]['region'], self.W + '#R2')


class DataAndQueryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p = Project(invalid=True)
        cls.journal = cls.p.ok('data_inspect', {'file': '序时账2025.xlsx'})['tables'][0]

    @classmethod
    def tearDownClass(cls):
        cls.p.close()

    def test_invalid_values_are_not_zero(self):
        env = self.p.call('data_import', {'candidate': self.journal['candidate'], 'kind': 'journal', 'fields': JOURNAL_FIELDS})
        self.assertFalse(env['ok'])
        self.assertEqual(env['error']['code'], 'DATA_INVALID_VALUES')
        self.assertEqual(env['error']['detail']['invalid'][0]['samples'][0]['value'], '12万')
        data = self.p.ok('data_import', {'candidate': self.journal['candidate'], 'kind': 'journal', 'fields': JOURNAL_FIELDS, 'excludeInvalid': True})
        self.assertEqual((data['rows'], data['excluded']), (22, 1))
        quality = self.p.ok('data_quality', {'table': data['table']})
        self.assertTrue(quality['passed'])
        view = self.p.ok('table_view', {'table': data['table'], 'locate': 6, 'limit': 2})
        self.assertEqual(view['highlight'], 6)
        self.assertEqual(view['rows'][0]['source']['sheet'], '序时账')

    def test_query_rules(self):
        table = self.p.ok('data_import', {'candidate': self.journal['candidate'], 'kind': 'journal', 'fields': JOURNAL_FIELDS, 'excludeInvalid': True})['table']
        name = table.replace('-', '')
        cases = {
            'DELETE FROM %s' % name: 'ARG_NOT_SELECT',
            'SELECT 1; SELECT 2': 'ARG_MULTIPLE_STATEMENTS',
            "SELECT * FROM read_csv('a.csv')": 'ARG_FUNCTION_NOT_ALLOWED',
            "SELECT * FROM 'a.csv'": 'REF_UNDECLARED_TABLE',
            'SELECT * FROM D999': 'REF_UNDECLARED_TABLE',
            'SELECT 不存在 FROM %s' % name: 'REF_UNKNOWN_COLUMN',
        }
        for sql, code in cases.items():
            env = self.p.call('query_validate', {'sql': sql, 'inputs': [table]})
            self.assertFalse(env['ok'], sql)
            self.assertEqual(env['error']['code'], code, sql)
            self.assertTrue(env['error']['retryable'])
        run = self.p.ok('query_run', {'sql': AGING_SQL.replace('D002', name), 'inputs': [table], 'purpose': '账龄'})
        self.assertEqual(run['head'], [[k, '%.2f' % v] for k, v in fx.AGING.items()])
        self.assertEqual(run['sums'], {'金额': '1835000.00'})


class FillTests(unittest.TestCase):
    """完整流程：口径核对 → 编译 → 写入 → 回读 → 追溯；最后模拟服务重启。"""

    @classmethod
    def setUpClass(cls):
        cls.p = Project()
        cls.W, cls.D1, cls.D2 = cls.p.imported()
        sql = AGING_SQL.replace('D002', cls.D2.replace('-', ''))
        cls.R1 = cls.p.ok('query_run', {'sql': sql, 'inputs': [cls.D2], 'purpose': '账龄'})['result']
        cls.R2 = cls.p.ok('query_run', {'sql': TOP_SQL.replace('D002', cls.D2.replace('-', '')), 'inputs': [cls.D2]})['result']
        cls.basis = {'sources': [{'table': cls.D2}], 'period': {'value': '2025-01-01 至 2025-12-31', 'basis': 'confirmed'},
                     'scope': {'include': '科目 1122', 'exclude': [], 'basis': 'inferred'}, 'amount': {'column': '借方金额-贷方金额', 'direction': '借方', 'basis': 'inferred'},
                     'grouping': {'description': '按凭证日期起算账龄', 'basis': 'confirmed'}, 'unit': {'scale': '1', 'basis': 'template'},
                     'control': {'label': '账龄合计 = 余额表 1122 期末', 'column': '金额', 'right': {'kind': 'balance', 'table': cls.D1, 'account': '1122', 'field': 'closing'}, 'basis': 'data'},
                     'steps': [{'sql': sql}]}
        cls.bindings = [
            {'region': cls.W + '#R1', 'cells': [{'cell': '应收账款!B%d' % (7 + i), 'from': {'result': cls.R1, 'row': i + 1, 'column': '金额'}, 'scale': '1'} for i in range(4)],
             'manual': ['应收账款!C7:C10', '应收账款!E7:E10']},
            {'region': cls.W + '#R2', 'cells': [c for i in range(5) for c in (
                {'cell': 'B%d' % (15 + i), 'from': {'result': cls.R2, 'row': i + 1, 'column': '客户'}},
                {'cell': 'C%d' % (15 + i), 'from': {'result': cls.R2, 'row': i + 1, 'column': '余额'}, 'scale': '1'})]},
        ]

    @classmethod
    def tearDownClass(cls):
        cls.p.close()

    def test_checks(self):
        self.assertEqual(self.p.ok('check_basis', {'region': self.W + '#R1', 'basis': self.basis}), {'passed': True, 'missing': []})
        missing = self.p.ok('check_basis', {'region': self.W + '#R9', 'basis': {**self.basis, 'period': {'value': '2025'}}})['missing']
        self.assertIn('期间的依据', missing)
        control = self.p.ok('check_control', {'region': self.W + '#R1', 'left': {'result': self.R1, 'column': '金额', 'op': 'sum'},
                                              'right': {'kind': 'balance', 'table': self.D1, 'account': '1122', 'field': 'closing'}, 'tolerance': '0'})
        self.assertEqual((control['passed'], control['left'], control['right']), (True, '1835000.00', '1835000.00'))
        conservation = self.p.ok('check_conservation', {'from': self.D2, 'result': self.R1, 'amount': '金额', 'excluded': None,
                                                        'sourceAmount': '"借方金额" - "贷方金额"', 'sourceFilter': "\"科目编码\" = '1122'"})
        self.assertTrue(conservation['passed'])
        wrong = self.p.ok('query_run', {'sql': 'SELECT \'全部\' AS 账龄, SUM("借方金额") AS 金额 FROM %s WHERE "科目编码" = \'1122\'' % self.D2.replace('-', ''), 'inputs': [self.D2]})['result']
        compare = self.p.ok('check_compare', {'left': self.R1, 'right': wrong, 'mapping': {'columns': {'金额': '金额'}}})
        self.assertFalse(compare['passed'])

    def test_fill_end_to_end(self):
        self.p.ok('check_basis', {'region': self.W + '#R1', 'basis': self.basis})
        plan = self.p.ok('fill_compile', {'task': TASK, 'workpaper': self.W, 'bindings': self.bindings, 'overwrite': []})
        self.assertEqual((plan['cells'], plan['manual'], plan['overwrites'], plan['expand']), (14, 8, [], []))
        self.assertEqual(plan['output'], '输出/%s/应收账款底稿-填写成果.xlsx' % TASK)
        preview = self.p.ok('fill_preview', {'plan': plan['plan'], 'range': '应收账款!B7:B8'})['cells']
        self.assertEqual([(c['cell'], c['value'], c['from']) for c in preview], [('应收账款!B7', '1200000', self.R1 + '#1'), ('应收账款!B8', '435000', self.R1 + '#2')])

        self.assertEqual(self.p.call('fill_apply', {'task': TASK, 'plan': plan['plan']})['error']['code'], 'NEEDS_INPUT')
        self.assertEqual(self.p.call('fill_apply', {'task': 'T-other', 'plan': plan['plan'], 'approvedCard': 'C-0006'})['error']['code'], 'ARG_TASK_MISMATCH')
        applied = self.p.ok('fill_apply', {'task': TASK, 'plan': plan['plan'], 'approvedCard': 'C-0006', 'excelRecalc': True})
        self.assertEqual(applied['checks'], {'readback': 'passed', 'formulas': 'passed', 'crossFoot': 'passed', 'untouched': 'passed',
                                             'completeness': 'passed', 'excel': 'skipped' if sys.platform != 'win32' else applied['checks']['excel']})
        self.assertEqual(len(applied['manualCells']), 8)
        self.assertTrue((self.p.root / applied['report']).is_file())
        self.assertTrue((self.p.root / '应收账款底稿.xlsx').is_file())

        values = openpyxl.load_workbook(self.p.root / applied['output'], data_only=True)['应收账款']
        self.assertEqual([values['B%d' % r].value for r in range(7, 12)], [1200000, 435000, 150000, 50000, 1835000])
        self.assertAlmostEqual(values['D7'].value, 1200000 / 1835000)
        self.assertEqual([values['B%d' % r].value for r in range(15, 20)], ['甲公司', '乙公司', '丙公司', '戊公司', '丁公司'])
        self.assertEqual(values['C20'].value, 1635000)
        formulas = openpyxl.load_workbook(self.p.root / applied['output'])['应收账款']
        self.assertEqual((formulas['B11'].value, formulas['D7'].value), ('=SUM(B7:B10)', '=IF($B$11=0,0,B7/$B$11)'))
        self.assertIn('A23:E25', [str(m) for m in formulas.merged_cells.ranges])

        trace = self.p.ok('trace_cell', {'output': applied['output'], 'cell': '应收账款!B7'})
        self.assertEqual((trace['value'], trace['text'], trace['from']['result']), ('1200000', '1,200,000.00', self.R1))
        self.assertEqual(trace['source']['table'], self.D2)
        self.assertIn({'label': '期间', 'value': '2025-01-01 至 2025-12-31', 'basis': 'confirmed'}, trace['basis'])
        self.assertEqual(self.p.ok('trace_cell', {'output': applied['output'], 'cell': '应收账款!C7'}), {'cell': '应收账款!C7', 'from': None, 'manual': True})

        # 服务重启后，之前的编号和结果仍然可用。
        self.p.restart()
        self.assertEqual(self.p.ok('result_get', {'result': self.R1})['rows'][0], ['1年以内', '1200000.00'])
        self.assertEqual(self.p.ok('trace_cell', {'output': applied['output'], 'cell': '应收账款!C15'})['value'], '500000')
        self.assertEqual(self.p.ok('fill_preview', {'plan': plan['plan']})['cells'][0]['cell'], '应收账款!B7')
        later = self.p.ok('query_run', {'sql': 'SELECT COUNT(*) AS n FROM %s' % self.D2.replace('-', ''), 'inputs': [self.D2]})['result']
        self.assertGreater(int(later.split('-')[1]), int(self.R2.split('-')[1]), '重启后编号继续递增，不会重复')

    def test_compile_errors_overwrite_and_expand(self):
        bad = [{'region': self.W + '#R1', 'cells': [
            {'cell': 'D7', 'from': {'result': self.R1, 'row': 1, 'column': '金额'}},
            {'cell': 'H7', 'from': {'result': self.R1, 'row': 1, 'column': '金额'}},
            {'cell': 'B8', 'from': {'result': self.R1, 'row': 9, 'column': '金额'}},
            {'cell': 'B9', 'from': {'result': self.R1, 'row': 1, 'column': '不存在'}},
            {'cell': 'B10', 'from': {'value': '1'}},
        ]}]
        env = self.p.call('fill_compile', {'task': TASK, 'workpaper': self.W, 'bindings': bad})
        self.assertEqual([e['code'] for e in env['error']['detail']['errors']],
                         ['REF_FORMULA_CELL', 'REF_OUTSIDE_REGION', 'REF_ROW_OUT_OF_RANGE', 'REF_UNKNOWN_COLUMN', 'ARG_INVALID'])

        longer = [{'region': self.W + '#R2', 'cells': [{'cell': 'C%d' % (15 + i), 'from': {'result': self.R2, 'row': min(i + 1, 5), 'column': '余额'}} for i in range(7)]}]
        plan = self.p.ok('fill_compile', {'task': TASK, 'workpaper': self.W, 'bindings': longer})
        self.assertEqual(plan['expand'], [{'region': self.W + '#R2', 'rows': 2, 'after': 19}])
        self.assertEqual(self.p.call('fill_apply', {'task': TASK, 'plan': plan['plan'], 'approvedCard': 'C-1'})['error']['code'], 'NEEDS_INPUT')

        book = openpyxl.load_workbook(self.p.root / '应收账款底稿.xlsx')
        book['应收账款']['B7'] = 999
        book.save(self.p.root / '已有数字的底稿.xlsx')
        W = self.p.ok('workpaper_parse', {'file': '已有数字的底稿.xlsx'})['workpaper']
        cells = [{'region': W + '#R1', 'cells': [{'cell': 'B%d' % (7 + i), 'from': {'result': self.R1, 'row': i + 1, 'column': '金额'}, 'scale': '10000'} for i in range(4)]}]
        plan = self.p.ok('fill_compile', {'task': TASK, 'workpaper': W, 'bindings': cells})
        self.assertEqual(plan['overwrites'], [{'cell': '应收账款!B7', 'old': '999', 'value': '120'}])
        self.assertEqual(self.p.call('fill_apply', {'task': TASK, 'plan': plan['plan'], 'approvedCard': 'C-1'})['error']['code'], 'NEEDS_INPUT')
        plan = self.p.ok('fill_compile', {'task': TASK, 'workpaper': W, 'bindings': cells, 'overwrite': ['应收账款!B7']})
        self.assertEqual(plan['overwrites'], [])
        applied = self.p.ok('fill_apply', {'task': TASK, 'plan': plan['plan'], 'approvedCard': 'C-1'})
        self.assertEqual(applied['checks']['crossFoot'], 'passed')
        values = openpyxl.load_workbook(self.p.root / applied['output'], data_only=True)['应收账款']
        self.assertEqual([values['B%d' % r].value for r in range(7, 12)], [120, 43.5, 15, 5, 183.5])


class ServerTests(unittest.TestCase):
    def test_new_tools_run_in_server_process(self):
        p = Project()
        try:
            init = {'jsonrpc': '2.0', 'id': 1, 'method': 'initialize', 'params': {'protocolVersion': '2025-06-18', 'capabilities': {}, 'clientInfo': {'name': 't', 'version': '1'}}}
            calls = [init,
                     {'jsonrpc': '2.0', 'id': 2, 'method': 'tools/call', 'params': {'name': 'workpaper_parse', 'arguments': {'file': '应收账款底稿.xlsx'}}},
                     {'jsonrpc': '2.0', 'id': 3, 'method': 'tools/call', 'params': {'name': 'workpaper_region_view', 'arguments': {'region': 'W-404#R1'}}}]
            env = {**os.environ, 'AUDIT_PROJECT_ROOT': str(p.root), 'AUDIT_STATE_DIR': str(p.state)}
            lines = ''.join(json.dumps(m, ensure_ascii=False) + '\n' for m in calls)
            out = subprocess.run([sys.executable, '-X', 'utf8', str(ROOT / 'src' / 'audit-mcp' / 'server.py')], input=lines.encode('utf-8'), capture_output=True, env=env, timeout=120)
            replies = {r['id']: r for r in map(json.loads, out.stdout.decode('utf-8').splitlines()) if 'id' in r}
            parsed = replies[2]['result']
            self.assertFalse(parsed['isError'])
            self.assertEqual(parsed['structuredContent']['data']['workpaper'], 'W-001')
            self.assertTrue((p.state / 'tools' / 'objects' / 'W-001.json').is_file(), '状态放在宿主给的目录里')
            self.assertFalse((p.root / '.audit-agent').exists())
            missing = replies[3]['result']
            self.assertTrue(missing['isError'])
            self.assertEqual(missing['structuredContent']['error']['code'], 'REF_NOT_FOUND')
        finally:
            p.close()


if __name__ == '__main__':
    unittest.main()
