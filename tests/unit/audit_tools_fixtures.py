"""审计工具测试用的合成底稿和资料：应收账款账龄、前五名客户，余额表和序时账。

数字设计成能手算核对：应收账款期末 1,835,000.00 = 账龄四段之和 = 序时账 1122 借方减贷方。
"""
import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side

THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 每个客户的应收发生额：(日期, 客户, 借方, 贷方)。账龄按凭证日期算到 2025-12-31。
SALES = [
    ('2025-12-31', '甲公司', '500.00', '0'),
    ('2025-06-30', '甲公司', '499500.00', '0'),
    ('2025-03-15', '乙公司', '400000.00', '0'),
    ('2025-09-01', '丙公司', '300000.00', '0'),
    ('2024-08-20', '丁公司', '200000.00', '0'),
    ('2024-05-10', '戊公司', '235000.00', '0'),
    ('2023-07-01', '己公司', '150000.00', '0'),
    ('2021-02-01', '庚公司', '50000.00', '0'),
    ('2025-10-10', '辛公司', '0', '0'),
    ('2025-11-11', '壬公司', '100000.00', '100000.00'),
]
AGING = {'1年以内': Decimal('1200000.00'), '1-2年': Decimal('435000.00'), '2-3年': Decimal('150000.00'), '3年以上': Decimal('50000.00')}
CLOSING = Decimal('1835000.00')


def workpaper(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '应收账款'
    ws['A1'] = '应收账款账龄分析表'
    ws.merge_cells('A1:E1')
    ws['A1'].font = Font(bold=True)
    ws['A2'] = '被审计单位：'
    ws['B2'].border = BOX
    ws['D2'] = '单位：元'
    ws['A3'] = '截止日：2025年12月31日'
    ws['A5'], ws['B5'], ws['D5'], ws['E5'] = '账龄', '期末余额', '占比', '备注'
    ws['B6'], ws['C6'] = '账面余额', '坏账准备'
    for span in ('A5:A6', 'B5:C5', 'D5:D6', 'E5:E6'):
        ws.merge_cells(span)
    for row, label in enumerate(AGING, start=7):
        ws['A%d' % row] = label
        ws['D%d' % row] = '=IF($B$11=0,0,B%d/$B$11)' % row
    ws['A11'] = '合计'
    ws['B11'], ws['C11'], ws['D11'] = '=SUM(B7:B10)', '=SUM(C7:C10)', '=SUM(D7:D10)'
    for row in range(5, 12):
        for col in 'ABCDE':
            ws['%s%d' % (col, row)].border = BOX
    ws['E11'].border = Border()
    ws['A13'] = '期末余额前五名'
    ws['A14'], ws['B14'], ws['C14'], ws['D14'] = '序号', '客户名称', '期末余额', '占应收账款比例'
    for i, row in enumerate(range(15, 20), start=1):
        ws['A%d' % row] = i
        ws['D%d' % row] = '=IF($B$11=0,0,C%d/$B$11)' % row
    ws['A20'] = '合计'
    ws['C20'], ws['D20'] = '=SUM(C15:C19)', '=SUM(D15:D19)'
    for row in range(14, 21):
        for col in 'ABCD':
            ws['%s%d' % (col, row)].border = BOX
    ws['A22'] = '审计说明：'
    ws.merge_cells('A23:E25')
    for row in range(23, 26):
        for col in 'ABCDE':
            ws['%s%d' % (col, row)].border = BOX
    ws.freeze_panes = 'B7'
    wb.save(path)


def balance(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '科目余额表'
    ws.append(['科目编码', '科目名称', '期初余额', '本期借方', '本期贷方', '期末余额'])
    ws.append(['1001', '库存现金', 5000.0, 12000.0, 15000.0, 2000.0])
    ws.append(['1122', '应收账款', 635000.0, 1300000.0, 100000.0, 1835000.0])
    ws.append(['2202', '应付账款', 300000.0, 50000.0, 70000.0, 320000.0])
    ws.append(['合计', None, None, 1362000.0, 185000.0, None])
    wb.save(path)


def journal(path, invalid=False):
    wb = Workbook()
    ws = wb.active
    ws.title = '序时账'
    ws.append(['日期', '凭证号', '科目编码', '摘要', '客户', '借方金额', '贷方金额'])
    # 1122 期初 635,000.00 由以前年度发生额构成；本年借方 1,300,000.00、贷方 100,000.00。
    n = 0
    for date, customer, debit, credit in SALES:
        n += 1
        ws.append([datetime.datetime.strptime(date, '%Y-%m-%d'), '记-%03d' % n, '1122', '销售商品', customer, float(debit), float(credit)])
        ws.append([datetime.datetime.strptime(date, '%Y-%m-%d'), '记-%03d' % n, '6001', '销售商品', customer, float(credit), float(debit)])
    ws.append([datetime.datetime(2025, 3, 1), '记-%03d' % (n + 1), '1001', '提现', None, 12000.0, 0.0])
    ws.append([datetime.datetime(2025, 3, 1), '记-%03d' % (n + 1), '1002', '提现', None, 0.0, 12000.0])
    if invalid:
        ws.append([datetime.datetime(2025, 4, 1), '记-%03d' % (n + 2), '6602', '办公费', None, '12万', 0.0])
    wb.save(path)


def build(root, invalid=False):
    workpaper(root / '应收账款底稿.xlsx')
    balance(root / '科目余额表2025.xlsx')
    journal(root / '序时账2025.xlsx', invalid=invalid)
